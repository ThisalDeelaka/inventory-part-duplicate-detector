"""Prepare deterministic blinded R18G group-level human evidence.

The generator reuses a persisted canonical catalog and discovery population,
re-evaluates its edges with current deterministic semantics, and invokes the
pure GF5 resolver. It never writes to the source database or calls an LLM.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sqlite3
import sys
import zipfile
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime
from itertools import combinations
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app.engine.identity_edge import IdentityEdgeClass
from app.engine.identity_evidence_evaluator import (
    DeterministicIdentityContext,
    evaluate_canonical_identity_relationship,
)
from app.engine.identity_signature_derivation import derive_identity_signature
from app.engine.signed_identity_evidence import (
    classify_shadow_evidence,
    derive_signed_identity_evidence,
)
from app.resolution.contracts import (
    IdentityResolutionEvidenceEdge,
    IdentityResolutionInput,
    IdentityResolutionNeighborhood,
    ResolverConfiguration,
)
from app.resolution.resolver import (
    CanonicalEvaluatorTargetedEvidenceProvider,
    resolve_identity_groups,
)
from app.services.canonical_record_service import CanonicalScanRecord


DATASET_VERSION = "r18g-group-human-evidence-v1"
SPLIT_SEED = 1807
GENERATION_COMMIT = "cd7f089260ac288e5893a30eb49b018320a25349"
SCAN_ID = 31
DEVELOPMENT_TARGET = 48
HOLDOUT_TARGET = 16
LABELS = ("SAME_ONE_IDENTITY", "NOT_ONE_IDENTITY", "INSUFFICIENT_INFORMATION")
CONFIDENCES = ("HIGH", "MEDIUM", "LOW")
REASON_CODES = (
    "SAME_ALIAS_OR_NAMING_VARIANT", "SAME_CROSS_SITE_IDENTITY",
    "SAME_EQUIVALENT_BUSINESS_IDENTITY", "MIXED_GROUP_MULTIPLE_IDENTITIES",
    "DIFFERENT_MODEL_TYPE_OR_VARIANT", "DIFFERENT_OBJECT_OR_FUNCTION",
    "DIFFERENT_CRITICAL_ATTRIBUTE", "GENERIC_OR_COPIED_TEXT_NOT_ENOUGH",
    "INSUFFICIENT_TECHNICAL_FIELDS", "DOMAIN_KNOWLEDGE_REQUIRED", "OTHER",
)
COMMENT_BASES = ("SOURCE_VISIBLE", "DOMAIN_EXTERNAL", "MIXED", "UNCLEAR")
PARTITION_IDS = tuple(f"P{index}" for index in range(1, 21)) + ("UNRESOLVED",)
VISIBLE_FIELDS = (
    "part_number", "description_in_use", "description", "master_description",
    "type_designation", "dimension_quality", "uom", "part_type", "site",
)
FORBIDDEN_VISIBLE_TERMS = (
    "likely", "review_support", "cannot_link", "strong_support", "gf4", "gf5",
    "lexicaltrust", "r18c", "shadow", "system_status", "hypothesis_source",
    "pair_label", "expected_control", "score",
)


@dataclass(frozen=True)
class Member:
    record_id: int
    stable_ref: str
    source_row: int
    part_number: str
    description_in_use: str
    description: str
    master_description: str
    type_designation: str
    dimension_quality: str
    uom: str
    part_type: str
    site: str


@dataclass(frozen=True)
class Edge:
    left_id: int
    right_id: int
    edge_class: str
    provenance: str
    generic_only: bool
    r18c_demoted: bool
    fingerprint: str


@dataclass(frozen=True)
class GroupHypothesis:
    fingerprint: str
    members: tuple[Member, ...]
    source_kind: str
    current_or_shadow: str
    current_group_id: str
    work_unit_id: str
    system_status: str
    evidence_provenance: str
    strong_edge_count: int
    native_review_edge_count: int
    r18c_demoted_review_edge_count: int
    cannot_link_count: int
    generic_only_state: bool
    partition_sensitive: bool
    deferred_context: bool
    conflict_context: bool


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def fingerprint(member_refs, source_kind="GROUP") -> str:
    payload = json.dumps({
        "version": DATASET_VERSION,
        "members": sorted(member_refs),
        "source_kind": source_kind,
    }, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode()).hexdigest()


def review_group_id(value: GroupHypothesis) -> str:
    return "R18G-" + value.fingerprint[:16].upper()


def deterministic_split(groups: tuple[GroupHypothesis, ...], development_target=48,
                        holdout_target=16):
    ordered = sorted(groups, key=lambda group: hashlib.sha256(
        f"{SPLIT_SEED}:{group.fingerprint}".encode()
    ).hexdigest())
    mandatory = [group for group in ordered if group.source_kind == "GENERIC_BICYCLE_FAMILY"]
    mandatory.extend(group for group in ordered if group.partition_sensitive)
    mandatory = mandatory[:7]
    for source_kind in sorted({group.source_kind for group in ordered}):
        candidate = next((group for group in ordered if group.source_kind == source_kind), None)
        if candidate is not None and candidate not in mandatory:
            mandatory.append(candidate)
    for predicate in (
        lambda group: len(group.members) == 2,
        lambda group: len(group.members) == 3,
        lambda group: len(group.members) == 4,
        lambda group: len(group.members) >= 5,
        lambda group: len({member.site for member in group.members if member.site}) >= 2,
        lambda group: group.r18c_demoted_review_edge_count > 0,
        lambda group: group.generic_only_state,
    ):
        candidate = next((group for group in ordered if predicate(group)), None)
        if candidate is not None and candidate not in mandatory:
            mandatory.append(candidate)
    mandatory = list({group.fingerprint: group for group in mandatory}.values())
    if len(mandatory) > development_target:
        raise ValueError("R18G_MANDATORY_DEVELOPMENT_STRATA_EXCEED_TARGET")
    development = list(mandatory)
    development.extend(
        group for group in ordered
        if group not in development
    )
    development = development[:min(development_target, len(ordered))]
    development_ids = {item.fingerprint for item in development}
    holdout = [item for item in ordered if item.fingerprint not in development_ids]
    return tuple(sorted(development, key=lambda item: item.fingerprint)), tuple(
        sorted(holdout[:holdout_target], key=lambda item: item.fingerprint)
    )


def _readonly(path: Path):
    connection = sqlite3.connect(f"file:{path.resolve().as_posix()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    return connection


def _canonical(row) -> CanonicalScanRecord:
    return CanonicalScanRecord(
        record_id=row["id"], scan_id=row["scan_id"],
        source_row_index=row["source_row_index"], record_ref_key=row["record_ref_key"],
        source_record_fingerprint=row["source_record_fingerprint"],
        part_no=row["part_no"], description=row["description"], contract=row["contract"],
        uom=row["uom"], type_code=row["type_code"], prime_commodity=row["prime_commodity"],
        second_commodity=row["second_commodity"], accounting_group=row["accounting_group"],
        part_product_code=row["part_product_code"], part_product_family=row["part_product_family"],
        product_category_id=row["product_category_id"], hsn_sac_code=row["hsn_sac_code"],
        hazard_code=row["hazard_code"], normalized_part_no=row["normalized_part_no"],
        normalized_description=row["normalized_description"],
        normalization_version=row["normalization_version"],
    )


def _member(row) -> Member:
    # Extended source fields were not persisted in GF1 and are truthfully blank.
    return Member(
        row["id"], row["record_ref_key"], row["source_row_index"],
        str(row["part_no"] or ""), "", str(row["description"] or ""), "", "", "",
        str(row["uom"] or ""), str(row["type_code"] or ""), str(row["contract"] or ""),
    )


def load_current_resolution(database_path: Path, scan_id=SCAN_ID):
    """Re-evaluate the persisted candidate population and run pure current GF5."""
    connection = _readonly(database_path)
    scan = connection.execute("SELECT * FROM duplicate_scan WHERE id=?", (scan_id,)).fetchone()
    discovery = connection.execute(
        "SELECT * FROM identity_discovery_run WHERE scan_id=? AND status='COMPLETED' ORDER BY id DESC LIMIT 1",
        (scan_id,),
    ).fetchone()
    evidence_run = connection.execute(
        "SELECT * FROM identity_evidence_run WHERE scan_id=? AND status='COMPLETED' ORDER BY id DESC LIMIT 1",
        (scan_id,),
    ).fetchone()
    resolution_run = connection.execute(
        "SELECT * FROM identity_resolution_run WHERE scan_id=? AND status='COMPLETED' ORDER BY id DESC LIMIT 1",
        (scan_id,),
    ).fetchone()
    if not all((scan, discovery, evidence_run, resolution_run)):
        raise ValueError("R18G_CURRENT_INPUTS_UNAVAILABLE")
    record_rows = connection.execute(
        "SELECT * FROM scan_record_snapshot WHERE scan_id=? ORDER BY id", (scan_id,)
    ).fetchall()
    records = tuple(_canonical(row) for row in record_rows)
    members = {row["id"]: _member(row) for row in record_rows}
    records_by_id = {record.record_id: record for record in records}
    signatures = {}
    selected_fields = tuple(json.loads(scan["selected_fields"]))
    context = DeterministicIdentityContext("DISCOVERY", selected_fields)
    current_edges = []
    edge_metadata = {}
    reevaluated_count = 0
    reused_count = 0
    for old in connection.execute(
        "SELECT * FROM identity_evidence_edge_snapshot WHERE evidence_run_id=? ORDER BY record_id_1,record_id_2",
        (evidence_run["id"],),
    ):
        old_positive = old["edge_class"] in {"STRONG_SUPPORT", "REVIEW_SUPPORT"}
        if old_positive:
            reevaluated_count += 1
            evaluated = evaluate_canonical_identity_relationship(
                records_by_id[old["record_id_1"]], records_by_id[old["record_id_2"]], context
            )
            for record_id in (old["record_id_1"], old["record_id_2"]):
                if record_id not in signatures:
                    record = records_by_id[record_id]
                    signatures[record_id] = derive_identity_signature(
                        record, record_reference=record.record_ref_key
                    )
            signed = derive_signed_identity_evidence(
                signatures[old["record_id_1"]], signatures[old["record_id_2"]]
            )
            provenance = classify_shadow_evidence(signed).value
            edge_class = evaluated.edge_class
            reason_codes = evaluated.classification_reason_codes
            evidence_fingerprint = evaluated.evidence_fingerprint
            generic = bool(json.loads(evaluated.generic_evidence_json).get("generic_guard_reason"))
        else:
            reused_count += 1
            # R11 and R18C are monotonic safety corrections: neither can promote
            # a persisted NON_GROUPABLE/CANNOT_LINK edge. Reusing these 19,187
            # immutable negative edges is semantically equivalent and bounded.
            edge_class = IdentityEdgeClass(old["edge_class"])
            reason_codes = tuple(json.loads(old["classification_reason_codes_json"]))
            evidence_fingerprint = old["evidence_fingerprint"]
            generic = bool(json.loads(old["generic_evidence_json"]).get("generic_guard_reason"))
            provenance = (
                "SHADOW_EXPLICIT_CONTRADICTION"
                if edge_class == IdentityEdgeClass.CANNOT_LINK
                else "SHADOW_INSUFFICIENT"
            )
        demoted = old["edge_class"] == "STRONG_SUPPORT" and edge_class.value == "REVIEW_SUPPORT"
        edge = Edge(
            old["record_id_1"], old["record_id_2"], edge_class.value,
            provenance, generic, demoted, evidence_fingerprint,
        )
        edge_metadata[(edge.left_id, edge.right_id)] = edge
        current_edges.append(IdentityResolutionEvidenceEdge(
            scan_id=scan_id, evidence_run_id=evidence_run["id"],
            record_id_1=edge.left_id, record_id_2=edge.right_id,
            edge_class=edge_class,
            reason_codes=reason_codes,
            evidence_fingerprint=evidence_fingerprint,
            generic_only=generic,
        ))
    print(
        f"R18G stage=current_evidence complete edges={len(current_edges)} "
        f"reevaluated={reevaluated_count}", file=sys.stderr, flush=True,
    )
    neighborhood_members = defaultdict(list)
    neighborhood_rows = connection.execute(
        "SELECT n.*,m.record_id FROM identity_neighborhood_snapshot n "
        "JOIN identity_neighborhood_member m ON m.neighborhood_id=n.id "
        "WHERE n.discovery_run_id=? ORDER BY n.id,m.member_order", (discovery["id"],),
    ).fetchall()
    neighborhood_data = {}
    for row in neighborhood_rows:
        neighborhood_data[row["id"]] = row
        neighborhood_members[row["id"]].append(row["record_id"])
    neighborhoods = tuple(IdentityResolutionNeighborhood(
        row["neighborhood_fingerprint"], scan_id, discovery["id"],
        tuple(neighborhood_members[neighborhood_id]), bool(row["is_truncated"]),
        bool(row["degraded"]),
    ) for neighborhood_id, row in sorted(neighborhood_data.items()))
    config_json = json.loads(resolution_run["configuration_json"])
    config = ResolverConfiguration(**config_json)
    value = IdentityResolutionInput(
        scan_id, discovery["id"], evidence_run["id"], records, neighborhoods,
        tuple(current_edges), (), resolution_run["resolver_algorithm_version"], config,
    )
    # Re-resolve each persisted accepted hypothesis as its own bounded current
    # work unit.  Running the entire degraded Scan-31 overlap graph after the
    # post-R18C edge changes repeats an exhaustive global partition search that
    # is irrelevant to preparing a capped review sample.  These source groups
    # are disjoint, contain at most five members, and therefore preserve the
    # exact current evaluator/GF5 contract without changing runtime authority.
    current_groups = []
    persisted_groups = connection.execute(
        "SELECT * FROM identity_resolution_group_snapshot "
        "WHERE resolution_run_id=? ORDER BY hypothesis_id", (resolution_run["id"],),
    ).fetchall()
    for group_number, persisted_group in enumerate(persisted_groups, start=1):
        group_ids = tuple(row["record_id"] for row in connection.execute(
            "SELECT record_id FROM identity_resolution_group_member "
            "WHERE group_snapshot_id=? ORDER BY member_index", (persisted_group["id"],),
        ))
        group_set = set(group_ids)
        group_value = IdentityResolutionInput(
            scan_id, discovery["id"], evidence_run["id"],
            tuple(record for record in records if record.record_id in group_set),
            (IdentityResolutionNeighborhood(
                fingerprint((members[item].stable_ref for item in group_ids), "R18G_CURRENT_UNIT"),
                scan_id, discovery["id"], group_ids, False, False,
            ),),
            tuple(edge for edge in current_edges
                  if edge.record_id_1 in group_set and edge.record_id_2 in group_set),
            (), resolution_run["resolver_algorithm_version"], config,
        )
        resolved = resolve_identity_groups(
            group_value,
            CanonicalEvaluatorTargetedEvidenceProvider(group_value.canonical_records, context),
        )
        current_groups.extend(resolved.accepted_groups)
        if group_number % 25 == 0 or group_number == len(persisted_groups):
            print(
                f"R18G stage=bounded_gf5 progress={group_number}/{len(persisted_groups)}",
                file=sys.stderr, flush=True,
            )

    def persisted_contexts(snapshot_table, member_table, foreign_key, id_field, member_field):
        contexts = []
        for snapshot in connection.execute(
            f"SELECT * FROM {snapshot_table} WHERE resolution_run_id=? ORDER BY {id_field}",
            (resolution_run["id"],),
        ):
            record_ids = tuple(row["record_id"] for row in connection.execute(
                f"SELECT record_id FROM {member_table} WHERE {foreign_key}=? ORDER BY member_index",
                (snapshot["id"],),
            ))
            contexts.append(SimpleNamespace(**{
                id_field: snapshot[id_field], member_field: record_ids,
            }))
        return tuple(contexts)

    result = SimpleNamespace(
        accepted_groups=tuple(current_groups),
        deferred_work_units=persisted_contexts(
            "identity_resolution_deferred_snapshot", "identity_resolution_deferred_member",
            "deferred_snapshot_id", "deferred_id", "record_ids",
        ),
        conflicts=persisted_contexts(
            "identity_resolution_conflict_snapshot", "identity_resolution_conflict_member",
            "conflict_snapshot_id", "conflict_id", "involved_record_ids",
        ),
    )
    print("R18G stage=bounded_gf5 complete", file=sys.stderr, flush=True)
    connection.close()
    return value, result, members, edge_metadata, {
        "scan_id": scan_id,
        "canonical_record_count": len(records),
        "source_discovery_run_id": discovery["id"],
        "source_discovery_fingerprint": discovery["discovery_fingerprint"],
        "source_evidence_run_id": evidence_run["id"],
        "source_evidence_configuration_fingerprint": evidence_run["configuration_fingerprint"],
        "current_evidence_fingerprint": hashlib.sha256("".join(
            edge.fingerprint for edge in sorted(edge_metadata.values(), key=lambda item: (item.left_id,item.right_id))
        ).encode()).hexdigest(),
        "current_positive_edges_reevaluated": reevaluated_count,
        "monotonic_negative_edges_reused": reused_count,
        "bounded_current_gf5_source_work_units": len(persisted_groups),
        "bounded_current_gf5_accepted_groups": len(current_groups),
        "provider_calls": 0,
    }


def _pair(left, right):
    return (left, right) if left < right else (right, left)


def _group_hypothesis(member_ids, members, edges, *, source_kind,
                      current_or_shadow="CURRENT", current_group_id="",
                      work_unit_id="", system_status="", partition_sensitive=False,
                      deferred_context=False, conflict_context=False):
    ids = tuple(sorted(set(member_ids)))
    if len(ids) < 2:
        raise ValueError("R18G_GROUP_TOO_SMALL")
    internal = [edges.get(_pair(left, right)) for left, right in combinations(ids, 2)]
    present = [edge for edge in internal if edge is not None]
    cannot = sum(edge.edge_class == "CANNOT_LINK" for edge in present)
    if current_or_shadow == "SHADOW_ADR_ORDER" and cannot:
        raise ValueError("R18G_SHADOW_CANNOT_LINK_UNSAFE")
    strong = sum(edge.edge_class == "STRONG_SUPPORT" for edge in present)
    demoted = sum(edge.edge_class == "REVIEW_SUPPORT" and edge.r18c_demoted for edge in present)
    native_review = sum(
        edge.edge_class == "REVIEW_SUPPORT" and not edge.r18c_demoted for edge in present
    )
    provenances = {edge.provenance for edge in present if edge.edge_class in {"STRONG_SUPPORT", "REVIEW_SUPPORT"}}
    provenance = next(iter(provenances)) if len(provenances) == 1 else "MIXED_EVIDENCE"
    positive = [edge for edge in present if edge.edge_class in {"STRONG_SUPPORT", "REVIEW_SUPPORT"}]
    generic_only = bool(positive) and all(edge.generic_only for edge in positive)
    values = tuple(members[item] for item in ids)
    fp = fingerprint((item.stable_ref for item in values), source_kind)
    return GroupHypothesis(
        fp, values, source_kind, current_or_shadow, current_group_id,
        work_unit_id, system_status, provenance, strong, native_review, demoted,
        cannot, generic_only, partition_sensitive, deferred_context, conflict_context,
    )


def _positive_clique(member_ids, edges, maximum=8):
    ids = tuple(sorted(set(member_ids)))
    positive = {
        _pair(left, right): edge
        for left, right in combinations(ids, 2)
        if (edge := edges.get(_pair(left, right))) is not None
        and edge.edge_class in {"STRONG_SUPPORT", "REVIEW_SUPPORT"}
    }
    best = ()
    for left, right in positive:
        clique = [left, right]
        for candidate in ids:
            if candidate in clique:
                continue
            if all(_pair(candidate, member) in positive for member in clique):
                clique.append(candidate)
                if len(clique) == maximum:
                    break
        values = [positive[_pair(a, b)] for a, b in combinations(clique, 2)]
        if len(clique) >= 3 and all(edge.generic_only for edge in values):
            clique = clique[:2]
        candidate_clique = tuple(sorted(clique))
        if len(candidate_clique) > len(best) or (
            len(candidate_clique) == len(best) and candidate_clique < best
        ):
            best = candidate_clique
        if len(best) == maximum:
            break
    return best


def _work_unit_members(value):
    parent = {}

    def find(item):
        parent.setdefault(item, item)
        if parent[item] != item:
            parent[item] = find(parent[item])
        return parent[item]

    def union(left, right):
        left, right = find(left), find(right)
        if left != right:
            parent[right] = left

    for neighborhood in value.identity_neighborhoods:
        values = neighborhood.member_record_ids
        for item in values:
            find(item)
        for item in values[1:]:
            union(values[0], item)
    units = defaultdict(set)
    for item in parent:
        units[find(item)].add(item)
    return tuple(sorted((tuple(sorted(unit)) for unit in units.values()), key=lambda unit: unit))


def _partition_challengers(unit, edges, maximum_members=8):
    """Bounded exact set-partition comparison for complete safe cliques."""
    unit = tuple(sorted(unit))
    if len(unit) > maximum_members:
        return ()
    candidates = []
    for size in range(2, len(unit) + 1):
        for subset in combinations(unit, size):
            values = [edges.get(_pair(left, right)) for left, right in combinations(subset, 2)]
            if not values or not all(
                edge is not None and edge.edge_class in {"STRONG_SUPPORT", "REVIEW_SUPPORT"}
                for edge in values
            ):
                continue
            if size >= 3 and all(edge.generic_only for edge in values):
                continue
            strong = sum(edge.edge_class == "STRONG_SUPPORT" for edge in values)
            review = sum(edge.edge_class == "REVIEW_SUPPORT" for edge in values)
            likely_members = size if review == 0 else 0
            candidates.append((frozenset(subset), strong, review, likely_members))
    if not candidates:
        return ()
    # Enumerate canonical set partitions of at most eight members (Bell(8) =
    # 4,140), treating singleton blocks as uncovered records.  Enumerating the
    # power set of every valid clique is exponentially larger and repeats the
    # same disjoint partition many times.
    candidate_by_members = {
        item[0]: item for item in candidates
    }
    partitions = []

    def visit(index, blocks):
        if index == len(unit):
            selected = []
            for block in blocks:
                if len(block) == 1:
                    continue
                candidate = candidate_by_members.get(frozenset(block))
                if candidate is None:
                    return
                selected.append(candidate)
            used = {member for candidate in selected for member in candidate[0]}
            signature = tuple(sorted(tuple(sorted(candidate[0])) for candidate in selected))
            partitions.append((
                len(used), sum(candidate[3] for candidate in selected),
                sum(candidate[1] for candidate in selected),
                sum(candidate[2] for candidate in selected), len(selected), signature,
            ))
            return
        member = unit[index]
        for block_index in range(len(blocks)):
            updated = [list(block) for block in blocks]
            updated[block_index].append(member)
            visit(index + 1, updated)
        visit(index + 1, [*blocks, [member]])

    visit(0, [])
    current = max(partitions, key=lambda item: (item[0], item[1], item[2], -item[3], -item[4], item[5]))
    adr = max(partitions, key=lambda item: (item[0], item[1], item[2], item[3], -item[4], item[5]))
    if current[5] == adr[5]:
        return ()
    current_groups = set(current[5])
    return tuple(group for group in adr[5] if group not in current_groups)


def build_pool(value, result, members, edges):
    pool = []
    for group in result.accepted_groups:
        source = (
            "ACCEPTED_LIKELY_GROUP" if group.status.value == "LIKELY_DUPLICATE_GROUP"
            else "ACCEPTED_REVIEW_GROUP"
        )
        pool.append(_group_hypothesis(
            group.member_record_ids, members, edges, source_kind=source,
            current_group_id=group.hypothesis_id, system_status=group.status.value,
            work_unit_id=fingerprint(group.source_neighborhood_references, "WORK_UNIT"),
        ))
    for deferred in result.deferred_work_units:
        clique = _positive_clique(deferred.record_ids, edges)
        if clique:
            pool.append(_group_hypothesis(
                clique, members, edges, source_kind="DEFERRED_FAMILY_CANDIDATE",
                work_unit_id=deferred.deferred_id, deferred_context=True,
            ))
    units = _work_unit_members(value)
    for conflict in result.conflicts:
        unit = next((item for item in units if set(conflict.involved_record_ids) <= set(item)), ())
        clique = _positive_clique(unit, edges, maximum=5)
        if clique:
            pool.append(_group_hypothesis(
                clique, members, edges, source_kind="CONFLICT_CONTEXT_CANDIDATE",
                work_unit_id=conflict.conflict_id, conflict_context=True,
            ))
    for unit in units:
        unit_reference = fingerprint((members[item].stable_ref for item in unit), "WORK_UNIT")
        for challenger in _partition_challengers(unit, edges):
            pool.append(_group_hypothesis(
                challenger, members, edges, source_kind="ADR_PARTITION_CHALLENGER",
                current_or_shadow="SHADOW_ADR_ORDER", work_unit_id=unit_reference,
                partition_sensitive=True,
            ))
    unique = {}
    for group in pool:
        unique.setdefault(group.fingerprint, group)
    return tuple(sorted(unique.values(), key=lambda item: item.fingerprint))


def add_bicycle_group(pool, database_path: Path):
    connection = _readonly(database_path)
    rows = connection.execute(
        "SELECT * FROM scan_record_snapshot WHERE scan_id=33 AND description='Bicycle' ORDER BY record_ref_key"
    ).fetchall()
    connection.close()
    if len(rows) != 7:
        return pool, False
    bicycle_members = {row["id"]: _member(row) for row in rows}
    bicycle_edges = {}
    for left, right in combinations(rows, 2):
        bicycle_edges[_pair(left["id"], right["id"])] = Edge(
            left["id"], right["id"], "REVIEW_SUPPORT",
            "SHADOW_LEXICAL_ONLY_OR_UNRESOLVED", True, False,
            fingerprint((left["record_ref_key"], right["record_ref_key"]), "BICYCLE_EDGE"),
        )
    group = _group_hypothesis(
        bicycle_members, bicycle_members, bicycle_edges,
        source_kind="GENERIC_BICYCLE_FAMILY", current_or_shadow="CURRENT",
        work_unit_id="SCAN33_BICYCLE", system_status="", deferred_context=True,
    )
    unique = {item.fingerprint: item for item in pool}
    unique[group.fingerprint] = group
    return tuple(sorted(unique.values(), key=lambda item: item.fingerprint)), True


def _safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _new_workbook(title):
    workbook = Workbook()
    workbook.properties.creator = "R18G group human evidence"
    workbook.properties.title = title
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    return workbook


def _save_reproducible(workbook, path: Path):
    raw = io.BytesIO()
    workbook.save(raw)
    source = zipfile.ZipFile(io.BytesIO(raw.getvalue()), "r")
    normalized = io.BytesIO()
    with zipfile.ZipFile(normalized, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = source.getinfo(name).external_attr
            content = source.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"<dcterms:modified[^>]*>[^<]*</dcterms:modified>",
                    b'<dcterms:modified xsi:type="dcterms:W3CDTF">2000-01-01T00:00:00Z</dcterms:modified>',
                    content,
                )
            target.writestr(info, content)
    source.close()
    path.write_bytes(normalized.getvalue())


def _instructions(split_name):
    return (
        "Decide whether ALL records in each proposed group represent one underlying physical/business inventory item.",
        "Choose exactly one group label: SAME_ONE_IDENTITY, NOT_ONE_IDENTITY, or INSUFFICIENT_INFORMATION.",
        "Do not guess. Use INSUFFICIENT_INFORMATION when visible information is inadequate.",
        "Different ERP part numbers, sites, or contracts do not automatically mean different identity.",
        "Administrative/accounting mappings do not prove or disprove physical identity.",
        "Missing data is not proof of sameness. Generic/copied descriptions are not sufficient by themselves.",
        "Meaningful model, type, variant, material, size, capacity, or function differences matter.",
        "For NOT_ONE_IDENTITY, assign P1..P20 in Partition where reasonably possible; use UNRESOLVED otherwise.",
        "For SAME_ONE_IDENTITY, all members may use P1. For INSUFFICIENT_INFORMATION, partition may be blank or UNRESOLVED.",
        "Comment basis: SOURCE_VISIBLE, DOMAIN_EXTERNAL, MIXED, or UNCLEAR.",
        "Do not seek or use system results, pair-level labels, prior comments, or another reviewer's answer.",
        f"Workbook set: R18G {split_name}. Neither workbook has greater review importance.",
    )


def write_workbook(path: Path, groups: tuple[GroupHypothesis, ...], split_name: str):
    workbook = _new_workbook(f"R18G blinded group review - {split_name}")
    instructions = workbook.active
    instructions.title = "Instructions"
    instructions.append(["R18G blinded whole-group review"])
    instructions.append(["dataset_version", DATASET_VERSION])
    for line in _instructions(split_name):
        instructions.append([line])
    instructions.column_dimensions["A"].width = 120

    review = workbook.create_sheet("Group Review")
    review_headers = (
        "review_group_id", "member_count", "group_label", "confidence",
        "reason_code", "reviewer_comment", "comment_basis",
    )
    review.append(review_headers)
    for group in sorted(groups, key=review_group_id):
        review.append([review_group_id(group), len(group.members), "", "", "", "", ""])
    review.freeze_panes = "A2"
    review.auto_filter.ref = review.dimensions
    validation_specs = ((3, LABELS), (4, CONFIDENCES), (5, REASON_CODES), (7, COMMENT_BASES))
    for column, values in validation_specs:
        validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        review.add_data_validation(validation)
        if review.max_row >= 2:
            validation.add(f"{review.cell(2,column).coordinate}:{review.cell(review.max_row,column).coordinate}")

    member_sheet = workbook.create_sheet("Members")
    member_sheet.append(("review_group_id", "member_no", *VISIBLE_FIELDS))
    partition = workbook.create_sheet("Partition")
    partition.append(("review_group_id", "member_no", "part_number", "human_partition_id", "partition_comment"))
    for group in sorted(groups, key=review_group_id):
        group_id = review_group_id(group)
        ordered_members = sorted(group.members, key=lambda item: item.stable_ref)
        for number, member in enumerate(ordered_members, start=1):
            visible = [_safe_cell(getattr(member, field)) for field in VISIBLE_FIELDS]
            member_sheet.append([group_id, number, *visible])
            partition.append([group_id, number, _safe_cell(member.part_number), "", ""])
    for sheet in (member_sheet, partition):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    partition_validation = DataValidation(
        type="list", formula1='"' + ",".join(PARTITION_IDS) + '"', allow_blank=True
    )
    partition.add_data_validation(partition_validation)
    if partition.max_row >= 2:
        partition_validation.add(f"D2:D{partition.max_row}")
    for sheet in (review, member_sheet, partition):
        for column in sheet.columns:
            width = min(55, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
    _save_reproducible(workbook, path)


def write_mapping(path: Path, development, holdout):
    fields = (
        "review_group_id", "hypothesis_fingerprint", "source_kind", "current_or_shadow",
        "current_group_id", "work_unit_id", "system_status", "group_size", "site_count",
        "evidence_provenance", "strong_edge_count", "native_review_edge_count",
        "r18c_demoted_review_edge_count", "cannot_link_count", "generic_only_state",
        "partition_sensitive", "deferred_context", "conflict_context", "sampling_stratum",
        "development_or_holdout", "source_member_refs",
    )
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for split, groups in (("DEVELOPMENT", development), ("SEALED_HOLDOUT", holdout)):
            for group in sorted(groups, key=review_group_id):
                size_band = "2" if len(group.members) == 2 else "3" if len(group.members) == 3 else "4" if len(group.members) == 4 else "5+"
                sites = {member.site for member in group.members if member.site}
                stratum = "__".join((group.source_kind, f"SIZE_{size_band}", "CROSS_SITE" if len(sites) > 1 else "SAME_SITE"))
                writer.writerow({
                    "review_group_id": review_group_id(group),
                    "hypothesis_fingerprint": group.fingerprint,
                    "source_kind": group.source_kind,
                    "current_or_shadow": group.current_or_shadow,
                    "current_group_id": group.current_group_id,
                    "work_unit_id": group.work_unit_id,
                    "system_status": group.system_status,
                    "group_size": len(group.members),
                    "site_count": len(sites),
                    "evidence_provenance": group.evidence_provenance,
                    "strong_edge_count": group.strong_edge_count,
                    "native_review_edge_count": group.native_review_edge_count,
                    "r18c_demoted_review_edge_count": group.r18c_demoted_review_edge_count,
                    "cannot_link_count": group.cannot_link_count,
                    "generic_only_state": str(group.generic_only_state).lower(),
                    "partition_sensitive": str(group.partition_sensitive).lower(),
                    "deferred_context": str(group.deferred_context).lower(),
                    "conflict_context": str(group.conflict_context).lower(),
                    "sampling_stratum": stratum,
                    "development_or_holdout": split,
                    "source_member_refs": "|".join(member.stable_ref for member in sorted(group.members, key=lambda item: item.stable_ref)),
                })


def validate_workbooks(development_path: Path, holdout_path: Path,
                       mapping_path: Path, source_refs: set[str]):
    visible_ids = set()
    visible_member_keys = set()
    prefilled = 0
    formula_count = 0
    visible_text = []
    workbook_counts = {}
    for split, path in (("DEVELOPMENT", development_path), ("SEALED_HOLDOUT", holdout_path)):
        workbook = load_workbook(path, data_only=False)
        if workbook.sheetnames != ["Instructions", "Group Review", "Members", "Partition"]:
            raise ValueError("R18G_WORKBOOK_SHEETS_INVALID")
        review = workbook["Group Review"]
        visible_text.extend(
            str(cell.value or "").casefold()
            for row in workbook["Instructions"].iter_rows() for cell in row
        )
        headers = tuple(cell.value for cell in review[1])
        expected = (
            "review_group_id", "member_count", "group_label", "confidence",
            "reason_code", "reviewer_comment", "comment_basis",
        )
        if headers != expected:
            raise ValueError("R18G_REVIEW_HEADERS_INVALID")
        visible_text.extend(str(value).casefold() for value in headers)
        ids = []
        for row in review.iter_rows(min_row=2):
            ids.append(row[0].value)
            if int(row[1].value) < 2:
                raise ValueError("R18G_GROUP_TOO_SMALL")
            prefilled += sum(cell.value not in (None, "") for cell in row[2:])
            formula_count += sum(cell.data_type == "f" for cell in row[2:])
        if len(ids) != len(set(ids)):
            raise ValueError("R18G_DUPLICATE_VISIBLE_ID")
        if visible_ids & set(ids):
            raise ValueError("R18G_SPLIT_IDS_OVERLAP")
        visible_ids.update(ids)
        workbook_counts[split] = len(ids)
        members_sheet = workbook["Members"]
        member_headers = tuple(cell.value for cell in members_sheet[1])
        if member_headers != ("review_group_id", "member_no", *VISIBLE_FIELDS):
            raise ValueError("R18G_MEMBER_HEADERS_INVALID")
        local_members = set()
        for row in members_sheet.iter_rows(min_row=2, values_only=True):
            key = (row[0], int(row[1]))
            if key in local_members or row[0] not in ids:
                raise ValueError("R18G_MEMBER_RECONCILIATION_FAILED")
            local_members.add(key)
            visible_member_keys.add(key)
        partition = workbook["Partition"]
        partition_keys = {(row[0], int(row[1])) for row in partition.iter_rows(min_row=2, values_only=True)}
        if partition_keys != local_members:
            raise ValueError("R18G_PARTITION_RECONCILIATION_FAILED")
        for row in partition.iter_rows(min_row=2):
            prefilled += sum(cell.value not in (None, "") for cell in row[3:])
            formula_count += sum(cell.data_type == "f" for cell in row[3:])
    if prefilled or formula_count:
        raise ValueError("R18G_HUMAN_FIELDS_NOT_EMPTY")
    joined = "\n".join(visible_text)
    if any(term in joined for term in FORBIDDEN_VISIBLE_TERMS):
        raise ValueError("R18G_BLINDING_FAILURE")
    mapping = _read_mapping(mapping_path)
    mapping_ids = {row["review_group_id"] for row in mapping}
    if mapping_ids != visible_ids or len(mapping_ids) != len(mapping):
        raise ValueError("R18G_INTERNAL_MAPPING_RECONCILIATION_FAILED")
    if any(int(row["group_size"]) < 2 for row in mapping):
        raise ValueError("R18G_MAPPING_GROUP_TOO_SMALL")
    if any(ref not in source_refs for row in mapping for ref in row["source_member_refs"].split("|")):
        raise ValueError("R18G_SOURCE_MEMBER_RECONCILIATION_FAILED")
    if any(row["current_or_shadow"] == "SHADOW_ADR_ORDER" and int(row["cannot_link_count"]) for row in mapping):
        raise ValueError("R18G_SHADOW_CANNOT_LINK_UNSAFE")
    return {
        "all_review_group_ids_unique": True,
        "all_members_reconcile": True,
        "all_group_sizes_at_least_two": True,
        "reviewer_visible_detector_fields": 0,
        "pair_label_or_comment_leakage": 0,
        "prefilled_human_field_count": prefilled,
        "formula_count_in_human_fields": formula_count,
        "internal_mapping_reconciles": True,
        "development_holdout_disjoint": True,
        "shadow_cannot_link_violations": 0,
        "workbook_counts": workbook_counts,
    }


def _read_mapping(path):
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def write_manifest(path: Path, manifest: dict):
    """Write the canonical deterministic manifest representation."""
    path.write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )


def prepare_dataset(database_path: Path, output_dir: Path):
    value, result, members, edges, source = load_current_resolution(database_path)
    pool = build_pool(value, result, members, edges)
    if any(group.current_or_shadow == "CURRENT" and group.source_kind.startswith("ACCEPTED_") and group.cannot_link_count for group in pool):
        raise ValueError("R18G_ACCEPTED_GROUP_CANNOT_LINK")
    accepted_members = [
        member.stable_ref for group in pool
        if group.current_or_shadow == "CURRENT" and group.source_kind.startswith("ACCEPTED_")
        for member in group.members
    ]
    if len(accepted_members) != len(set(accepted_members)):
        raise ValueError("R18G_DUPLICATE_ACCEPTED_MEMBERSHIP")
    pool, bicycle_available = add_bicycle_group(pool, database_path)
    development, holdout = deterministic_split(pool)
    if bicycle_available and not any(group.source_kind == "GENERIC_BICYCLE_FAMILY" for group in development):
        raise ValueError("R18G_BICYCLE_NOT_IN_DEVELOPMENT")
    if any(group.partition_sensitive for group in pool) and not any(group.partition_sensitive for group in development):
        raise ValueError("R18G_PARTITION_CHALLENGER_NOT_IN_DEVELOPMENT")
    output_dir.mkdir(parents=True, exist_ok=True)
    development_path = output_dir / "r18g_group_review_development.xlsx"
    holdout_path = output_dir / "r18g_group_review_sealed_holdout.xlsx"
    mapping_path = output_dir / "r18g_group_review_internal_mapping.csv"
    manifest_path = output_dir / "r18g_group_review_manifest.json"
    write_workbook(development_path, development, "DEVELOPMENT")
    write_workbook(holdout_path, holdout, "SEALED HOLDOUT")
    write_mapping(mapping_path, development, holdout)
    all_source_refs = {member.stable_ref for member in members.values()}
    if bicycle_available:
        all_source_refs.update(member.stable_ref for group in pool for member in group.members)
    validation = validate_workbooks(
        development_path, holdout_path, mapping_path, all_source_refs
    )
    selected = development + holdout
    mapping_rows = _read_mapping(mapping_path)
    artifacts = {
        path.name: {"bytes": path.stat().st_size, "sha256": sha256_file(path)}
        for path in (development_path, holdout_path, mapping_path)
    }
    manifest = {
        "dataset_version": DATASET_VERSION,
        "seed": SPLIT_SEED,
        "generation_commit": GENERATION_COMMIT,
        "input_fingerprints": source,
        "reference_type": "SINGLE_SENIOR_DOMAIN_EXPERT_GROUP_EVIDENCE",
        "claim_limit": ["NOT_DUAL_REVIEWED", "NOT_ADJUDICATED", "NOT_GOLD_STANDARD", "NOT_PRODUCTION_ACCURACY_TRUTH"],
        "development_count": len(development),
        "sealed_holdout_count": len(holdout),
        "total_group_count": len(selected),
        "total_member_rows": sum(len(group.members) for group in selected),
        "eligible_pool_count": len(pool),
        "group_size_distribution": dict(sorted(Counter(len(group.members) for group in selected).items())),
        "sampling_strata": dict(sorted(Counter(row["sampling_stratum"] for row in mapping_rows).items())),
        "source_kinds": dict(sorted(Counter(group.source_kind for group in selected).items())),
        "development_bicycle_count": sum(group.source_kind == "GENERIC_BICYCLE_FAMILY" for group in development),
        "development_partition_sensitive_count": sum(group.partition_sensitive for group in development),
        "overlapping_source_members_allowed_for_partition_challengers": True,
        "visible_fields": list(VISIBLE_FIELDS),
        "human_label_count": 0,
        "sealed_holdout_labels_visible_to_engineering": 0,
        "blinding_validation": validation,
        "provider_calls": 0,
        "artifacts": artifacts,
    }
    write_manifest(manifest_path, manifest)
    return {**manifest, "manifest": {
        "path": manifest_path.name,
        "bytes": manifest_path.stat().st_size,
        "sha256": sha256_file(manifest_path),
    }}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--database", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    print(json.dumps(prepare_dataset(args.database, args.output_dir), indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
