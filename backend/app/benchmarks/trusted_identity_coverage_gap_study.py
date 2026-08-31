"""Offline-only R17 trusted-identity coverage counterfactuals.

The strategies in this module are diagnostic hypotheses.  They do not create
SignedIdentityEvidence, change GF4, or participate in a production scan path.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from itertools import combinations
from pathlib import Path
from time import perf_counter

from openpyxl import load_workbook

from app.benchmarks.signed_identity_evidence_shadow_audit import (
    _current_edge_class,
    _engine_record,
    _source_row_and_ref,
)
from app.engine.generic_description_guard import is_generic_description
from app.engine.identity_signature import IdentitySignature, IdentitySourceField
from app.engine.identity_signature_derivation import derive_identity_signature
from app.engine.normalizer import normalize_description, normalize_part_no_with_dictionary
from app.engine.signed_identity_evidence import (
    ShadowEvidenceBucket,
    classify_shadow_evidence,
    derive_signed_identity_evidence,
)


R17_COUNTERFACTUAL_VERSION = "trusted-identity-gap-study-v1"
STRATEGIES = ("A", "B", "C", "D", "E", "F")

_STRUCTURAL_WORDS = frozenset({
    "can", "description", "exercise", "generic", "item", "model", "part",
    "sales", "sample", "test", "type",
})
_SHORT_PREFIX = re.compile(r"^[a-z]{1,2}$")


@dataclass(frozen=True)
class CounterfactualResult:
    promoted: bool
    reasons: tuple[str, ...]
    normalized_span: str = ""


def _value(record: dict, *names: str) -> str:
    for name in names:
        if name in record:
            return str(record.get(name) or "")
    return ""


def _description(record: dict) -> str:
    return _value(record, "DESCRIPTION", "Part Description")


def _part_number(record: dict) -> str:
    return _value(record, "PART_NO", "Part No")


def _description_tokens(record: dict) -> tuple[str, ...]:
    return tuple(normalize_description(_description(record)).split())


def _semantic_units(record: dict) -> tuple[str, ...]:
    return tuple(
        token for token in _description_tokens(record)
        if token not in _STRUCTURAL_WORDS and (
            len(token) >= 3 or (any(c.isalpha() for c in token)
                                and any(c.isdigit() for c in token))
        )
    )


def _bounded_identity_span(record: dict) -> str:
    tokens = _description_tokens(record)
    units = _semantic_units(record)
    if not 2 <= len(tokens) <= 8 or len(units) < 2:
        return ""
    if is_generic_description(_description(record)):
        return ""
    return " ".join(units)


def _compound_head(record: dict) -> str:
    words = tuple(
        token for token in _semantic_units(record)
        if token.isalpha() and len(token) >= 3
    )
    return " ".join(words[-2:]) if len(words) >= 2 else ""


def _part_stem(record: dict) -> tuple[str, ...]:
    tokens = list(normalize_part_no_with_dictionary(_part_number(record)).split())
    removed = 0
    while tokens and removed < 2 and _SHORT_PREFIX.fullmatch(tokens[0]):
        tokens.pop(0)
        removed += 1
    return tuple(tokens)


def _observations(signature: IdentitySignature):
    return (
        signature.object_construct_observations
        + signature.assembly_component_role_observations
        + signature.model_type_observations
        + signature.variant_observations
        + signature.critical_attribute_observations
    )


def _shared(signature_a: IdentitySignature, signature_b: IdentitySignature):
    left = {
        (item.semantic_category.value, item.semantic_key, item.normalized_value): item
        for item in _observations(signature_a)
    }
    right = {
        (item.semantic_category.value, item.semantic_key, item.normalized_value): item
        for item in _observations(signature_b)
    }
    return [(key, left[key], right[key]) for key in sorted(set(left) & set(right))]


def _has_part_source(item) -> bool:
    return any(source.source_field == IdentitySourceField.PART_NUMBER for source in item.sources)


def _existing_fact_refinement(
    left: dict, right: dict, left_signature: IdentitySignature,
    right_signature: IdentitySignature,
) -> tuple[bool, tuple[str, ...]]:
    shared = _shared(left_signature, right_signature)
    models = [row for row in shared if row[0][1] == "structural_alphanumeric_model"]
    if not models:
        return False, ()
    roles = [row for row in shared if row[0][0] == "ASSEMBLY_COMPONENT_ROLE"]
    part_variants = [
        row for row in shared
        if row[0][0] == "VARIANT_IDENTITY"
        and _has_part_source(row[1]) and _has_part_source(row[2])
    ]
    full_part_equal = bool(
        normalize_part_no_with_dictionary(_part_number(left))
        and normalize_part_no_with_dictionary(_part_number(left))
        == normalize_part_no_with_dictionary(_part_number(right))
    )
    model_with_part = any(_has_part_source(a) and _has_part_source(b) for _key, a, b in models)
    if (model_with_part and (full_part_equal or part_variants)) or roles:
        return True, ("EXISTING_STRUCTURAL_MODEL_WITH_CORROBORATING_TYPED_FACT",)
    return False, ()


def _part_supports_phrase(record: dict, span: str) -> bool:
    if not span:
        return False
    words = tuple(span.split())
    normalized_part = normalize_part_no_with_dictionary(_part_number(record))
    ordered_part_tokens = tuple(normalized_part.split())
    part_tokens = set(ordered_part_tokens)
    if any(word in part_tokens for word in words if len(word) >= 4):
        return True
    compact_part = "".join(ordered_part_tokens)
    compact_span = "".join(words)
    if compact_span and compact_span in compact_part:
        return True
    acronym = "".join(word[0] for word in words if word and word[0].isalpha())
    return len(acronym) >= 2 and acronym in part_tokens


def evaluate_strategy(
    strategy: str,
    left: dict,
    right: dict,
    *,
    left_signature: IdentitySignature | None = None,
    right_signature: IdentitySignature | None = None,
) -> CounterfactualResult:
    """Evaluate one named offline strategy without constructing authority."""
    if strategy not in STRATEGIES:
        raise ValueError(f"unknown R17 strategy: {strategy}")
    left_signature = left_signature or derive_identity_signature(left)
    right_signature = right_signature or derive_identity_signature(right)
    left_span = _bounded_identity_span(left)
    right_span = _bounded_identity_span(right)
    span = left_span if left_span and left_span == right_span else ""
    exact_description = bool(
        normalize_description(_description(left))
        and normalize_description(_description(left))
        == normalize_description(_description(right))
        and 2 <= len(_description_tokens(left)) <= 8
        and not is_generic_description(_description(left))
    )
    stems_equal = bool(
        _part_stem(left) and _part_stem(left) == _part_stem(right)
        and any(len(token) >= 3 for token in _part_stem(left))
    )
    heads_equal = bool(
        _compound_head(left)
        and _compound_head(left) == _compound_head(right)
    )
    refined, refinement_reasons = _existing_fact_refinement(
        left, right, left_signature, right_signature
    )

    if strategy == "A":
        return CounterfactualResult(bool(span), ("BOUNDED_IDENTITY_PHRASE_EQUAL",) if span else (), span)
    if strategy == "B":
        return CounterfactualResult(exact_description, ("SOURCE_FIELD_EXACT_SPAN_EQUAL",) if exact_description else (), normalize_description(_description(left)) if exact_description else "")
    if strategy == "C":
        return CounterfactualResult(stems_equal, ("STRUCTURED_PART_NUMBER_STEM_EQUAL",) if stems_equal else ())
    if strategy == "D":
        return CounterfactualResult(heads_equal, ("COMPOUND_HEAD_EQUAL",) if heads_equal else (), _compound_head(left) if heads_equal else "")
    if strategy == "E":
        return CounterfactualResult(refined, refinement_reasons)

    phrase_corroborated = bool(
        span and _part_supports_phrase(left, span) and _part_supports_phrase(right, span)
    )
    promoted = refined or stems_equal or phrase_corroborated
    reasons = tuple(sorted(set(
        refinement_reasons
        + (("STRUCTURED_PART_NUMBER_STEM_EQUAL",) if stems_equal else ())
        + (("SOURCE_CORROBORATED_BOUNDED_IDENTITY_PHRASE",) if phrase_corroborated else ())
    )))
    return CounterfactualResult(promoted, reasons, span if phrase_corroborated else "")


def _explicit_numeric_variant_risk(left: dict, right: dict) -> bool:
    left_numbers = tuple(re.findall(r"\d+", normalize_part_no_with_dictionary(_part_number(left))))
    right_numbers = tuple(re.findall(r"\d+", normalize_part_no_with_dictionary(_part_number(right))))
    return bool(left_numbers and right_numbers and left_numbers != right_numbers)


def _promotion_label(
    result: CounterfactualResult, current_edge_class: str, left: dict, right: dict
) -> str:
    if _explicit_numeric_variant_risk(left, right):
        return "UNSAFE_PROMOTION"
    if "EXISTING_STRUCTURAL_MODEL_WITH_CORROBORATING_TYPED_FACT" in result.reasons:
        return "SAFE_TRUSTED_IDENTITY"
    if current_edge_class == "STRONG_SUPPORT":
        return "PLAUSIBLE_BUT_NOT_TRUSTED"
    return "INSUFFICIENT_DATA"


def audit_counterfactuals(csv_path: Path, workbook_path: Path) -> dict:
    """Audit A-F over exactly the accepted R12 complete-pair edges."""
    started = perf_counter()
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        source_rows = list(csv.DictReader(handle))
    book = load_workbook(workbook_path, data_only=False, read_only=True)
    sheet = book["Group Data"]
    headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    groups = defaultdict(list)
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values, strict=True))
        source_index, stable_ref = _source_row_and_ref(row["Source Row / Stable Record Reference"])
        groups[row["Canonical Group ID"]].append(
            (source_index, stable_ref, _engine_record(source_rows[source_index]), row)
        )

    signatures = {}
    results = {strategy: [] for strategy in STRATEGIES}
    baseline = Counter()
    for members in groups.values():
        for left_item, right_item in combinations(members, 2):
            left_index, left_ref, left, left_context = left_item
            right_index, right_ref, right, right_context = right_item
            for ref, record in ((left_ref, left), (right_ref, right)):
                signatures.setdefault(ref, derive_identity_signature(record, record_reference=ref))
            evidence = derive_signed_identity_evidence(signatures[left_ref], signatures[right_ref])
            bucket = classify_shadow_evidence(evidence).value
            baseline[bucket] += 1
            current = _current_edge_class(left, right)
            for strategy in STRATEGIES:
                outcome = evaluate_strategy(
                    strategy, left, right,
                    left_signature=signatures[left_ref], right_signature=signatures[right_ref],
                )
                if outcome.promoted and bucket != ShadowEvidenceBucket.SHADOW_TRUSTED_IDENTITY_PRESENT.value:
                    results[strategy].append({
                        "left_row": left_index,
                        "right_row": right_index,
                        "left_ref": left_ref,
                        "right_ref": right_ref,
                        "left_part_no": left["PART_NO"],
                        "right_part_no": right["PART_NO"],
                        "left_description": left["DESCRIPTION"],
                        "right_description": right["DESCRIPTION"],
                        "current_edge_class": current,
                        "group_status": left_context["Status"],
                        "reasons": outcome.reasons,
                        "inspection_label": _promotion_label(
                            outcome, current, left, right
                        ),
                    })

    strategy_summary = {}
    baseline_trusted = baseline[ShadowEvidenceBucket.SHADOW_TRUSTED_IDENTITY_PRESENT.value]
    baseline_lexical = baseline[ShadowEvidenceBucket.SHADOW_LEXICAL_ONLY_OR_UNRESOLVED.value]
    for strategy, promoted in results.items():
        labels = Counter(item["inspection_label"] for item in promoted)
        strategy_summary[strategy] = {
            "newly_promoted": len(promoted),
            "counterfactual_trusted": baseline_trusted + len(promoted),
            "counterfactual_lexical_only": baseline_lexical - len(promoted),
            "promotion_labels": dict(sorted(labels.items())),
            "promotions": promoted,
        }
    return {
        "version": R17_COUNTERFACTUAL_VERSION,
        "pair_source": "R12 read-only accepted-group complete-pair edges",
        "group_count": len(groups),
        "real_pairs_evaluated": sum(baseline.values()),
        "baseline_buckets": dict(sorted(baseline.items())),
        "strategies": strategy_summary,
        "wall_time_seconds": perf_counter() - started,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("csv_path", type=Path)
    parser.add_argument("workbook_path", type=Path)
    parser.add_argument("--summary-only", action="store_true")
    args = parser.parse_args()
    result = audit_counterfactuals(args.csv_path, args.workbook_path)
    if args.summary_only:
        for value in result["strategies"].values():
            value.pop("promotions", None)
    print(json.dumps(result, ensure_ascii=True, sort_keys=True, indent=2))


if __name__ == "__main__":
    main()
