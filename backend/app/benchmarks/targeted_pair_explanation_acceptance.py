"""Provider-free acceptance for targeted pair-explanation evidence persistence."""

from __future__ import annotations

import argparse
from dataclasses import replace
import json
import sqlite3
import tempfile
import time
from collections import Counter
from io import BytesIO
from pathlib import Path

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from openpyxl import load_workbook

from app.benchmarks.acceptance_provenance import request_provenance
from app.benchmarks.match_strength_v2_shadow import (
    approved_csv,
    provider_disabled_configuration,
)
from app.benchmarks.scan_determinism_audit import ScanAudit
from app.db.database import Base
from app.db.models import (
    G2V2ProjectionRun,
    HumanIdentityConstraint,
    IdentityEvidenceEdgeSnapshot,
    IdentityResolutionRun,
)
from app.g2_v2.contracts import G2V2EvidenceOrigin
from app.match_strength.service import MatchStrengthProjectionService
from app.identity_read.deterministic_explanations import (
    project_group_explanation,
)
from app.resolution.pair_explanation import (
    PairExplanationAvailability,
    project_proposal_pair_explanation,
    project_targeted_pair_explanation,
)
from app.services.identity_read_service import IdentityReadService
from app.services.identity_read_xlsx_export_service import (
    SHEET_ORDER,
    authority_selected_system_groups_to_xlsx,
)
from app.services.identity_resolution_service import load_persisted_resolution_result
from app.services.scan_runner import ScanRunner


MODES = {
    "site_selected": {
        "selected_fields": ["CONTRACT", "UNIT_MEAS"],
        "request_fingerprint": (
            "fab4c9e6de6bad086bd14b4582d16b9ad64ecb2330b7663f2cf7cab5fff3faff"
        ),
        "artifact": "site_selected.acceptance_provenance.json",
    },
    "site_unselected": {
        "selected_fields": ["UNIT_MEAS"],
        "request_fingerprint": (
            "6742e55aeecdd734e7e37875abbcf79f8c8ec8a1531d4c06bca6f3ebed49a85a"
        ),
        "artifact": "site_unselected.acceptance_provenance.json",
    },
}


def _provider_calls(session, scan_id: int) -> int:
    resolution = session.query(IdentityResolutionRun).filter_by(scan_id=scan_id).one()
    return int(resolution.provider_request_count or 0)


def _stages(database: Path, scan_id: int):
    connection = sqlite3.connect(database)
    connection.row_factory = sqlite3.Row
    try:
        return {
            item.stage.split("_", 1)[0]: {
                "name": item.stage,
                "count": item.count,
                "sha256": item.fingerprint,
            }
            for item in ScanAudit(connection, scan_id).stages()
        }
    finally:
        connection.close()


def _run_mode(repository_root: Path, records, mode: str, settings: dict) -> dict:
    with tempfile.TemporaryDirectory(
        prefix=f"targeted-pair-explanation-{mode}-"
    ) as directory:
        database = Path(directory) / "acceptance.sqlite3"
        engine = create_engine(
            f"sqlite:///{database.as_posix()}",
            connect_args={"check_same_thread": False},
        )
        Base.metadata.create_all(engine)
        session = sessionmaker(bind=engine)()
        try:
            scan, _ = ScanRunner(
                session, provider_disabled_configuration()
            ).run(
                records,
                f"Targeted explanation acceptance: {mode}",
                settings["selected_fields"],
                75.0,
                source_type="CANONICAL_DB_SNAPSHOT",
                sensitive_mode=True,
                scan_mode="SAME_SITE_DUPLICATE",
            )
            snapshot = IdentityReadService(session).load_identity_read_snapshot(scan.id)
            resolution = session.query(IdentityResolutionRun).filter_by(
                scan_id=scan.id
            ).one()
            projection = session.query(G2V2ProjectionRun).filter_by(
                scan_id=scan.id
            ).one()
            persisted = load_persisted_resolution_result(session, resolution.id)
            targeted = {
                item.evidence_fingerprint: item
                for item in persisted.targeted_evidence_results
            }
            proposal_references = {
                item.evidence_fingerprint
                for group in snapshot.groups
                for item in group.internal_evidence
                if item.evidence_origin == G2V2EvidenceOrigin.PROPOSAL_EVIDENCE
            }
            proposal_rows = session.query(IdentityEvidenceEdgeSnapshot).filter(
                IdentityEvidenceEdgeSnapshot.evidence_fingerprint.in_(
                    proposal_references
                )
            ).all()
            proposals = {item.evidence_fingerprint: item for item in proposal_rows}

            coverage = Counter()
            by_origin = Counter()
            group_complete = {}
            relationship_count = 0
            for group in snapshot.groups:
                complete = True
                for item in group.internal_evidence:
                    relationship_count += 1
                    if item.evidence_origin == G2V2EvidenceOrigin.PROPOSAL_EVIDENCE:
                        source = proposals.get(item.evidence_fingerprint)
                        if source is None:
                            availability = "UNKNOWN"
                        else:
                            explanation = project_proposal_pair_explanation(
                                source,
                                record_reference_1=item.stable_record_reference_1,
                                record_reference_2=item.stable_record_reference_2,
                            )
                            availability = explanation.availability.value
                        by_origin[f"PROPOSAL_{availability}"] += 1
                    else:
                        source = targeted.get(item.evidence_fingerprint)
                        if source is None:
                            availability = "UNKNOWN"
                        else:
                            explanation = project_targeted_pair_explanation(source)
                            availability = explanation.availability.value
                        by_origin[f"TARGETED_{availability}"] += 1
                    coverage[availability] += 1
                    complete = complete and (
                        availability == PairExplanationAvailability.COMPLETE.value
                    )
                group_complete[group.versioned_group_key] = complete

            strengths = MatchStrengthProjectionService(session).project_groups(
                snapshot.groups
            )
            projection_started = time.perf_counter()
            read_models = []
            for group in snapshot.groups:
                sources = {}
                for item in group.internal_evidence:
                    if item.evidence_origin == G2V2EvidenceOrigin.PROPOSAL_EVIDENCE:
                        source = proposals[item.evidence_fingerprint]
                        sources[item.evidence_fingerprint] = project_proposal_pair_explanation(
                            source,
                            record_reference_1=item.stable_record_reference_1,
                            record_reference_2=item.stable_record_reference_2,
                        )
                    else:
                        sources[item.evidence_fingerprint] = project_targeted_pair_explanation(
                            targeted[item.evidence_fingerprint]
                        )
                read_models.append(project_group_explanation(
                    group, strengths[group.versioned_group_key], sources,
                    include_details=True,
                ))
            projection_elapsed = time.perf_counter() - projection_started
            all_pair_details = [
                pair
                for model in read_models
                for pair in model.pair_explanations
            ]
            two_member_strong = [
                model for group, model in zip(snapshot.groups, read_models)
                if group.member_count == 2
                and group.status.value == "LIKELY_DUPLICATE_GROUP"
            ][:5]
            two_member_review = [
                model for group, model in zip(snapshot.groups, read_models)
                if group.member_count == 2
                and group.status.value == "POSSIBLE_DUPLICATE_GROUP_REVIEW"
            ][:5]
            scored_pairs = sorted(
                (pair for pair in all_pair_details if pair.deterministic_score is not None),
                key=lambda pair: (pair.deterministic_score, pair.relationship_id),
            )
            score_strata = {
                "lowest_10": scored_pairs[:10],
                "highest_10": scored_pairs[-10:],
                "nearest_90_10": sorted(
                    scored_pairs,
                    key=lambda pair: (abs(pair.deterministic_score - 90), pair.relationship_id),
                )[:10],
                "nearest_60_10": sorted(
                    scored_pairs,
                    key=lambda pair: (abs(pair.deterministic_score - 60), pair.relationship_id),
                )[:10],
            }
            rendered_items = [
                item
                for pair in all_pair_details
                for item in (
                    pair.supporting_items + pair.weakening_items
                    + pair.contradiction_items + pair.safety_items
                )
            ]
            rendering_inspection = {
                "strong_two_member_groups": len(two_member_strong),
                "review_two_member_groups": len(two_member_review),
                "multi_member_groups": sum(group.member_count >= 3 for group in snapshot.groups),
                "pairs_with_score": sum(pair.deterministic_score is not None for pair in all_pair_details),
                "pairs_with_signed_relationship": sum(bool(pair.signed_relationship) for pair in all_pair_details),
                "pairs_with_supporting_evidence": sum(bool(pair.supporting_items) for pair in all_pair_details),
                "reason_codes_visible": sum(
                    len(pair.decision_reason_codes) for pair in all_pair_details
                ),
                "evidence_items": len(rendered_items),
                "evidence_items_with_source_field": sum(
                    bool(item.source_field) for item in rendered_items
                ),
                "unsafe_causal_summary_count": sum(
                    any(term in model.group_summary.lower() for term in (
                        "caused the group", "decisive pair", "primary cause"
                    ))
                    for model in read_models
                ),
                "score_strata_counts": {
                    key: len(value) for key, value in score_strata.items()
                },
                "two_member_examples": [
                    {
                        "summary": model.group_summary,
                        "score": model.relationships[0].deterministic_score,
                        "relationship": model.relationships[0].signed_relationship,
                        "reason_codes": list(model.pair_explanations[0].decision_reason_codes),
                    }
                    for model in two_member_strong[:2] + two_member_review[:2]
                ],
            }
            from app.api.routes_identity_groups import _read_group
            baseline_payloads = [
                _read_group(
                    group, match_strength=strengths[group.versioned_group_key]
                )
                for group in snapshot.groups
            ]
            explanation_payloads = [
                _read_group(
                    group, match_strength=strengths[group.versioned_group_key],
                    deterministic_explanation=replace(
                        read_model, pair_explanations=()
                    ),
                )
                for group, read_model in zip(snapshot.groups, read_models)
            ]
            payload_before = len(json.dumps(baseline_payloads, sort_keys=True))
            payload_after = len(json.dumps(explanation_payloads, sort_keys=True))
            largest_group_payload = max(
                (
                    len(json.dumps(_read_group(
                        group,
                        match_strength=strengths[group.versioned_group_key],
                        deterministic_explanation=read_model,
                        detail=True,
                    ), sort_keys=True))
                    for group, read_model in zip(snapshot.groups, read_models)
                ),
                default=0,
            )
            bands = Counter(
                result.match_band.value if result.match_band else "UNSCORED"
                for result in strengths.values()
            )
            crossovers = [
                group
                for group in snapshot.groups
                if group.status.value == "POSSIBLE_DUPLICATE_GROUP_REVIEW"
                and strengths[group.versioned_group_key].match_band is not None
                and strengths[group.versioned_group_key].match_band.value == "HIGH_MATCH"
            ]
            exact_crossover = sum(
                bool([
                    item
                    for item in group.internal_evidence
                    if item.edge_class.value == "REVIEW_SUPPORT"
                    and item.reason_codes
                ])
                for group in crossovers
            )
            current_stages = _stages(database, scan.id)
            baseline = json.loads((
                repository_root
                / "artifacts"
                / "request_scoped_site_constraint"
                / settings["artifact"]
            ).read_text(encoding="utf-8"))
            stage_matches = {
                key: current_stages[key] == baseline["semantic_fingerprints"][key]
                for key in current_stages
            }
            cross_site_groups = sum(
                len({member.contract for member in group.members}) > 1
                for group in snapshot.groups
            )
            request = request_provenance(
                scan_mode="SAME_SITE_DUPLICATE",
                selected_fields=settings["selected_fields"],
                threshold=75.0,
                candidate_mode="group_first_primary",
                sensitive_mode=True,
                source_type="CANONICAL_DB_SNAPSHOT",
                feature_flags={
                    "hybrid_retrieval_enabled": True,
                    "local_embedding_enabled": True,
                },
                semantic_options={
                    "identity_discovery_mode": "DISCOVERY",
                    "visible_projection_contract": "G2_V2",
                },
            )
            xlsx_result = None
            if mode == "site_selected":
                xlsx_started = time.perf_counter()
                workbook_bytes = authority_selected_system_groups_to_xlsx(
                    session, scan.id
                )
                xlsx_seconds = time.perf_counter() - xlsx_started
                workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)
                formula_count = sum(
                    cell.data_type == "f"
                    for sheet in workbook.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                )
                review_sheet = workbook["Review Groups"]
                xlsx_result = {
                    "seconds": round(xlsx_seconds, 6),
                    "bytes": len(workbook_bytes),
                    "sheets": list(workbook.sheetnames),
                    "sheet_order_matches": tuple(workbook.sheetnames) == SHEET_ORDER,
                    "formula_count": formula_count,
                    "review_group_rows": review_sheet.max_row - 1,
                    "maximum_row_height": max(
                        (review_sheet.row_dimensions[index].height or 15)
                        for index in range(2, review_sheet.max_row + 1)
                    ),
                }
            result = {
                "request_fingerprint": request["sha256"],
                "request_fingerprint_matches": (
                    request["sha256"] == settings["request_fingerprint"]
                ),
                "detector_counts": {
                    "records": resolution.source_record_count,
                    "groups": resolution.accepted_group_count,
                    "stronger": resolution.likely_group_count,
                    "review": resolution.review_group_count,
                    "conflicts": resolution.conflict_count,
                    "deferred": resolution.deferred_work_unit_count,
                    "grouped_records": (
                        projection.canonical_record_count
                        - projection.unassigned_record_count
                    ),
                    "unassigned": resolution.unassigned_record_count,
                    "cross_site_groups": cross_site_groups,
                    "provider_calls": _provider_calls(session, scan.id),
                    "human_constraints": session.query(
                        HumanIdentityConstraint
                    ).count(),
                },
                "match_strength": {
                    "scored": len(strengths) - bands["UNSCORED"],
                    "unscored": bands["UNSCORED"],
                    "high": bands["HIGH_MATCH"],
                    "moderate": bands["MODERATE_MATCH"],
                    "borderline": bands["BORDERLINE_MATCH"],
                },
                "explanation_coverage": {
                    "group_summaries": len(read_models),
                    "relationship_map_entries": sum(
                        len(item.relationships) for item in read_models
                    ),
                    "rendered_pair_explanations": sum(
                        len(item.pair_explanations) for item in read_models
                    ),
                    "projection_seconds": round(projection_elapsed, 6),
                    "list_payload_bytes_before": payload_before,
                    "list_payload_bytes_after": payload_after,
                    "largest_group_payload_bytes": largest_group_payload,
                    "relationships": relationship_count,
                    "full": coverage["COMPLETE"],
                    "partial": coverage["PARTIAL_LEGACY"],
                    "generic_only": 0,
                    "requires_rerun": 0,
                    "unknown": coverage["UNKNOWN"],
                    "by_origin": dict(sorted(by_origin.items())),
                    "two_member_groups_full": sum(
                        group.member_count == 2
                        and group_complete[group.versioned_group_key]
                        for group in snapshot.groups
                    ),
                    "two_member_groups_total": sum(
                        group.member_count == 2 for group in snapshot.groups
                    ),
                    "multi_member_groups_full": sum(
                        group.member_count >= 3
                        and group_complete[group.versioned_group_key]
                        for group in snapshot.groups
                    ),
                    "multi_member_groups_total": sum(
                        group.member_count >= 3 for group in snapshot.groups
                    ),
                },
                "rendering_inspection": rendering_inspection,
                "crossover_reason_coverage": {
                    "groups": len(crossovers),
                    "exact": exact_crossover,
                    "partial": len(crossovers) - exact_crossover,
                    "generic_only": 0,
                },
                "semantic_stage_matches": stage_matches,
                "semantic_fingerprints": current_stages,
                "xlsx": xlsx_result,
            }
        finally:
            session.close()
            engine.dispose()
    return result


def run_acceptance(repository_root: Path) -> dict:
    records = approved_csv(repository_root)
    result = {"contract": "targeted-pair-explanation-acceptance-v1"}
    for mode, settings in MODES.items():
        result[mode] = _run_mode(repository_root, records, mode, settings)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repository-root", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = run_acceptance(args.repository_root.resolve())
    rendered = json.dumps(result, indent=2, sort_keys=True) + "\n"
    if args.output:
        args.output.write_text(rendered, encoding="utf-8")
    print(rendered, end="")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
