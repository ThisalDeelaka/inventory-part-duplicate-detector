"""Offline validation and adjudication packaging for R18 human workbooks.

This module validates human-entered files against the original blinded
workbooks. It does not create labels, inspect detector-private mappings, or
participate in production scan execution.
"""

from __future__ import annotations

import argparse
import json
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app.benchmarks.human_identity_review_dataset import (
    CONFIDENCES,
    LABELS,
    REASON_CODES,
    _save_reproducible_xlsx,
)


INPUT_HEADERS = (
    "identity_label",
    "confidence",
    "reason_code",
    "reviewer_comment",
)
ADJUDICATION_INPUT_HEADERS = (
    "adjudicated_label",
    "adjudicator_confidence",
    "adjudication_reason",
)
FORBIDDEN_BLINDED_HEADER_TERMS = (
    "score",
    "similarity",
    "detector",
    "status",
    "gf4",
    "gf5",
    "group",
    "llm",
    "known_control",
    "expected",
    "stratum",
)


class HumanIdentityReviewExecutionError(ValueError):
    """A submitted human-review artifact failed closed validation."""


@dataclass(frozen=True)
class ReviewerSubmissionValidation:
    reviewer_code: str
    pair_count: int
    completed_label_count: int
    declared_complete: bool
    source_fields_unchanged: bool = True
    pair_ids_reconciled: bool = True


def _fail(code: str) -> None:
    raise HumanIdentityReviewExecutionError(code)


def _sheet_values(sheet) -> tuple[tuple[object, ...], ...]:
    return tuple(tuple(cell.value for cell in row) for row in sheet.iter_rows())


def _load_review(path: Path):
    workbook = load_workbook(path, data_only=False, read_only=False)
    if tuple(workbook.sheetnames) != ("Instructions", "Review"):
        _fail("R18A_REVIEWER_SHEETS_INVALID")
    sheet = workbook["Review"]
    if any(cell.data_type == "f" for row in sheet.iter_rows() for cell in row):
        _fail("R18A_FORMULA_CELL_FORBIDDEN")
    headers = tuple(cell.value for cell in sheet[1])
    if not headers or len(headers) != len(set(headers)):
        _fail("R18A_REVIEWER_HEADERS_INVALID")
    if any(
        term in str(header).casefold()
        for header in headers
        for term in FORBIDDEN_BLINDED_HEADER_TERMS
    ):
        _fail("R18A_BLINDING_INVALID")
    if "review_pair_id" not in headers or any(name not in headers for name in INPUT_HEADERS):
        _fail("R18A_REVIEWER_HEADERS_INVALID")
    return workbook, sheet, headers


def _reviewer_code(workbook: Workbook) -> str:
    values = _sheet_values(workbook["Instructions"])
    matches = [row[1] for row in values if len(row) >= 2 and row[0] == "reviewer_code"]
    if len(matches) != 1 or not isinstance(matches[0], str):
        _fail("R18A_REVIEWER_CODE_INVALID")
    return matches[0]


def _rows_by_id(sheet, headers) -> dict[str, tuple[object, ...]]:
    pair_index = headers.index("review_pair_id")
    rows: dict[str, tuple[object, ...]] = {}
    for cells in sheet.iter_rows(min_row=2):
        row = tuple(cell.value for cell in cells)
        pair_id = row[pair_index]
        if not isinstance(pair_id, str) or not pair_id:
            _fail("R18A_PAIR_ID_MISSING")
        if pair_id in rows:
            _fail("R18A_DUPLICATE_PAIR_ID")
        rows[pair_id] = row
    return rows


def validate_reviewer_submission(
    reference_path: Path,
    submission_path: Path,
    *,
    expected_reviewer_code: str,
    declared_complete: bool,
) -> ReviewerSubmissionValidation:
    """Validate a reviewer file without altering or interpreting its labels."""
    reference_book, reference_sheet, reference_headers = _load_review(reference_path)
    submitted_book, submitted_sheet, submitted_headers = _load_review(submission_path)
    if submitted_headers != reference_headers:
        _fail("R18A_REVIEWER_HEADERS_CHANGED")
    if _reviewer_code(reference_book) != expected_reviewer_code:
        _fail("R18A_REFERENCE_REVIEWER_CODE_MISMATCH")
    if _reviewer_code(submitted_book) != expected_reviewer_code:
        _fail("R18A_SUBMISSION_REVIEWER_CODE_MISMATCH")
    if _sheet_values(reference_book["Instructions"]) != _sheet_values(
        submitted_book["Instructions"]
    ):
        _fail("R18A_INSTRUCTIONS_CHANGED")

    reference_rows = _rows_by_id(reference_sheet, reference_headers)
    submitted_rows = _rows_by_id(submitted_sheet, submitted_headers)
    if set(submitted_rows) - set(reference_rows):
        _fail("R18A_FOREIGN_PAIR_ID")
    if set(reference_rows) - set(submitted_rows):
        _fail("R18A_EXPECTED_PAIR_ID_MISSING")

    input_indexes = {name: submitted_headers.index(name) for name in INPUT_HEADERS}
    source_indexes = tuple(
        index
        for index, name in enumerate(submitted_headers)
        if name not in INPUT_HEADERS
    )
    completed = 0
    for pair_id, submitted in submitted_rows.items():
        reference = reference_rows[pair_id]
        if any(submitted[index] != reference[index] for index in source_indexes):
            _fail("R18A_SOURCE_FIELDS_CHANGED")
        label = submitted[input_indexes["identity_label"]]
        confidence = submitted[input_indexes["confidence"]]
        reason = submitted[input_indexes["reason_code"]]
        comment = submitted[input_indexes["reviewer_comment"]]
        if label not in (None, "") and label not in LABELS:
            _fail("R18A_LABEL_INVALID")
        if confidence not in (None, "") and confidence not in CONFIDENCES:
            _fail("R18A_CONFIDENCE_INVALID")
        if reason not in (None, "") and reason not in REASON_CODES:
            _fail("R18A_REASON_CODE_INVALID")
        if comment not in (None, "") and not isinstance(comment, str):
            _fail("R18A_COMMENT_INVALID")
        fields_present = tuple(value not in (None, "") for value in (label, confidence, reason))
        if any(fields_present) and not all(fields_present):
            _fail("R18A_PARTIAL_LABEL_INVALID")
        if all(fields_present):
            completed += 1
        elif declared_complete:
            _fail("R18A_COMPLETE_WORKBOOK_HAS_MISSING_LABEL")

    return ReviewerSubmissionValidation(
        reviewer_code=expected_reviewer_code,
        pair_count=len(reference_rows),
        completed_label_count=completed,
        declared_complete=declared_complete,
    )


def create_disagreement_adjudication_workbook(
    reviewer_a_reference: Path,
    reviewer_a_submission: Path,
    reviewer_b_reference: Path,
    reviewer_b_submission: Path,
    output_path: Path,
) -> int:
    """Create a blinded, disagreement-only adjudication workbook."""
    validate_reviewer_submission(
        reviewer_a_reference,
        reviewer_a_submission,
        expected_reviewer_code="REVIEWER_A",
        declared_complete=True,
    )
    validate_reviewer_submission(
        reviewer_b_reference,
        reviewer_b_submission,
        expected_reviewer_code="REVIEWER_B",
        declared_complete=True,
    )
    _a_book, a_sheet, headers = _load_review(reviewer_a_submission)
    _b_book, b_sheet, b_headers = _load_review(reviewer_b_submission)
    if b_headers != headers:
        _fail("R18A_REVIEWER_HEADERS_CHANGED")
    a_rows = _rows_by_id(a_sheet, headers)
    b_rows = _rows_by_id(b_sheet, headers)
    if set(a_rows) != set(b_rows):
        _fail("R18A_REVIEWER_PAIR_SET_MISMATCH")

    input_indexes = {name: headers.index(name) for name in INPUT_HEADERS}
    source_headers = tuple(name for name in headers if name not in INPUT_HEADERS)
    source_indexes = tuple(headers.index(name) for name in source_headers)
    output_headers = source_headers + (
        "reviewer_a_label",
        "reviewer_a_confidence",
        "reviewer_a_reason",
        "reviewer_a_note",
        "reviewer_b_label",
        "reviewer_b_confidence",
        "reviewer_b_reason",
        "reviewer_b_note",
    ) + ADJUDICATION_INPUT_HEADERS

    disagreements = []
    for pair_id in sorted(a_rows):
        a_row = a_rows[pair_id]
        b_row = b_rows[pair_id]
        if a_row[input_indexes["identity_label"]] == b_row[input_indexes["identity_label"]]:
            continue
        disagreements.append(
            tuple(a_row[index] for index in source_indexes)
            + tuple(a_row[input_indexes[name]] for name in INPUT_HEADERS)
            + tuple(b_row[input_indexes[name]] for name in INPUT_HEADERS)
            + ("", "", "")
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Adjudication"
    sheet.append(output_headers)
    for row in disagreements:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for name, values in (
        ("adjudicated_label", LABELS),
        ("adjudicator_confidence", CONFIDENCES),
    ):
        column = output_headers.index(name) + 1
        validation = DataValidation(
            type="list", formula1='"' + ",".join(values) + '"', allow_blank=True
        )
        sheet.add_data_validation(validation)
        if disagreements:
            validation.add(
                f"{sheet.cell(2, column).coordinate}:{sheet.cell(sheet.max_row, column).coordinate}"
            )
    _save_reproducible_xlsx(workbook, output_path)
    return len(disagreements)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--reference", required=True, type=Path)
    parser.add_argument("--submission", required=True, type=Path)
    parser.add_argument("--reviewer-code", required=True, choices=("REVIEWER_A", "REVIEWER_B"))
    parser.add_argument("--declared-complete", action="store_true")
    arguments = parser.parse_args()
    result = validate_reviewer_submission(
        arguments.reference,
        arguments.submission,
        expected_reviewer_code=arguments.reviewer_code,
        declared_complete=arguments.declared_complete,
    )
    print(json.dumps(asdict(result), indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
