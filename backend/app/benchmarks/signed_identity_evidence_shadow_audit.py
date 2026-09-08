"""Offline-only R16 audit over a read-only accepted-group workbook."""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from itertools import combinations
from pathlib import Path
from time import perf_counter

from openpyxl import load_workbook

from app.engine.identity_discriminator import evaluate_identity_discriminators
from app.engine.identity_edge import classify_identity_edge
from app.engine.identity_signature_derivation import derive_identity_signature
from app.engine.scoring import score_candidate
from app.engine.signed_identity_evidence import (
    classify_shadow_evidence,
    derive_signed_identity_evidence,
)


def _source_row_and_ref(value):
    left, right = str(value).split(" / ", 1)
    return int(re.search(r"\d+", left).group(0)), right


def _engine_record(row):
    return {
        "PART_NO": row["Part No"],
        "DESCRIPTION": row["Part Description"],
        "MASTER_DESCRIPTION": row.get("Master Part Description", ""),
        "TYPE_DESIGNATION": row.get("Type Designation", ""),
        "DIMENSION_QUALITY": row.get("Dimension/ Quality", ""),
        "CONTRACT": row["Site"],
        "UNIT_MEAS": row["Inventory UoM"],
        "TYPE_CODE": row["Part Type"],
        "PRIME_COMMODITY": row["Commodity Group 1"],
        "SECOND_COMMODITY": row["Commodity Group 2"],
        "ACCOUNTING_GROUP": row["Accounting Group"],
        "PART_PRODUCT_CODE": row["Product Code"],
        "PART_PRODUCT_FAMILY": row["Product Family"],
        "PRODUCT_CATEGORY_ID": row["Product Category"],
        "HSN_SAC_CODE": row["HSN/SAC Code"],
        "HAZARD_CODE": row["Safety Code"],
    }


def _current_edge_class(left, right):
    score = score_candidate(
        left, right, ["CONTRACT", "UNIT_MEAS"], "SAME_SITE_DUPLICATE",
        allow_uom_mapping_review=True,
    )
    discriminator = evaluate_identity_discriminators(
        left["PART_NO"], left["DESCRIPTION"], left["UNIT_MEAS"],
        right["PART_NO"], right["DESCRIPTION"], right["UNIT_MEAS"],
    )
    evidence = dict(score)
    evidence["critical_mismatches"] = list(score.get("critical_mismatches") or []) + list(
        discriminator.protected_conflicts
    )
    return classify_identity_edge(evidence).edge_class.value


def _evidence_summary(evidence):
    return {
        "fingerprint": evidence.evidence_fingerprint,
        "bucket": classify_shadow_evidence(evidence).value,
        "facts": [
            {
                "channel": item.channel.value,
                "semantic_key": item.semantic_key,
                "normalized_matches": list(item.normalized_matches),
                "reason_code": item.reason_code,
                "sources_1": [source.source_field.value for source in item.source_observations_1],
                "sources_2": [source.source_field.value for source in item.source_observations_2],
            }
            for item in evidence.facts
        ],
    }


def audit_accepted_group_edges(csv_path: Path, workbook_path: Path, *, target_pairs=()):
    started = perf_counter()
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        source_rows = list(csv.DictReader(handle))
    book = load_workbook(workbook_path, data_only=False, read_only=True)
    sheet = book["Group Data"]
    headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    groups = defaultdict(list)
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values, strict=True))
        source_index, stable_ref = _source_row_and_ref(
            row["Source Row / Stable Record Reference"]
        )
        source = source_rows[source_index]
        groups[row["Canonical Group ID"]].append((stable_ref, _engine_record(source)))

    signature_cache = {}
    pair_rows = []
    comparison_seconds = 0.0
    target_keys = {frozenset(pair) for pair in target_pairs}
    target_results = {}
    for members in groups.values():
        for (left_ref, left), (right_ref, right) in combinations(members, 2):
            for reference, record in ((left_ref, left), (right_ref, right)):
                if reference not in signature_cache:
                    signature_cache[reference] = derive_identity_signature(
                        record, record_reference=reference
                    )
            pair_started = perf_counter()
            evidence = derive_signed_identity_evidence(
                signature_cache[left_ref], signature_cache[right_ref]
            )
            comparison_seconds += perf_counter() - pair_started
            current = _current_edge_class(left, right)
            bucket = classify_shadow_evidence(evidence).value
            pair_rows.append((current, bucket))
            part_key = frozenset((left["PART_NO"], right["PART_NO"]))
            if part_key in target_keys:
                target_results[" / ".join(sorted(part_key))] = {
                    "current_edge_class": current,
                    **_evidence_summary(evidence),
                }

    bucket_counts = Counter(bucket for _current, bucket in pair_rows)
    lexical_bucket = "SHADOW_LEXICAL_ONLY_OR_UNRESOLVED"
    return {
        "pair_source": "R12 read-only Group Data accepted-group complete-pair edges",
        "group_count": len(groups),
        "real_pairs_evaluated": len(pair_rows),
        "shadow_bucket_counts": dict(sorted(bucket_counts.items())),
        "current_strong_shadow_lexical_only": sum(
            current == "STRONG_SUPPORT" and bucket == lexical_bucket
            for current, bucket in pair_rows
        ),
        "current_review_shadow_lexical_only": sum(
            current == "REVIEW_SUPPORT" and bucket == lexical_bucket
            for current, bucket in pair_rows
        ),
        "accepted_group_edges_shadow_lexical_only": sum(
            bucket == lexical_bucket for _current, bucket in pair_rows
        ),
        "existing_cannot_link_pairs_in_source": sum(
            current == "CANNOT_LINK" for current, _bucket in pair_rows
        ),
        "target_pairs": target_results,
        "pair_comparison_seconds": comparison_seconds,
        "pair_comparisons_per_second": (
            len(pair_rows) / comparison_seconds if comparison_seconds else None
        ),
        "wall_time_seconds": perf_counter() - started,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("csv_path", type=Path)
    parser.add_argument("workbook_path", type=Path)
    parser.add_argument("--target-pair", action="append", default=[])
    args = parser.parse_args()
    target_pairs = tuple(tuple(value.split("::", 1)) for value in args.target_pair)
    print(json.dumps(
        audit_accepted_group_edges(
            args.csv_path, args.workbook_path, target_pairs=target_pairs
        ),
        ensure_ascii=True, sort_keys=True, indent=2,
    ))


if __name__ == "__main__":
    main()
