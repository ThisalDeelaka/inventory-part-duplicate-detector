"""Forensic R18B reference recovery and deterministic quality evaluation.

This offline module never supplies human labels or changes detector behavior.
It recovers literal human fields by logical header name, certifies reviewer-
visible source cells against the immutable template, and joins only frozen R18
evidence snapshots.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sqlite3
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from app.benchmarks.human_identity_review_dataset import (
    CONFIDENCES,
    LABELS,
    REASON_CODES,
    VISIBLE_RECORD_FIELDS,
)


RECOVERY_VERSION = "r18b-primary-expert-recovery-v1"
EVALUATION_VERSION = "r18b-deterministic-quality-evaluation-v1"
REFERENCE_TYPE = "SINGLE_SENIOR_DOMAIN_EXPERT_RECOVERED"
EXPECTED_HEADERS = (
    "review_pair_id",
    *(f"record_a_{field}" for field in VISIBLE_RECORD_FIELDS),
    *(f"record_b_{field}" for field in VISIBLE_RECORD_FIELDS),
    "identity_label",
    "confidence",
    "reason_code",
    "reviewer_comment",
)
SOURCE_HEADERS = EXPECTED_HEADERS[1:19]
HUMAN_HEADERS = EXPECTED_HEADERS[19:]
SYSTEM_CLASSES = (
    "STRONG_SUPPORT",
    "REVIEW_SUPPORT",
    "NON_GROUPABLE",
    "CANNOT_LINK",
)


class R18BRecoveryError(ValueError):
    """A forensic recovery gate failed closed."""


@dataclass(frozen=True)
class RecoveredRow:
    review_pair_id: str
    identity_label: str
    confidence: str
    reason_code: str
    reviewer_comment: str


@dataclass(frozen=True)
class RecoveryResult:
    rows: tuple[RecoveredRow, ...]
    header_row: int
    logical_column_mapping: dict[str, int]
    extra_columns: tuple[dict[str, object], ...]
    source_cells_compared: int
    source_cells_different: int
    submission_sha256: str
    original_template_sha256: str


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _find_header(sheet) -> tuple[int, tuple[object, ...]]:
    matches = []
    for row_number, cells in enumerate(
        sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 50)), start=1
    ):
        values = tuple(cell.value for cell in cells)
        if "review_pair_id" in values:
            matches.append((row_number, values))
    if len(matches) != 1:
        raise R18BRecoveryError("R18B_R1_SCHEMA_MAPPING_AMBIGUOUS")
    return matches[0]


def _logical_mapping(headers: tuple[object, ...]) -> dict[str, int]:
    mapping = {}
    for name in EXPECTED_HEADERS:
        indexes = [index for index, value in enumerate(headers) if value == name]
        if len(indexes) != 1:
            raise R18BRecoveryError("R18B_R1_SCHEMA_MAPPING_AMBIGUOUS")
        mapping[name] = indexes[0]
    return mapping


def _read_expected_rows(sheet, header_row, mapping, expected_ids):
    rows = {}
    max_column = max(mapping.values()) + 1
    for cells in sheet.iter_rows(
        min_row=header_row + 1,
        max_row=header_row + len(expected_ids),
        max_col=max(sheet.max_column, max_column),
    ):
        pair_id = cells[mapping["review_pair_id"]].value
        if not isinstance(pair_id, str) or not pair_id:
            raise R18BRecoveryError("R18B_R1_PAIR_ID_INTEGRITY_FAILED")
        if pair_id in rows:
            raise R18BRecoveryError("R18B_R1_PAIR_ID_INTEGRITY_FAILED")
        rows[pair_id] = cells
    if set(rows) != expected_ids:
        raise R18BRecoveryError("R18B_R1_PAIR_ID_INTEGRITY_FAILED")
    return rows


def recover_primary_reference(
    original_template: Path,
    submission: Path,
    expected_ids: set[str],
) -> RecoveryResult:
    """Recover a submission without modifying it or trusting physical columns."""
    original = load_workbook(original_template, data_only=False, read_only=True)
    submitted = load_workbook(submission, data_only=False, read_only=True)
    if "Review" not in original.sheetnames or "Review" not in submitted.sheetnames:
        raise R18BRecoveryError("R18B_R1_SCHEMA_MAPPING_AMBIGUOUS")
    original_sheet = original["Review"]
    submitted_sheet = submitted["Review"]
    original_header_row, original_headers = _find_header(original_sheet)
    submitted_header_row, submitted_headers = _find_header(submitted_sheet)
    original_mapping = _logical_mapping(original_headers)
    submitted_mapping = _logical_mapping(submitted_headers)
    original_rows = _read_expected_rows(
        original_sheet, original_header_row, original_mapping, expected_ids
    )
    submitted_rows = _read_expected_rows(
        submitted_sheet, submitted_header_row, submitted_mapping, expected_ids
    )

    extra_indexes = tuple(
        index for index, value in enumerate(submitted_headers)
        if value not in EXPECTED_HEADERS
    )
    extra_columns = []
    for index in extra_indexes:
        formula_count = int(
            submitted_sheet.cell(submitted_header_row, index + 1).data_type == "f"
        ) + sum(
            submitted_rows[pair_id][index].data_type == "f"
            for pair_id in expected_ids
        )
        extra_columns.append({
            "one_based_column": index + 1,
            "header": str(submitted_headers[index] or ""),
            "formula_cell_count": formula_count,
            "authoritative": False,
        })

    source_cells_different = 0
    recovered = []
    for pair_id in sorted(expected_ids):
        source = original_rows[pair_id]
        candidate = submitted_rows[pair_id]
        for name in SOURCE_HEADERS:
            original_cell = source[original_mapping[name]]
            submitted_cell = candidate[submitted_mapping[name]]
            if original_cell.data_type == "f" or submitted_cell.data_type == "f":
                raise R18BRecoveryError("R18B_R1_FORMULA_CONTAMINATION")
            source_cells_different += original_cell.value != submitted_cell.value
        values = {}
        for name in HUMAN_HEADERS:
            cell = candidate[submitted_mapping[name]]
            if cell.data_type == "f":
                raise R18BRecoveryError("R18B_R1_FORMULA_CONTAMINATION")
            if cell.value is not None and not isinstance(cell.value, str):
                raise R18BRecoveryError("R18B_R1_HUMAN_INPUT_INVALID")
            values[name] = "" if cell.value is None else cell.value
        if values["identity_label"] not in LABELS:
            raise R18BRecoveryError("R18B_R1_HUMAN_INPUT_INVALID")
        if values["confidence"] not in CONFIDENCES:
            raise R18BRecoveryError("R18B_R1_HUMAN_INPUT_INVALID")
        if values["reason_code"] not in REASON_CODES:
            raise R18BRecoveryError("R18B_R1_HUMAN_INPUT_INVALID")
        recovered.append(RecoveredRow(pair_id, **values))
    if source_cells_different:
        raise R18BRecoveryError("R18B_R1_REVIEWER_VISIBLE_SOURCE_CHANGED")
    return RecoveryResult(
        rows=tuple(recovered),
        header_row=submitted_header_row,
        logical_column_mapping={name: index + 1 for name, index in submitted_mapping.items()},
        extra_columns=tuple(extra_columns),
        source_cells_compared=len(expected_ids) * len(SOURCE_HEADERS),
        source_cells_different=0,
        submission_sha256=sha256_file(submission),
        original_template_sha256=sha256_file(original_template),
    )


def write_recovered_reference(path: Path, recovery: RecoveryResult) -> str:
    headers = (
        "review_pair_id", "identity_label", "confidence", "reason_code",
        "reviewer_comment", "reference_type", "submission_sha256",
        "original_template_sha256", "recovery_version",
    )
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        for row in recovery.rows:
            writer.writerow({
                **row.__dict__,
                "reference_type": REFERENCE_TYPE,
                "submission_sha256": recovery.submission_sha256,
                "original_template_sha256": recovery.original_template_sha256,
                "recovery_version": RECOVERY_VERSION,
            })
    return sha256_file(path)


def _matrix(rows):
    counts = {
        label: {edge_class: 0 for edge_class in SYSTEM_CLASSES}
        for label in LABELS
    }
    for row in rows:
        counts[row["expert_label"]][row["system_edge_class"]] += 1
    row_percentages = {}
    for label, values in counts.items():
        total = sum(values.values())
        row_percentages[label] = {
            key: round(value * 100 / total, 6) if total else None
            for key, value in values.items()
        }
    column_percentages = {}
    for edge_class in SYSTEM_CLASSES:
        total = sum(counts[label][edge_class] for label in LABELS)
        column_percentages[edge_class] = {
            label: round(counts[label][edge_class] * 100 / total, 6) if total else None
            for label in LABELS
        }
    return counts, row_percentages, column_percentages


def _agreement_category(label, edge_class):
    if edge_class not in SYSTEM_CLASSES:
        return "NOT_COMPARABLE", "NOT_APPLICABLE"
    if label == "SAME_IDENTITY":
        if edge_class in ("STRONG_SUPPORT", "REVIEW_SUPPORT"):
            return "SUPPORT_ALIGNED", "NONE"
        if edge_class == "CANNOT_LINK":
            return "FALSE_SEPARATION", "SEVERE"
        return "MISSED_IDENTITY_OPPORTUNITY", "MATERIAL"
    if label == "DIFFERENT_IDENTITY":
        if edge_class in ("CANNOT_LINK", "NON_GROUPABLE"):
            return "SAFETY_ALIGNED", "NONE"
        if edge_class == "STRONG_SUPPORT":
            return "FALSE_STRONG_SUPPORT", "SEVERE"
        return "NOISY_REVIEW_SUPPORT", "REVIEW_BURDEN"
    if edge_class == "NON_GROUPABLE":
        return "CAUTIOUS_WITH_INSUFFICIENT", "NONE"
    return "OVERCOMMIT_SIGNAL", "CAUTION"


def _reason_inconsistent(label, reason):
    if label == "SAME_IDENTITY":
        return reason.startswith("DIFFERENT_") or reason == "COPIED_OR_GENERIC_TEXT_NOT_ENOUGH"
    if label == "DIFFERENT_IDENTITY":
        return reason.startswith("SAME_")
    return reason not in {"INSUFFICIENT_FIELDS", "OTHER", "COPIED_OR_GENERIC_TEXT_NOT_ENOUGH"}


def evaluate(
    reference: RecoveryResult,
    secondary: RecoveryResult,
    mapping_path: Path,
    database_path: Path,
    output_csv: Path,
) -> dict[str, object]:
    with mapping_path.open(encoding="utf-8", newline="") as handle:
        mapping_rows = tuple(csv.DictReader(handle))
    by_id = {row.review_pair_id: row for row in reference.rows}
    secondary_by_id = {row.review_pair_id: row for row in secondary.rows}
    mapping_by_id = {row["review_pair_id"]: row for row in mapping_rows}
    if len(mapping_by_id) != len(mapping_rows) or set(mapping_by_id) != set(by_id):
        raise R18BRecoveryError("R18B_R1_PAIR_ID_INTEGRITY_FAILED")
    if set(secondary_by_id) != set(by_id):
        raise R18BRecoveryError("R18B_R1_PAIR_ID_INTEGRITY_FAILED")

    connection = sqlite3.connect(
        f"file:{database_path.resolve().as_posix()}?mode=ro", uri=True
    )
    evidence = {
        row[0]: row[1:]
        for row in connection.execute(
            "SELECT evidence_fingerprint, generic_evidence_json, "
            "technical_evidence_json, classification_reason_codes_json, "
            "protected_conflicts_json "
            "FROM identity_evidence_edge_snapshot WHERE evidence_run_id = 7"
        )
    }
    connection.close()

    rows = []
    for pair_id in sorted(by_id):
        expert = by_id[pair_id]
        machine = mapping_by_id[pair_id]
        edge_class = machine["current_gf4_edge_class"]
        if edge_class not in SYSTEM_CLASSES:
            # Three diagnostic-only historical controls have no frozen GF4 edge.
            edge_class = "HISTORICAL_DIAGNOSTIC"
        snapshot = evidence.get(machine["evidence_fingerprint"])
        generic_reason = ""
        system_reason_codes = ""
        technical_contradiction = False
        if snapshot:
            generic_reason = str(json.loads(snapshot[0]).get("generic_guard_reason") or "")
            system_reason_codes = "|".join(json.loads(snapshot[2]))
            technical_contradiction = bool(json.loads(snapshot[3]))
        shadow = machine["r16_shadow_bucket"]
        agreement, severity = _agreement_category(expert.identity_label, edge_class)
        rows.append({
            "review_pair_id": pair_id,
            "panel_membership": machine["panel_membership"],
            "sampling_stratum": machine["evaluation_stratum"],
            "diagnostic_family": machine["diagnostic_family"],
            "expert_label": expert.identity_label,
            "expert_confidence": expert.confidence,
            "expert_reason": expert.reason_code,
            "system_edge_class": edge_class,
            "shadow_bucket": shadow,
            "generic_only": str(bool(generic_reason)).lower(),
            "generic_guard_reason": generic_reason,
            "trusted_identity_present": str(shadow in {
                "SHADOW_TRUSTED_IDENTITY_PRESENT", "SHADOW_MIXED_EVIDENCE"
            }).lower(),
            "lexical_only_or_unresolved": str(
                shadow == "SHADOW_LEXICAL_ONLY_OR_UNRESOLVED"
            ).lower(),
            "description_dominant_support": str(
                shadow == "SHADOW_LEXICAL_ONLY_OR_UNRESOLVED"
                and edge_class in {"STRONG_SUPPORT", "REVIEW_SUPPORT"}
            ).lower(),
            "cannot_link_present": machine["protected_conflict_present"],
            "technical_contradiction_present": str(technical_contradiction).lower(),
            "system_reason_codes": system_reason_codes,
            "agreement_category": agreement,
            "severity_category": severity,
        })

    headers = tuple(rows[0])
    with output_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)

    evaluation_rows = [row for row in rows if "EVALUATION" in row["panel_membership"].split("|")]
    diagnostic_rows = [row for row in rows if "DIAGNOSTIC" in row["panel_membership"].split("|")]
    matrix, row_pct, column_pct = _matrix(evaluation_rows)
    confidence = {
        band: {
            label: dict(Counter(
                row["system_edge_class"] for row in evaluation_rows
                if row["expert_confidence"] == band and row["expert_label"] == label
            ))
            for label in LABELS
        }
        for band in CONFIDENCES
    }
    provenance = {}
    for bucket in sorted({row["shadow_bucket"] for row in evaluation_rows}):
        selected = [row for row in evaluation_rows if row["shadow_bucket"] == bucket]
        provenance[bucket] = {
            "sample_count": len(selected),
            "human_labels": dict(Counter(row["expert_label"] for row in selected)),
            "high_confidence_count": sum(row["expert_confidence"] == "HIGH" for row in selected),
            "system_classes": dict(Counter(row["system_edge_class"] for row in selected)),
        }
    generic = {}
    generic_scope = [row for row in evaluation_rows if row["generic_only"] == "true" or row["lexical_only_or_unresolved"] == "true"]
    for label in LABELS:
        for band in CONFIDENCES:
            selected = [row for row in generic_scope if row["expert_label"] == label and row["expert_confidence"] == band]
            generic[f"{label}__{band}"] = dict(Counter(row["system_edge_class"] for row in selected))

    a_b_agreement = sum(
        by_id[pair_id].identity_label == secondary_by_id[pair_id].identity_label
        for pair_id in by_id
    )
    secondary_status = (
        "DEGRADED_ONE_CLASS_RESPONSE"
        if len({row.identity_label for row in secondary.rows}) == 1
        else "MULTI_CLASS_RESPONSE"
    )
    agreement_by_a_label = {
        label: {
            "agreed": sum(
                by_id[pair_id].identity_label == secondary_by_id[pair_id].identity_label
                for pair_id in by_id if by_id[pair_id].identity_label == label
            ),
            "total": sum(row.identity_label == label for row in reference.rows),
        }
        for label in LABELS
    }
    agreement_by_a_confidence = {
        band: {
            "agreed": sum(
                by_id[pair_id].identity_label == secondary_by_id[pair_id].identity_label
                for pair_id in by_id if by_id[pair_id].confidence == band
            ),
            "total": sum(row.confidence == band for row in reference.rows),
        }
        for band in CONFIDENCES
    }
    disagreements = Counter()
    for pair_id in by_id:
        a_label = by_id[pair_id].identity_label
        b_label = secondary_by_id[pair_id].identity_label
        if a_label != b_label:
            disagreements[f"{a_label}__{b_label}"] += 1

    population_counts = {
        "CANNOT_LINK": 139,
        "NON_GROUPABLE__SHADOW_INSUFFICIENT": 14121,
        "NON_GROUPABLE__SHADOW_LEXICAL_ONLY_OR_UNRESOLVED": 4911,
        "NON_GROUPABLE__SHADOW_TRUSTED_IDENTITY_PRESENT": 16,
        "REVIEW_SUPPORT__SHADOW_INSUFFICIENT": 18,
        "REVIEW_SUPPORT__SHADOW_LEXICAL_ONLY_OR_UNRESOLVED": 879,
        "REVIEW_SUPPORT__SHADOW_MIXED_EVIDENCE": 1,
        "REVIEW_SUPPORT__SHADOW_TRUSTED_IDENTITY_PRESENT": 41,
        "STRONG_SUPPORT__SHADOW_LEXICAL_ONLY_OR_UNRESOLVED": 261,
        "STRONG_SUPPORT__SHADOW_MIXED_EVIDENCE": 1,
        "STRONG_SUPPORT__SHADOW_TRUSTED_IDENTITY_PRESENT": 7,
    }
    sample_counts = Counter(row["sampling_stratum"] for row in evaluation_rows)
    weighted = defaultdict(float)
    covered_population = 0
    unsampled = []
    for stratum, population in population_counts.items():
        sample = sample_counts[stratum]
        if not sample:
            unsampled.append({"stratum": stratum, "population": population})
            continue
        covered_population += population
        weight = population / sample
        for row in evaluation_rows:
            if row["sampling_stratum"] == stratum:
                weighted[f"{row['expert_label']}__{row['system_edge_class']}"] += weight

    known_controls = [
        {
            "diagnostic_family": row["diagnostic_family"],
            "expert_label": row["expert_label"],
            "expert_confidence": row["expert_confidence"],
            "expert_reason": row["expert_reason"],
            "system_edge_class": row["system_edge_class"],
            "shadow_bucket": row["shadow_bucket"],
            "agreement_category": row["agreement_category"],
        }
        for row in diagnostic_rows if row["diagnostic_family"]
    ]
    return {
        "evaluation_version": EVALUATION_VERSION,
        "unique_pair_count": len(rows),
        "evaluation_pair_count": len(evaluation_rows),
        "diagnostic_pair_count": len(diagnostic_rows),
        "panel_overlap_count": sum(row["panel_membership"] == "DIAGNOSTIC|EVALUATION" for row in rows),
        "matrix_counts": matrix,
        "matrix_row_percentages": row_pct,
        "matrix_column_percentages": column_pct,
        "confidence_stratification": confidence,
        "generic_lexical_analysis": generic,
        "provenance_analysis": provenance,
        "expert_label_counts_all_unique_pairs": dict(Counter(row.identity_label for row in reference.rows)),
        "expert_confidence_counts_all_unique_pairs": dict(Counter(row.confidence for row in reference.rows)),
        "secondary_label_counts": dict(Counter(row.identity_label for row in secondary.rows)),
        "secondary_confidence_counts": dict(Counter(row.confidence for row in secondary.rows)),
        "secondary_reason_counts": dict(Counter(row.reason_code for row in secondary.rows)),
        "secondary_internal_inconsistency_count": sum(_reason_inconsistent(row.identity_label, row.reason_code) for row in secondary.rows),
        "secondary_review_status": secondary_status,
        "a_b_raw_agreement_count": a_b_agreement,
        "a_b_raw_agreement_percentage": round(a_b_agreement * 100 / len(rows), 6),
        "a_b_disagreement_types": dict(disagreements),
        "agreement_by_a_label": agreement_by_a_label,
        "agreement_by_a_confidence": agreement_by_a_confidence,
        "sampling_stratum_counts": dict(sample_counts),
        "weighted_full_population_estimate": "NOT_VALID" if unsampled else "VALID",
        "unsampled_strata": unsampled,
        "covered_candidate_population": covered_population,
        "covered_population_weighted_matrix_estimate": {
            key: round(value, 6) for key, value in sorted(weighted.items())
        },
        "known_controls": known_controls,
        "evaluation_csv_sha256": sha256_file(output_csv),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--original-a", required=True, type=Path)
    parser.add_argument("--submitted-a", required=True, type=Path)
    parser.add_argument("--original-b", required=True, type=Path)
    parser.add_argument("--submitted-b", required=True, type=Path)
    parser.add_argument("--mapping", required=True, type=Path)
    parser.add_argument("--database", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    arguments = parser.parse_args()
    with arguments.mapping.open(encoding="utf-8", newline="") as handle:
        expected_ids = {row["review_pair_id"] for row in csv.DictReader(handle)}
    primary = recover_primary_reference(arguments.original_a, arguments.submitted_a, expected_ids)
    secondary = recover_primary_reference(arguments.original_b, arguments.submitted_b, expected_ids)
    arguments.output_dir.mkdir(parents=True, exist_ok=True)
    reference_path = arguments.output_dir / "r18b_recovered_primary_expert_reference.csv"
    reference_sha = write_recovered_reference(reference_path, primary)
    evaluation_path = arguments.output_dir / "r18b_deterministic_quality_evaluation.csv"
    result = evaluate(primary, secondary, arguments.mapping, arguments.database, evaluation_path)
    result["recovery"] = {
        "classification": "R18B_R1_PRIMARY_EXPERT_REFERENCE_RECOVERED",
        "header_row": primary.header_row,
        "logical_column_mapping": primary.logical_column_mapping,
        "extra_columns": primary.extra_columns,
        "source_cells_compared": primary.source_cells_compared,
        "source_cells_different": primary.source_cells_different,
        "submission_sha256": primary.submission_sha256,
        "original_template_sha256": primary.original_template_sha256,
        "derived_reference_sha256": reference_sha,
        "derived_reference_pair_count": len(primary.rows),
    }
    summary_path = (
        arguments.output_dir / "r18b_deterministic_quality_evaluation_summary.json"
    )
    summary_path.write_text(
        json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
