"""Offline R18 pair-review dataset generation.

This module consumes a persisted candidate-edge population.  It never searches
for candidates, changes detector authority, calls a provider, or supplies a
human answer.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sqlite3
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app.engine.identity_signature_derivation import derive_identity_signature
from app.engine.signed_identity_evidence import (
    classify_shadow_evidence,
    derive_signed_identity_evidence,
)


R18_DATASET_VERSION = "gf12-a2-r18-human-identity-pairs-v1"
R18_SAMPLING_SEED = 1201
R18_SCAN_ID = 31
PROTECTED_CSV_SIZE = 3_265_800
PROTECTED_CSV_SHA256 = "8740777d660a16661585e6141de7be3309219b7758dd12486f3abb1a05b1110b"

LABELS = ("SAME_IDENTITY", "DIFFERENT_IDENTITY", "INSUFFICIENT_INFORMATION")
CONFIDENCES = ("HIGH", "MEDIUM", "LOW")
REASON_CODES = (
    "SAME_EXACT_OR_ALIAS_IDENTITY",
    "SAME_MODEL_AND_OBJECT",
    "DIFFERENT_OBJECT",
    "DIFFERENT_MODEL_OR_TYPE",
    "DIFFERENT_VARIANT_OR_SIDE",
    "DIFFERENT_CRITICAL_ATTRIBUTE",
    "COPIED_OR_GENERIC_TEXT_NOT_ENOUGH",
    "PART_NUMBER_EVIDENCE",
    "DESCRIPTION_EVIDENCE",
    "INSUFFICIENT_FIELDS",
    "OTHER",
)

EVALUATION_TARGETS = (
    ("STRONG_SUPPORT__SHADOW_TRUSTED_IDENTITY_PRESENT", 7),
    ("STRONG_SUPPORT__SHADOW_LEXICAL_ONLY_OR_UNRESOLVED", 42),
    ("STRONG_SUPPORT__SHADOW_MIXED_EVIDENCE", 1),
    ("REVIEW_SUPPORT__SHADOW_TRUSTED_IDENTITY_PRESENT", 41),
    ("REVIEW_SUPPORT__SHADOW_LEXICAL_ONLY_OR_UNRESOLVED", 57),
    ("REVIEW_SUPPORT__SHADOW_MIXED_EVIDENCE", 1),
    ("CANNOT_LINK", 50),
    ("NON_GROUPABLE__SHADOW_TRUSTED_IDENTITY_PRESENT", 16),
    ("NON_GROUPABLE__SHADOW_LEXICAL_ONLY_OR_UNRESOLVED", 45),
    ("NON_GROUPABLE__SHADOW_INSUFFICIENT", 40),
)

# Zero-based source rows. H10-H13 are the four R12 false-group pairs.
DIAGNOSTIC_PAIRS = (
    ("R17_CONTACT_CLEANER", 5194, 5195),
    ("R17_TURBINE_LUBRICATING_OIL", 5246, 5263),
    ("R17_FRANCIS_TURBINE_LOWER_BEARING", 5262, 5283),
    ("R17_PUMP_X500", 5294, 5295),
    ("R17_FAN_BLADE", 2306, 2385),
    ("R17_F30", 2526, 3045),
    ("R17_B38", 5126, 5131),
    ("H1_CARBON_STICK_PENCIL", 367, 369),
    ("H2_RIM_TYRE", 2869, 3554),
    ("H3_TABLE_NAIL", 5136, 5134),
    ("H4_CONDITION_DISCOUNT", 3848, 3802),
    ("H5_CLUTCH_DUST_CAP", 4129, 4127),
    ("H6_CLUTCH_COIL_SPRING", 4129, 4634),
    ("H7_DUST_CAP_COIL_SPRING", 4127, 4634),
    ("H8_DIRECTIONAL_SIDE", 4132, 4635),
    ("H9_BUFFER_MIRROR", 1512, 1511),
    ("H10_R12_EXERCISE", 1434, 1433),
    ("H11_R12_MODEL_VARIANT", 4991, 4925),
    ("H12_R12_OBJECT_CLASS", 5135, 5012),
    ("H13_R12_COPIED_COIL_TEXT", 3983, 3941),
)

VISIBLE_RECORD_FIELDS = (
    "part_number", "description_in_use", "description", "master_description",
    "type_designation", "dimension_quality", "uom", "part_type", "site",
)


@dataclass(frozen=True)
class ReviewRecord:
    record_id: int
    stable_ref: str
    source_row: int
    part_number: str
    description_in_use: str
    description: str
    master_description: str
    type_designation: str
    dimension_quality: str
    uom: str
    part_type: str
    site: str


@dataclass(frozen=True)
class CandidatePair:
    left: ReviewRecord
    right: ReviewRecord
    edge_class: str
    shadow_bucket: str
    protected_conflict: bool
    resolution_context: str
    evidence_fingerprint: str
    panel_memberships: tuple[str, ...] = ()
    diagnostic_family: str = ""
    evaluation_stratum: str = ""

    @property
    def canonical_refs(self) -> tuple[str, str]:
        return tuple(sorted((self.left.stable_ref, self.right.stable_ref)))  # type: ignore[return-value]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _safe_cell(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _engine_record(record: ReviewRecord) -> dict[str, str]:
    return {
        "record_ref_key": record.stable_ref,
        "PART_NO": record.part_number,
        "DESCRIPTION": record.description,
        "MASTER_DESCRIPTION": record.master_description,
        "TYPE_DESIGNATION": record.type_designation,
        "DIMENSION_QUALITY": record.dimension_quality,
        "CONTRACT": record.site,
        "UNIT_MEAS": record.uom,
        "TYPE_CODE": record.part_type,
    }


def _review_pair_id(refs: tuple[str, str]) -> str:
    value = hashlib.sha256((R18_DATASET_VERSION + "|" + "|".join(refs)).encode()).hexdigest()
    return "R18-" + value[:16].upper()


def _pair_key(pair: CandidatePair) -> tuple[str, str]:
    return pair.canonical_refs


def evaluation_stratum(pair: CandidatePair) -> str:
    if pair.edge_class == "CANNOT_LINK":
        return "CANNOT_LINK"
    return f"{pair.edge_class}__{pair.shadow_bucket}"


def _selection_rank(pair: CandidatePair, stratum: str) -> str:
    payload = f"{R18_DATASET_VERSION}|{R18_SAMPLING_SEED}|{stratum}|{'|'.join(pair.canonical_refs)}"
    return hashlib.sha256(payload.encode()).hexdigest()


def select_evaluation_panel(
    population: tuple[CandidatePair, ...],
    targets: tuple[tuple[str, int], ...] = EVALUATION_TARGETS,
) -> tuple[tuple[CandidatePair, ...], dict[str, dict[str, int]]]:
    grouped: dict[str, list[CandidatePair]] = defaultdict(list)
    for pair in population:
        grouped[evaluation_stratum(pair)].append(pair)
    selected: list[CandidatePair] = []
    counts: dict[str, dict[str, int]] = {}
    for stratum, target in targets:
        ranked = sorted(grouped[stratum], key=lambda item: (_selection_rank(item, stratum), item.canonical_refs))
        if len(ranked) < target:
            raise ValueError(f"R18_EVALUATION_STRATUM_SHORTAGE:{stratum}:{len(ranked)}:{target}")
        chosen = ranked[:target]
        selected.extend(replace(item, evaluation_stratum=stratum) for item in chosen)
        counts[stratum] = {"population": len(ranked), "sample": len(chosen)}
    if len({_pair_key(item) for item in selected}) != len(selected):
        raise ValueError("R18_EVALUATION_PAIR_DUPLICATE")
    return tuple(selected), counts


def combine_panels(
    diagnostics: tuple[CandidatePair, ...], evaluation: tuple[CandidatePair, ...]
) -> tuple[CandidatePair, ...]:
    combined: dict[tuple[str, str], CandidatePair] = {}
    for pair, panel in ((*((item, "DIAGNOSTIC") for item in diagnostics),
                         *((item, "EVALUATION") for item in evaluation))):
        key = _pair_key(pair)
        if key not in combined:
            combined[key] = replace(pair, panel_memberships=(panel,))
        else:
            current = combined[key]
            combined[key] = replace(
                current,
                panel_memberships=tuple(sorted(set(current.panel_memberships + (panel,)))),
                diagnostic_family=current.diagnostic_family or pair.diagnostic_family,
                evaluation_stratum=current.evaluation_stratum or pair.evaluation_stratum,
            )
    return tuple(sorted(combined.values(), key=lambda item: _review_pair_id(item.canonical_refs)))


def _readonly_connection(path: Path) -> sqlite3.Connection:
    connection = sqlite3.connect(f"file:{path.resolve().as_posix()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    return connection


def load_candidate_population(
    database_path: Path, source_csv_path: Path, *, scan_id: int = R18_SCAN_ID
) -> tuple[tuple[CandidatePair, ...], tuple[CandidatePair, ...], dict[str, object]]:
    if source_csv_path.stat().st_size != PROTECTED_CSV_SIZE or sha256_file(source_csv_path) != PROTECTED_CSV_SHA256:
        raise ValueError("R18_PROTECTED_CSV_INTEGRITY_FAILURE")
    with source_csv_path.open(encoding="utf-8-sig", newline="") as handle:
        source_rows = list(csv.DictReader(handle))

    connection = _readonly_connection(database_path)
    try:
        scan = connection.execute(
            "select id, status, total_records from duplicate_scan where id=?", (scan_id,)
        ).fetchone()
        if not scan or scan["status"] != "COMPLETED" or scan["total_records"] != len(source_rows):
            raise ValueError("R18_SOURCE_SCAN_INVALID")
        run = connection.execute(
            "select * from identity_evidence_run where scan_id=? and status='COMPLETED' order by id desc limit 1",
            (scan_id,),
        ).fetchone()
        if not run:
            raise ValueError("R18_SOURCE_EVIDENCE_RUN_MISSING")
        discovery = connection.execute(
            "select * from identity_discovery_run where scan_id=? and status='COMPLETED' order by id desc limit 1",
            (scan_id,),
        ).fetchone()
        if not discovery or discovery["proposal_count"] <= 0:
            raise ValueError("R18_SOURCE_EDGE_POPULATION_UNAVAILABLE")

        db_records = {
            row["id"]: row for row in connection.execute(
                "select * from scan_record_snapshot where scan_id=?", (scan_id,)
            )
        }
        by_source_row: dict[int, ReviewRecord] = {}
        by_id: dict[int, ReviewRecord] = {}
        for record_id, row in db_records.items():
            source_index = int(row["source_row_index"])
            source = source_rows[source_index]
            item = ReviewRecord(
                record_id=record_id,
                stable_ref=row["record_ref_key"],
                source_row=source_index,
                part_number=str(source.get("Part No") or ""),
                description_in_use=str(source.get("Part Description in Use") or ""),
                description=str(source.get("Part Description") or row["description"] or ""),
                master_description=str(source.get("Master Part Description") or ""),
                type_designation=str(source.get("Type Designation") or ""),
                dimension_quality=str(source.get("Dimension/ Quality") or ""),
                uom=str(source.get("Inventory UoM") or row["uom"] or ""),
                part_type=str(source.get("Part Type") or row["type_code"] or ""),
                site=str(source.get("Site") or row["contract"] or ""),
            )
            by_source_row[source_index] = item
            by_id[record_id] = item

        conflict_sets = _resolution_member_sets(connection, "conflict", scan_id)
        deferred_sets = _resolution_member_sets(connection, "deferred", scan_id)
        accepted_pairs = {
            tuple(sorted((row["record_id_1"], row["record_id_2"])))
            for row in connection.execute(
                "select record_id_1, record_id_2 from g2_v2_internal_evidence "
                "where projection_run_id=(select max(id) from g2_v2_projection_run where scan_id=?)",
                (scan_id,),
            )
        }
        signatures: dict[int, object] = {}
        population = []
        edge_lookup: dict[tuple[int, int], CandidatePair] = {}
        for edge in connection.execute(
            "select * from identity_evidence_edge_snapshot where scan_id=? and evidence_run_id=?",
            (scan_id, run["id"]),
        ):
            left, right = by_id[edge["record_id_1"]], by_id[edge["record_id_2"]]
            for item in (left, right):
                if item.record_id not in signatures:
                    signatures[item.record_id] = derive_identity_signature(
                        _engine_record(item), record_reference=item.stable_ref
                    )
            shadow = classify_shadow_evidence(derive_signed_identity_evidence(
                signatures[left.record_id], signatures[right.record_id]
            )).value
            ids = tuple(sorted((left.record_id, right.record_id)))
            context = "ACCEPTED_GROUP_EDGE" if ids in accepted_pairs else _pair_resolution_context(ids, conflict_sets, deferred_sets)
            item = CandidatePair(
                left=left,
                right=right,
                edge_class=edge["edge_class"],
                shadow_bucket=shadow,
                protected_conflict=bool(json.loads(edge["protected_conflicts_json"] or "[]")),
                resolution_context=context,
                evidence_fingerprint=edge["evidence_fingerprint"],
            )
            population.append(item)
            edge_lookup[ids] = item

        diagnostics = []
        for family, left_row, right_row in DIAGNOSTIC_PAIRS:
            if left_row not in by_source_row or right_row not in by_source_row:
                raise ValueError(f"R18_DIAGNOSTIC_SOURCE_ROW_MISSING:{family}")
            left, right = by_source_row[left_row], by_source_row[right_row]
            key = tuple(sorted((left.record_id, right.record_id)))
            base = edge_lookup.get(key)
            if base is None:
                left_sig = derive_identity_signature(_engine_record(left), record_reference=left.stable_ref)
                right_sig = derive_identity_signature(_engine_record(right), record_reference=right.stable_ref)
                shadow = classify_shadow_evidence(derive_signed_identity_evidence(left_sig, right_sig)).value
                base = CandidatePair(left, right, "HISTORICAL_DIAGNOSTIC", shadow, False, "DIAGNOSTIC_ONLY", "")
            diagnostics.append(replace(base, diagnostic_family=family))

        metadata = {
            "scan_id": scan_id,
            "scan_status": scan["status"],
            "source_record_count": len(by_id),
            "candidate_population_count": len(population),
            "discovery_run_id": discovery["id"],
            "discovery_fingerprint": discovery["discovery_fingerprint"],
            "evidence_run_id": run["id"],
            "evidence_configuration_fingerprint": run["configuration_fingerprint"],
            "protected_csv_size": PROTECTED_CSV_SIZE,
            "protected_csv_sha256": PROTECTED_CSV_SHA256,
        }
        if len(population) != discovery["proposal_count"]:
            raise ValueError("R18_SOURCE_EDGE_POPULATION_RECONCILIATION_FAILURE")
        return tuple(population), tuple(diagnostics), metadata
    finally:
        connection.close()


def _resolution_member_sets(connection: sqlite3.Connection, kind: str, scan_id: int) -> tuple[frozenset[int], ...]:
    table = f"identity_resolution_{kind}_snapshot"
    member = f"identity_resolution_{kind}_member"
    key = f"{kind}_snapshot_id"
    rows = connection.execute(
        f"select m.{key}, m.record_id from {member} m join {table} s on s.id=m.{key} "
        "where s.scan_id=? order by m." + key + ", m.record_id",
        (scan_id,),
    )
    groups: dict[int, set[int]] = defaultdict(set)
    for row in rows:
        groups[row[0]].add(row[1])
    return tuple(frozenset(value) for value in groups.values())


def _pair_resolution_context(ids: tuple[int, int], conflicts: tuple[frozenset[int], ...], deferred: tuple[frozenset[int], ...]) -> str:
    pair = set(ids)
    if any(pair <= members for members in conflicts):
        return "CONFLICT"
    if any(pair <= members for members in deferred):
        return "DEFERRED"
    return "UNASSIGNED_OR_OTHER_CANDIDATE"


def _instructions() -> tuple[str, ...]:
    return (
        "Question: Do these two records represent the same underlying physical inventory item?",
        "SAME_IDENTITY: evidence supports the same item despite wording, part-number, or site differences.",
        "DIFFERENT_IDENTITY: evidence supports a different object, model/type, variant/side, or critical specification.",
        "INSUFFICIENT_INFORMATION: the visible source data cannot responsibly establish same or different.",
        "Do not guess. Matching descriptions alone do not prove identity; different part numbers alone do not prove difference.",
        "Complete the identity_label, confidence, and one primary reason_code. Reviewer comment is optional.",
    )


def _review_headers() -> tuple[str, ...]:
    headers = ["review_pair_id"]
    for side in ("record_a", "record_b"):
        headers.extend(f"{side}_{field}" for field in VISIBLE_RECORD_FIELDS)
    headers.extend(("identity_label", "confidence", "reason_code", "reviewer_comment"))
    return tuple(headers)


def _visible_row(pair: CandidatePair) -> list[object]:
    records = sorted((pair.left, pair.right), key=lambda item: item.stable_ref)
    row: list[object] = [_review_pair_id(pair.canonical_refs)]
    for item in records:
        row.extend(getattr(item, field) for field in VISIBLE_RECORD_FIELDS)
    row.extend(("", "", "", ""))
    return [_safe_cell(value) for value in row]


def _new_workbook(title: str) -> Workbook:
    workbook = Workbook()
    workbook.properties.creator = "GF12A2-R18"
    workbook.properties.title = title
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    return workbook


def _save_reproducible_xlsx(workbook: Workbook, path: Path) -> None:
    raw = io.BytesIO()
    workbook.save(raw)
    source = zipfile.ZipFile(io.BytesIO(raw.getvalue()), "r")
    normalized = io.BytesIO()
    with zipfile.ZipFile(normalized, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = source.getinfo(name).external_attr
            content = source.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"<dcterms:modified[^>]*>[^<]*</dcterms:modified>",
                    b'<dcterms:modified xsi:type="dcterms:W3CDTF">2000-01-01T00:00:00Z</dcterms:modified>',
                    content,
                )
            target.writestr(info, content)
    source.close()
    path.write_bytes(normalized.getvalue())


def write_reviewer_workbook(path: Path, pairs: tuple[CandidatePair, ...], reviewer_code: str) -> None:
    workbook = _new_workbook(f"R18 blinded identity review - {reviewer_code}")
    guide = workbook.active
    guide.title = "Instructions"
    guide.append(["reviewer_code", reviewer_code])
    guide.append(["dataset_version", R18_DATASET_VERSION])
    for line in _instructions():
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    review = workbook.create_sheet("Review")
    review.append(_review_headers())
    for pair in pairs:
        review.append(_visible_row(pair))
    review.freeze_panes = "A2"
    review.auto_filter.ref = review.dimensions
    for column in review.columns:
        review.column_dimensions[column[0].column_letter].width = min(45, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
    label_col = review.max_column - 3
    confidence_col = review.max_column - 2
    reason_col = review.max_column - 1
    for col, values in ((label_col, LABELS), (confidence_col, CONFIDENCES), (reason_col, REASON_CODES)):
        validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        review.add_data_validation(validation)
        validation.add(f"{review.cell(2, col).coordinate}:{review.cell(review.max_row, col).coordinate}")
    _save_reproducible_xlsx(workbook, path)


def write_adjudication_workbook(path: Path, pairs: tuple[CandidatePair, ...]) -> None:
    workbook = _new_workbook("R18 empty adjudication template")
    sheet = workbook.active
    sheet.title = "Adjudication"
    headers = ("review_pair_id", "reviewer_a_label", "reviewer_b_label", "agreement", "adjudicated_label", "adjudicator", "adjudication_reason")
    sheet.append(headers)
    for pair in pairs:
        sheet.append([_review_pair_id(pair.canonical_refs), "", "", "", "", "", ""])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    _save_reproducible_xlsx(workbook, path)


def write_internal_mapping(path: Path, pairs: tuple[CandidatePair, ...]) -> None:
    headers = (
        "review_pair_id", "record_a_stable_ref", "record_b_stable_ref", "record_a_source_row",
        "record_b_source_row", "panel_membership", "diagnostic_family", "evaluation_stratum",
        "current_gf4_edge_class", "r16_shadow_bucket", "protected_conflict_present",
        "resolution_context", "evidence_fingerprint",
    )
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(headers)
        for pair in pairs:
            records = sorted((pair.left, pair.right), key=lambda item: item.stable_ref)
            writer.writerow([
                _review_pair_id(pair.canonical_refs), records[0].stable_ref, records[1].stable_ref,
                records[0].source_row, records[1].source_row, "|".join(pair.panel_memberships),
                pair.diagnostic_family, pair.evaluation_stratum, pair.edge_class, pair.shadow_bucket,
                str(pair.protected_conflict).lower(), pair.resolution_context, pair.evidence_fingerprint,
            ])


def validate_artifacts(output_dir: Path, expected_pairs: int) -> dict[str, object]:
    reviewer_paths = [output_dir / "human_identity_review_reviewer_a.xlsx", output_dir / "human_identity_review_reviewer_b.xlsx"]
    contents = []
    leakage_terms = {"score", "gf4", "group_id", "shadow", "expected", "stratum", "status"}
    for path in reviewer_paths:
        workbook = load_workbook(path, data_only=False, read_only=False)
        sheet = workbook["Review"]
        headers = tuple(cell.value for cell in sheet[1])
        if len(headers) != len(set(headers)) or headers != _review_headers():
            raise ValueError("R18_REVIEWER_HEADERS_INVALID")
        if any(any(term in str(header).casefold() for term in leakage_terms) for header in headers):
            raise ValueError("R18_BLINDING_FAILURE")
        rows = tuple(tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2))
        if len(rows) != expected_pairs:
            raise ValueError("R18_REVIEWER_ROW_COUNT_INVALID")
        label_indexes = [headers.index(name) for name in ("identity_label", "confidence", "reason_code", "reviewer_comment")]
        if any(row[index] not in (None, "") for row in rows for index in label_indexes):
            raise ValueError("R18_PREFILLED_LABEL_FAILURE")
        if any(cell.data_type == "f" for row in sheet.iter_rows() for cell in row):
            raise ValueError("R18_FORMULA_INJECTION_FAILURE")
        contents.append(rows)
    if contents[0] != contents[1]:
        raise ValueError("R18_REVIEWER_CONTENT_MISMATCH")
    with (output_dir / "human_identity_review_internal_mapping.csv").open(encoding="utf-8", newline="") as handle:
        mapping = list(csv.DictReader(handle))
    review_ids = {row[0] for row in contents[0]}
    if len(mapping) != expected_pairs or {row["review_pair_id"] for row in mapping} != review_ids:
        raise ValueError("R18_MAPPING_RECONCILIATION_FAILURE")
    return {
        "review_pair_count": expected_pairs,
        "prefilled_label_count": 0,
        "detector_output_leakage_count": 0,
        "formula_cell_count": 0,
        "reviewer_pair_content_identical": True,
        "mapping_reconciled": True,
    }


def generate(database_path: Path, source_csv_path: Path, output_dir: Path, *, scan_id: int = R18_SCAN_ID) -> dict[str, object]:
    population, diagnostics, source = load_candidate_population(database_path, source_csv_path, scan_id=scan_id)
    evaluation, sampling_counts = select_evaluation_panel(population)
    combined = combine_panels(diagnostics, evaluation)
    output_dir.mkdir(parents=True, exist_ok=True)
    reviewer_a = output_dir / "human_identity_review_reviewer_a.xlsx"
    reviewer_b = output_dir / "human_identity_review_reviewer_b.xlsx"
    mapping = output_dir / "human_identity_review_internal_mapping.csv"
    adjudication = output_dir / "human_identity_review_adjudication.xlsx"
    manifest = output_dir / "human_identity_review_manifest.json"
    write_reviewer_workbook(reviewer_a, combined, "REVIEWER_A")
    write_reviewer_workbook(reviewer_b, combined, "REVIEWER_B")
    write_internal_mapping(mapping, combined)
    write_adjudication_workbook(adjudication, combined)
    validation = validate_artifacts(output_dir, len(combined))
    artifacts = {
        path.name: {"sha256": sha256_file(path), "size_bytes": path.stat().st_size}
        for path in (reviewer_a, reviewer_b, mapping, adjudication)
    }
    panel_counts = Counter(panel for pair in combined for panel in pair.panel_memberships)
    overlap_count = sum(pair.panel_memberships == ("DIAGNOSTIC", "EVALUATION") for pair in combined)
    payload: dict[str, object] = {
        "dataset_version": R18_DATASET_VERSION,
        "classification": "R18_HUMAN_REVIEW_DATASET_PREPARED",
        "dataset_authorization": "EXPLICIT_R18_USER_AUTHORIZATION",
        "sampling_seed": R18_SAMPLING_SEED,
        "selection_rule": "SHA256(dataset_version|seed|stratum|canonical_record_refs), ascending",
        "source": source,
        "source_population_counts": dict(sorted(Counter(evaluation_stratum(pair) for pair in population).items())),
        "evaluation_sampling_counts": sampling_counts,
        "diagnostic_panel_size": len(diagnostics),
        "evaluation_panel_size": len(evaluation),
        "unique_review_pair_count": len(combined),
        "panel_membership_counts": dict(sorted(panel_counts.items())),
        "panel_overlap_count": overlap_count,
        "human_labels_completed": 0,
        "reviewer_visible_fields": list(_review_headers()),
        "artifacts": artifacts,
        "validation": validation,
        "provider_calls": 0,
        "normal_g2_v2_scan": "NOT_REQUIRED",
    }
    manifest.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    payload["manifest"] = {"sha256": sha256_file(manifest), "size_bytes": manifest.stat().st_size}
    return payload


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--database", required=True, type=Path)
    parser.add_argument("--source-csv", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--scan-id", type=int, default=R18_SCAN_ID)
    arguments = parser.parse_args()
    print(json.dumps(generate(arguments.database, arguments.source_csv, arguments.output_dir, scan_id=arguments.scan_id), indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
