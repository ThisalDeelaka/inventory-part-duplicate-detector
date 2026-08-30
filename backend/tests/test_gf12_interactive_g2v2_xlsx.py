"""R6 focused contracts for interactive authority and Excel package compatibility."""

from __future__ import annotations

import io
import json
import re
import zipfile
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.worksheet.cell_range import CellRange

from app.core.config import Settings
from app.db.models import (
    DuplicateCandidate,
    G2V2ProjectionRun,
    IdentityGroupProjectionRun,
    ScanOrchestrationRun,
    ScanOrchestrationStageResultRow,
    ShadowComparisonRun,
)
from app.identity_read.key_codec import serialize_versioned_identity_group_key
from app.llm.runtime import get_llm_settings
from app.main import app
from app.orchestration.contracts import (
    ProductScanAuthority,
    ScanOrchestrationMode,
    orchestration_mode_for_product_authority,
)
from app.services.identity_read_service import IdentityReadService


_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _csv_bytes():
    return (
        "Part No,Part Description,Site,Inventory UoM,Part Type,Commodity Group 1\n"
        "R6-A,SKF 6205 BEARING,S1,PCS,PURCHASED,MECH\n"
        "R6-B,SKF BEARING 6205,S1,PCS,PURCHASED,MECH\n"
        "R6-C,UNRELATED FILTER ELEMENT,S1,PCS,PURCHASED,FILTER\n"
    ).encode()


def _request_data(*, authority="current_product", scan_mode="SAME_SITE_DUPLICATE"):
    result = {
        "scan_name": "R6 browser-equivalent current product",
        "threshold": "75",
        "selected_fields": json.dumps(["CONTRACT", "UNIT_MEAS"]),
        "column_mapping": "{}",
        "sensitive_mode": "true",
        "scan_mode": scan_mode,
    }
    if authority is not None:
        result["product_authority"] = authority
    return result


def _configuration(mode="legacy_primary"):
    return Settings(
        identity_orchestration_mode=mode,
        group_first_shadow_comparison_enabled=True,
        llm_demo_enabled=False,
        llm_provider="none",
        group_llm_provider="none",
        groq_api_key="",
        anthropic_api_key="",
    )


def _upload(client, db, *, authority="current_product", configured="legacy_primary"):
    app.dependency_overrides[get_llm_settings] = lambda: _configuration(configured)
    try:
        response = client.post(
            "/api/scans/upload",
            files={"file": ("r6.csv", _csv_bytes(), "text/csv")},
            data=_request_data(authority=authority),
        )
    finally:
        app.dependency_overrides.pop(get_llm_settings, None)
    assert response.status_code == 200, response.text
    scan_id = response.json()["scan_id"]
    return scan_id, db.query(ScanOrchestrationRun).filter_by(scan_id=scan_id).one()


def _package(payload: bytes):
    archive = zipfile.ZipFile(io.BytesIO(payload))
    return {
        "archive": archive,
        "grouped": ElementTree.fromstring(archive.read("xl/worksheets/sheet2.xml")),
        "flat": ElementTree.fromstring(archive.read("xl/worksheets/sheet3.xml")),
        "table": ElementTree.fromstring(archive.read("xl/tables/table1.xml")),
        "rels": ElementTree.fromstring(
            archive.read("xl/worksheets/_rels/sheet3.xml.rels")
        ),
    }


def test_r6_2_to_r6_10_current_product_request_overrides_legacy_config(db, client):
    scan_id, run = _upload(client, db, configured="legacy_primary")
    assert run.policy_version == "group-first-orchestration-policy-v2"
    assert run.mode == "group_first_primary"
    assert run.primary_identity_pipeline == "GROUP_FIRST_GF1_GF6"
    assert run.visible_projection_contract == "G2_V2"
    assert run.compatibility_projection_required is False
    assert run.status == "COMPLETED"
    assert run.primary_identity_ready is True and run.visible_product_ready is True

    stages = db.query(ScanOrchestrationStageResultRow).filter_by(
        orchestration_run_id=run.id
    ).all()
    assert {row.stage_id: row.status for row in stages} == {
        "CANONICAL_CATALOG": "SUCCEEDED",
        "DISCOVERY": "SUCCEEDED",
        "SIGNED_EVIDENCE": "SUCCEEDED",
        "GROUP_RESOLUTION": "SUCCEEDED",
        "G2_V2_PROJECTION": "SUCCEEDED",
        "LEGACY_PAIR_COMPATIBILITY": "NOT_APPLICABLE",
        "G1_COMPATIBILITY_PROJECTION": "NOT_APPLICABLE",
        "G2_V1_COMPATIBILITY_PROJECTION": "NOT_APPLICABLE",
        "SHADOW_COMPARISON": "NOT_APPLICABLE",
    }
    assert db.query(DuplicateCandidate).filter_by(scan_id=scan_id).count() == 0
    assert db.query(IdentityGroupProjectionRun).filter_by(scan_id=scan_id).count() == 0
    assert db.query(ShadowComparisonRun).filter_by(scan_id=scan_id).count() == 0
    assert db.query(G2V2ProjectionRun).filter_by(scan_id=scan_id).count() == 1
    assert IdentityReadService(db).load_identity_read_snapshot(
        scan_id
    ).projection_contract.value == "G2_V2"


def test_r6_current_product_is_route_default_and_invalid_values_are_rejected(db, client):
    scan_id, run = _upload(
        client, db, authority=None, configured="legacy_primary"
    )
    assert run.mode == "group_first_primary"
    assert IdentityReadService(db).load_identity_read_snapshot(
        scan_id
    ).projection_contract.value == "G2_V2"
    invalid = client.post(
        "/api/scans/upload",
        files={"file": ("r6.csv", _csv_bytes(), "text/csv")},
        data=_request_data(authority="arbitrary_pipeline"),
    )
    assert invalid.status_code == 422


def test_r6_11_explicit_legacy_compatibility_remains_v1(db, client):
    scan_id, run = _upload(
        client, db,
        authority="legacy_compatibility",
        configured="group_first_primary",
    )
    assert run.mode == "legacy_primary"
    assert run.visible_projection_contract == "G2_V1"
    assert run.compatibility_projection_required is True
    assert db.query(IdentityGroupProjectionRun).filter_by(scan_id=scan_id).count() == 1
    assert IdentityReadService(db).load_identity_read_snapshot(
        scan_id
    ).projection_contract.value == "G2_V1"


def test_product_authority_contract_is_typed_and_distinct_from_business_scope():
    assert orchestration_mode_for_product_authority(
        ProductScanAuthority.CURRENT_PRODUCT
    ) == ScanOrchestrationMode.GROUP_FIRST_PRIMARY
    assert orchestration_mode_for_product_authority(
        ProductScanAuthority.LEGACY_COMPATIBILITY
    ) == ScanOrchestrationMode.LEGACY_PRIMARY
    assert {item.value for item in ProductScanAuthority} == {
        "current_product", "legacy_compatibility",
    }


def test_r6_16_to_r6_18_excel_table_owns_the_only_group_data_filter(db, client):
    scan_id, _ = _upload(client, db)
    response = client.get(
        f"/api/scans/{scan_id}/identity-read/system-groups/export.xlsx"
    )
    assert response.status_code == 200
    package = _package(response.content)
    archive, flat, table, rels = (
        package["archive"], package["flat"], package["table"], package["rels"]
    )
    assert archive.testzip() is None
    assert "xl/tables/table1.xml" in archive.namelist()
    assert flat.findall(f"{{{_MAIN_NS}}}autoFilter") == []
    filters = table.findall(f"{{{_MAIN_NS}}}autoFilter")
    assert len(filters) == 1 and filters[0].attrib["ref"] == table.attrib["ref"]

    table_parts = flat.findall(
        f"{{{_MAIN_NS}}}tableParts/{{{_MAIN_NS}}}tablePart"
    )
    assert len(table_parts) == 1
    relationship_id = table_parts[0].attrib[f"{{{_DOC_REL_NS}}}id"]
    relationships = {
        item.attrib["Id"]: item.attrib["Target"]
        for item in rels.findall(f"{{{_PKG_REL_NS}}}Relationship")
    }
    assert relationships[relationship_id] in {
        "/xl/tables/table1.xml", "../tables/table1.xml",
    }


def test_excel_table_range_columns_name_and_relationship_are_valid(db, client):
    scan_id, _ = _upload(client, db)
    payload = client.get(
        f"/api/scans/{scan_id}/identity-read/system-groups/export.xlsx"
    ).content
    package = _package(payload)
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    flat = workbook["Group Data"]
    expected_ref = f"A1:T{flat.max_row}"
    assert package["table"].attrib["ref"] == expected_ref
    assert package["table"].attrib["name"] == "SystemGroupData"
    assert package["table"].attrib["displayName"] == "SystemGroupData"
    assert re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", "SystemGroupData")

    columns_node = package["table"].find(f"{{{_MAIN_NS}}}tableColumns")
    columns = columns_node.findall(f"{{{_MAIN_NS}}}tableColumn")
    assert int(columns_node.attrib["count"]) == len(columns) == 20
    assert [int(item.attrib["id"]) for item in columns] == list(range(1, 21))
    assert len({item.attrib["name"] for item in columns}) == 20
    assert tuple(item.attrib["name"] for item in columns) == tuple(
        cell.value for cell in flat[1]
    )


def test_r6_19_merged_ranges_are_valid_and_non_overlapping(db, client):
    scan_id, _ = _upload(client, db)
    payload = client.get(
        f"/api/scans/{scan_id}/identity-read/system-groups/export.xlsx"
    ).content
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    ranges = [CellRange(str(item)) for item in workbook["Duplicate Groups"].merged_cells]
    for index, left in enumerate(ranges):
        assert left.min_row >= 2 and left.min_col <= 6
        for right in ranges[index + 1:]:
            assert not (
                left.min_row <= right.max_row and right.min_row <= left.max_row
                and left.min_col <= right.max_col and right.min_col <= left.max_col
            )


def test_r6_21_round_trip_retains_table_and_single_filter(db, client):
    scan_id, _ = _upload(client, db)
    payload = client.get(
        f"/api/scans/{scan_id}/identity-read/system-groups/export.xlsx"
    ).content
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    output = io.BytesIO()
    workbook.save(output)
    reopened = load_workbook(io.BytesIO(output.getvalue()), data_only=False)
    assert tuple(reopened["Group Data"].tables) == ("SystemGroupData",)
    assert reopened["Group Data"].auto_filter.ref is None
    round_trip = _package(output.getvalue())
    assert len(round_trip["table"].findall(f"{{{_MAIN_NS}}}autoFilter")) == 1
    assert round_trip["flat"].findall(f"{{{_MAIN_NS}}}autoFilter") == []


def test_r6_22_to_r6_25_api_csv_xlsx_exact_scan_parity(db, client):
    scan_id, _ = _upload(client, db)
    snapshot = IdentityReadService(db).load_identity_read_snapshot(scan_id)
    csv_response = client.get(
        f"/api/scans/{scan_id}/identity-read/system-groups/export.csv"
    )
    xlsx_response = client.get(
        f"/api/scans/{scan_id}/identity-read/system-groups/export.xlsx"
    )
    assert csv_response.status_code == xlsx_response.status_code == 200
    import csv
    csv_rows = list(csv.DictReader(io.StringIO(csv_response.text)))
    workbook = load_workbook(io.BytesIO(xlsx_response.content), data_only=False)
    headers = [cell.value for cell in workbook["Group Data"][1]]
    values = [dict(zip(headers, row, strict=True)) for row in workbook["Group Data"].iter_rows(
        min_row=2, values_only=True
    )]
    expected_keys = {
        serialize_versioned_identity_group_key(group.versioned_group_key)
        for group in snapshot.groups
    }
    assert {row["group_key"] for row in csv_rows} == expected_keys
    assert {row["Canonical Group ID"] for row in values} == expected_keys
    assert {
        row["stable_record_reference"] for row in csv_rows
    } == {
        str(row["Source Row / Stable Record Reference"]).split(" / ")[-1]
        for row in values
    }
    summary = {
        row[0].value: row[1].value
        for row in workbook["Summary"].iter_rows(min_row=2, max_col=2)
    }
    assert summary["Scan identifier"] == scan_id
    assert summary["Projection contract"] == "G2_V2"
