from pathlib import Path

from openpyxl import load_workbook

from app.benchmarks.human_identity_review_dataset import (
    CandidatePair,
    ReviewRecord,
    combine_panels,
    select_evaluation_panel,
    validate_artifacts,
    write_adjudication_workbook,
    write_internal_mapping,
    write_reviewer_workbook,
)


def record(index: int, value: str = "item") -> ReviewRecord:
    return ReviewRecord(index, f"ref-{index:04d}", index, f"PN-{index}", value, value, value, "", "", "PCS", "Purchased", "S1")


def pair(index: int, edge: str, shadow: str) -> CandidatePair:
    return CandidatePair(record(index * 2), record(index * 2 + 1), edge, shadow, edge == "CANNOT_LINK", "TEST", f"evidence-{index}")


def test_selection_is_deterministic_stratified_and_deduplicated():
    population = tuple(pair(i, "REVIEW_SUPPORT", "SHADOW_LEXICAL_ONLY_OR_UNRESOLVED") for i in range(10))
    targets = (("REVIEW_SUPPORT__SHADOW_LEXICAL_ONLY_OR_UNRESOLVED", 5),)
    first, counts = select_evaluation_panel(population, targets)
    second, _ = select_evaluation_panel(tuple(reversed(population)), targets)
    assert first == second
    assert counts == {"REVIEW_SUPPORT__SHADOW_LEXICAL_ONLY_OR_UNRESOLVED": {"population": 10, "sample": 5}}
    assert len({item.canonical_refs for item in first}) == 5


def test_panel_overlap_is_one_review_row_with_dual_membership():
    item = pair(1, "CANNOT_LINK", "SHADOW_MIXED_EVIDENCE")
    combined = combine_panels((item,), (item,))
    assert len(combined) == 1
    assert combined[0].panel_memberships == ("DIAGNOSTIC", "EVALUATION")


def test_reviewer_workbooks_are_blinded_empty_safe_and_consistent(tmp_path: Path):
    items = combine_panels((pair(1, "CANNOT_LINK", "SHADOW_MIXED_EVIDENCE"),), ())
    output = tmp_path
    write_reviewer_workbook(output / "human_identity_review_reviewer_a.xlsx", items, "REVIEWER_A")
    write_reviewer_workbook(output / "human_identity_review_reviewer_b.xlsx", items, "REVIEWER_B")
    write_internal_mapping(output / "human_identity_review_internal_mapping.csv", items)
    write_adjudication_workbook(output / "human_identity_review_adjudication.xlsx", items)
    result = validate_artifacts(output, 1)
    assert result["prefilled_label_count"] == 0
    assert result["detector_output_leakage_count"] == 0
    sheet = load_workbook(output / "human_identity_review_reviewer_a.xlsx", data_only=False)["Review"]
    assert not any(cell.data_type == "f" for row in sheet.iter_rows() for cell in row)


def test_formula_like_source_values_are_written_as_text(tmp_path: Path):
    unsafe = record(1, "=HYPERLINK(\"bad\")")
    item = CandidatePair(unsafe, record(2), "CANNOT_LINK", "SHADOW_MIXED_EVIDENCE", True, "TEST", "e")
    items = combine_panels((item,), ())
    path = tmp_path / "safe.xlsx"
    write_reviewer_workbook(path, items, "REVIEWER_A")
    sheet = load_workbook(path, data_only=False)["Review"]
    assert not any(cell.data_type == "f" for row in sheet.iter_rows() for cell in row)
    assert any(str(cell.value).startswith("'=") for row in sheet.iter_rows() for cell in row if cell.value)
