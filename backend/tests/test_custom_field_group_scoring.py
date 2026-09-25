"""Selected Purchase/Sales conditions and custom fields must reach the group-first pipeline
and the Excel export, not just the scan record."""
import io
import json

import openpyxl

from app.services.identity_read_xlsx_export_service import _SELECTED_HEADER_FILL

PUMP = b"Hydraulic Pump 10 bar"


def _custom(client, mode, label="Manufacturer Code"):
    response = client.post("/api/config/custom-fields", json={"display_label": label, "mode": mode})
    assert response.status_code == 200, response.text


def _scan(client, csv, selected, part_type="INVENTORY"):
    response = client.post(
        "/api/scans/upload",
        files={"file": ("parts.csv", csv, "text/csv")},
        data={"threshold": "50", "selected_fields": json.dumps(selected), "part_type": part_type},
    )
    assert response.status_code == 200, response.text
    return response.json()["scan_id"]


def _sheet(client, scan_id, name):
    response = client.get(f"/api/scans/{scan_id}/identity-read/system-groups/export.xlsx")
    assert response.status_code == 200
    rows = list(openpyxl.load_workbook(io.BytesIO(response.content))[name].iter_rows(values_only=True))
    header_at = next(i for i, row in enumerate(rows) if row and "Part Number" in row)
    header = [str(value) for value in rows[header_at]]
    return header, [row for row in rows[header_at + 1:] if row and row[header.index("Part Number")]]


def _members(client, scan_id):
    header, rows = _sheet(client, scan_id, "Review Groups")
    return sorted(row[header.index("Part Number")] for row in rows)


def _strength(client, scan_id):
    header, rows = _sheet(client, scan_id, "Review Groups")
    values = [row[header.index("Match Strength")] for row in rows if row[header.index("Match Strength")] is not None]
    return values[0] if values else None


def _two_parts(manufacturer_a, manufacturer_b):
    return (
        b"Part No,Item Description,Site,Manufacturer Code\n"
        b"A1," + PUMP + b",10," + manufacturer_a + b"\n"
        b"A2," + PUMP + b",10," + manufacturer_b + b"\n"
    )


def test_strict_custom_field_mismatch_keeps_parts_out_of_one_group(client):
    _custom(client, "STRICT")
    scan_id = _scan(client, _two_parts(b"ACME", b"OTHER"), ["MANUFACTURER_CODE"])
    assert _members(client, scan_id) == []


def test_strict_custom_field_match_still_groups(client):
    _custom(client, "STRICT")
    scan_id = _scan(client, _two_parts(b"ACME", b"ACME"), ["MANUFACTURER_CODE"])
    assert _members(client, scan_id) == ["A1", "A2"]


def test_strict_custom_field_applies_even_when_not_ticked(client):
    _custom(client, "STRICT")
    scan_id = _scan(client, _two_parts(b"ACME", b"OTHER"), [])
    assert _members(client, scan_id) == []


def test_supporting_custom_field_match_raises_strength_over_mismatch(client):
    _custom(client, "SUPPORTING")
    same = _strength(client, _scan(client, _two_parts(b"ACME", b"ACME"), ["MANUFACTURER_CODE"]))
    differ = _strength(client, _scan(client, _two_parts(b"ACME", b"OTHER"), ["MANUFACTURER_CODE"]))
    assert same is not None and differ is not None
    assert same > differ


def test_selected_purchase_condition_reaches_group_scoring(client):
    def csv(buyer_b):
        return (
            b"Part No,Item Description,Site,Buyer Id\n"
            b"A1," + PUMP + b",10,ALAIN\n"
            b"A2," + PUMP + b",10," + buyer_b + b"\n"
        )
    same = _strength(client, _scan(client, csv(b"ALAIN"), ["BUYER_ID"], "PURCHASE"))
    differ = _strength(client, _scan(client, csv(b"BOB"), ["BUYER_ID"], "PURCHASE"))
    assert same > differ


def test_custom_fields_do_not_change_scans_that_do_not_use_them(client):
    _custom(client, "STRICT")
    csv = b"Part No,Item Description,Site\nA1," + PUMP + b",10\nA2," + PUMP + b",10\n"
    assert _members(client, _scan(client, csv, [])) == ["A1", "A2"]


def _workbook(client, scan_id, kind="system-groups"):
    response = client.get(f"/api/scans/{scan_id}/identity-read/{kind}/export.xlsx")
    assert response.status_code == 200
    return openpyxl.load_workbook(io.BytesIO(response.content))


def test_custom_field_values_appear_as_selected_columns_in_the_workbook(client):
    _custom(client, "SUPPORTING")
    scan_id = _scan(client, _two_parts(b"ACME", b"ACME"), ["MANUFACTURER_CODE"])
    for sheet in ("Review Groups", "Detailed Data"):
        header, rows = _sheet(client, scan_id, sheet)
        column = header.index("Manufacturer Code")
        assert [row[column] for row in rows] == ["ACME", "ACME"]
    # Detailed Data lists selected conditions before the unselected Inventory columns;
    # Review Groups keeps its reviewer layout and adds them after the Inventory block.
    detailed, _ = _sheet(client, scan_id, "Detailed Data")
    assert detailed.index("Manufacturer Code") < detailed.index("Commodity Group 01")
    review, _ = _sheet(client, scan_id, "Review Groups")
    assert review.index("Site") < review.index("Manufacturer Code") < review.index("Review Consideration")


def test_custom_column_is_highlighted_like_other_selected_conditions(client):
    _custom(client, "SUPPORTING")
    scan_id = _scan(client, _two_parts(b"ACME", b"ACME"), ["MANUFACTURER_CODE"])
    sheet = _workbook(client, scan_id)["Detailed Data"]
    header = [cell.value for cell in sheet[1]]
    custom = sheet.cell(1, header.index("Manufacturer Code") + 1).fill.fgColor.rgb
    unselected = sheet.cell(1, header.index("Commodity Group 01") + 1).fill.fgColor.rgb
    assert custom.endswith(_SELECTED_HEADER_FILL.fgColor.rgb[-6:])
    assert custom != unselected


def test_strict_custom_column_is_shown_even_when_not_ticked(client):
    _custom(client, "STRICT")
    scan_id = _scan(client, _two_parts(b"ACME", b"ACME"), [])
    header, _rows = _sheet(client, scan_id, "Detailed Data")
    assert "Manufacturer Code" in header


def test_conditions_card_names_custom_fields_by_label(client):
    _custom(client, "STRICT")
    scan_id = _scan(client, _two_parts(b"ACME", b"ACME"), [])
    overview = _workbook(client, scan_id)["Overview"]
    text = " ".join(str(cell.value) for row in overview.iter_rows() for cell in row if cell.value)
    assert "Manufacturer Code" in text


def test_selected_purchase_conditions_get_columns_and_unselected_do_not(client):
    csv = (
        b"Part No,Item Description,Site,Buyer Id,Purchase Group\n"
        b"A1," + PUMP + b",10,ALAIN,PG1\n"
        b"A2," + PUMP + b",10,ALAIN,PG1\n"
    )
    scan_id = _scan(client, csv, ["BUYER_ID"], "PURCHASE")
    header, rows = _sheet(client, scan_id, "Detailed Data")
    assert [row[header.index("Buyer Id")] for row in rows] == ["ALAIN", "ALAIN"]
    assert "Purchase Group" not in header


def test_workbook_columns_are_unchanged_when_no_extra_conditions_are_used(client):
    csv = b"Part No,Item Description,Site\nA1," + PUMP + b",10\nA2," + PUMP + b",10\n"
    header, _rows = _sheet(client, _scan(client, csv, ["CONTRACT"]), "Detailed Data")
    assert header[5:8] == ["Part Number", "Description", "Site"]
    assert header[-1] == "HSN/SAC Code" and len(header) == 18


def test_reviewed_identities_workbook_carries_the_same_columns(client):
    _custom(client, "SUPPORTING")
    scan_id = _scan(client, _two_parts(b"ACME", b"ACME"), ["MANUFACTURER_CODE"])
    sheet = _workbook(client, scan_id, "reviewed-identities")["Detailed Data"]
    assert "Manufacturer Code" in [cell.value for cell in sheet[1]]
