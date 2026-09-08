from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.benchmarks.human_identity_review_dataset import (
    CandidatePair,
    ReviewRecord,
    write_reviewer_workbook,
)
from app.benchmarks.human_identity_review_execution import (
    HumanIdentityReviewExecutionError,
    create_disagreement_adjudication_workbook,
    validate_reviewer_submission,
)


def _record(index: int) -> ReviewRecord:
    return ReviewRecord(
        index, f"ref-{index}", index, f"PN-{index}", f"Item {index}",
        f"Item {index}", f"Item {index}", "", "", "PCS", "Purchased", "S1"
    )


def _pairs() -> tuple[CandidatePair, ...]:
    return tuple(
        CandidatePair(
            _record(index * 2), _record(index * 2 + 1), "REVIEW_SUPPORT",
            "SHADOW_LEXICAL_ONLY_OR_UNRESOLVED", False, "TEST", f"edge-{index}",
        )
        for index in range(2)
    )


def _workbooks(tmp_path: Path):
    a = tmp_path / "a-reference.xlsx"
    b = tmp_path / "b-reference.xlsx"
    write_reviewer_workbook(a, _pairs(), "REVIEWER_A")
    write_reviewer_workbook(b, _pairs(), "REVIEWER_B")
    return a, b


def _fill(path: Path, labels=("SAME_IDENTITY", "DIFFERENT_IDENTITY")) -> None:
    workbook = load_workbook(path)
    sheet = workbook["Review"]
    headers = [cell.value for cell in sheet[1]]
    for row, label in zip(range(2, sheet.max_row + 1), labels):
        sheet.cell(row, headers.index("identity_label") + 1, label)
        sheet.cell(row, headers.index("confidence") + 1, "HIGH")
        sheet.cell(row, headers.index("reason_code") + 1, "DESCRIPTION_EVIDENCE")
    workbook.save(path)


def test_blank_original_is_valid_draft_but_not_complete(tmp_path: Path):
    a, _b = _workbooks(tmp_path)
    result = validate_reviewer_submission(
        a, a, expected_reviewer_code="REVIEWER_A", declared_complete=False
    )
    assert result.pair_count == 2
    assert result.completed_label_count == 0
    with pytest.raises(HumanIdentityReviewExecutionError, match="MISSING_LABEL"):
        validate_reviewer_submission(
            a, a, expected_reviewer_code="REVIEWER_A", declared_complete=True
        )


def test_complete_submission_accepts_only_valid_independent_human_fields(tmp_path: Path):
    a, _b = _workbooks(tmp_path)
    submitted = tmp_path / "a-submitted.xlsx"
    submitted.write_bytes(a.read_bytes())
    _fill(submitted)
    result = validate_reviewer_submission(
        a, submitted, expected_reviewer_code="REVIEWER_A", declared_complete=True
    )
    assert result.completed_label_count == 2
    assert result.source_fields_unchanged is True


@pytest.mark.parametrize(
    ("mutation", "error"),
    (
        ("source", "SOURCE_FIELDS_CHANGED"),
        ("label", "LABEL_INVALID"),
        ("confidence", "CONFIDENCE_INVALID"),
        ("foreign", "FOREIGN_PAIR_ID"),
        ("duplicate", "DUPLICATE_PAIR_ID"),
    ),
)
def test_submission_fails_closed_on_integrity_changes(tmp_path: Path, mutation, error):
    a, _b = _workbooks(tmp_path)
    submitted = tmp_path / "submitted.xlsx"
    submitted.write_bytes(a.read_bytes())
    workbook = load_workbook(submitted)
    sheet = workbook["Review"]
    headers = [cell.value for cell in sheet[1]]
    if mutation == "source":
        sheet.cell(2, headers.index("record_a_description") + 1, "changed")
    elif mutation == "label":
        sheet.cell(2, headers.index("identity_label") + 1, "MAYBE")
    elif mutation == "confidence":
        sheet.cell(2, headers.index("confidence") + 1, "CERTAIN")
    elif mutation == "foreign":
        sheet.cell(2, headers.index("review_pair_id") + 1, "FOREIGN")
    else:
        sheet.cell(3, headers.index("review_pair_id") + 1, sheet.cell(2, 1).value)
    workbook.save(submitted)
    with pytest.raises(HumanIdentityReviewExecutionError, match=error):
        validate_reviewer_submission(
            a, submitted, expected_reviewer_code="REVIEWER_A", declared_complete=False
        )


def test_adjudication_package_contains_disagreements_only_and_no_detector_fields(tmp_path: Path):
    a_ref, b_ref = _workbooks(tmp_path)
    a = tmp_path / "a.xlsx"
    b = tmp_path / "b.xlsx"
    a.write_bytes(a_ref.read_bytes())
    b.write_bytes(b_ref.read_bytes())
    _fill(a, ("SAME_IDENTITY", "DIFFERENT_IDENTITY"))
    _fill(b, ("DIFFERENT_IDENTITY", "DIFFERENT_IDENTITY"))
    output = tmp_path / "adjudication.xlsx"

    assert create_disagreement_adjudication_workbook(a_ref, a, b_ref, b, output) == 1
    sheet = load_workbook(output, data_only=False)["Adjudication"]
    headers = tuple(cell.value for cell in sheet[1])
    assert sheet.max_row == 2
    assert "reviewer_a_label" in headers
    assert "reviewer_b_label" in headers
    assert "adjudicated_label" in headers
    assert "adjudicator_confidence" in headers
    assert not any(term in str(header).casefold() for header in headers for term in ("score", "gf4", "gf5", "llm", "expected"))
