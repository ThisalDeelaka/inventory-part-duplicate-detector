import csv
import io

import pandas as pd
from openpyxl import load_workbook
from unittest.mock import patch

from app.match_strength.contracts import MatchBand, MatchStrengthStatus
from app.match_strength.service import (
    MatchStrengthProjectionService,
    match_strength_distribution,
    match_strength_payload,
)
from app.services.identity_read_service import IdentityReadService
from app.services.scan_runner import ScanRunner
from test_group_first_scan_orchestration import configuration


def test_persisted_projection_scores_without_evaluator_rerun(db, client):
    records = pd.DataFrame([
        {"PART_NO": "A", "DESCRIPTION": "SKF 6205 BEARING", "CONTRACT": "S1", "UNIT_MEAS": "PCS"},
        {"PART_NO": "B", "DESCRIPTION": "SKF BEARING 6205", "CONTRACT": "S1", "UNIT_MEAS": "PCS"},
    ])
    scan, _ = ScanRunner(db, configuration("group_first_primary")).run(
        records, "Match Strength persisted projection", ["CONTRACT", "UNIT_MEAS"], 60
    )
    scan_id = scan.id
    snapshot = IdentityReadService(db).load_identity_read_snapshot(scan_id)
    with patch(
        "app.engine.identity_evidence_evaluator.evaluate_canonical_identity_relationship",
        side_effect=AssertionError("read projection must not evaluate relationships"),
    ):
        projected = MatchStrengthProjectionService(db).project_groups(snapshot.groups)
    assert projected
    assert all(result.status == MatchStrengthStatus.SCORED for result in projected.values())

    summary = client.get(f"/api/scans/{scan_id}/identity-read/summary")
    listing = client.get(f"/api/scans/{scan_id}/identity-read/groups")
    assert summary.status_code == listing.status_code == 200
    summary_payload = summary.json()
    listed = listing.json()["items"]
    assert sum(summary_payload["match_strength_distribution"].values()) == summary_payload["group_count"]
    assert all(item["match_strength_status"] == "SCORED" for item in listed)
    assert all(item["match_strength_version"] == "DETERMINISTIC_MATCH_STRENGTH_V2" for item in listed)

    detail = client.get(
        f"/api/scans/{scan_id}/identity-read/groups/{listed[0]['versioned_group_key']}"
    ).json()
    assert detail["match_strength"] == listed[0]["match_strength"]
    assert detail["match_band"] == listed[0]["match_band"]

    csv_response = client.get(
        f"/api/scans/{scan_id}/identity-read/system-groups/export.csv"
    )
    csv_rows = list(csv.DictReader(io.StringIO(csv_response.text)))
    assert csv_rows and all(row["match_strength"] for row in csv_rows)
    assert all(row["match_strength_version"] == "DETERMINISTIC_MATCH_STRENGTH_V2" for row in csv_rows)

    xlsx_response = client.get(
        f"/api/scans/{scan_id}/identity-read/system-groups/export.xlsx"
    )
    workbook = load_workbook(io.BytesIO(xlsx_response.content), data_only=False)
    assert workbook.sheetnames == [
        "Overview", "Review Groups", "Group Index", "Detailed Data",
        "Technical Reference",
    ]
    assert "Match Strength" in tuple(cell.value for cell in workbook["Group Index"][1])
    assert "Match Strength Version" in tuple(
        cell.value for cell in workbook["Technical Reference"][1]
    )
    assert not any(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row
    )


def test_distribution_and_high_review_crossover_are_independent():
    from app.match_strength import MatchStrengthEvidence, project_match_strength

    high_review = project_match_strength(
        ("a", "b"),
        (MatchStrengthEvidence("a", "b", "REVIEW_SUPPORT", 99.23, True),),
    )
    moderate = project_match_strength(
        ("c", "d"),
        (MatchStrengthEvidence("c", "d", "REVIEW_SUPPORT", 75.0, True),),
    )
    assert match_strength_distribution((high_review, moderate)) == {
        "HIGH_MATCH": 1, "MODERATE_MATCH": 1,
        "BORDERLINE_MATCH": 0, "UNSCORED": 0,
    }
    payload = match_strength_payload(
        high_review, group_status="POSSIBLE_DUPLICATE_GROUP_REVIEW"
    )
    assert high_review.match_band == MatchBand.HIGH_MATCH
    assert payload["safety_status_crossover"] is True
    assert "safety rules require review" in payload["safety_status_message"]
