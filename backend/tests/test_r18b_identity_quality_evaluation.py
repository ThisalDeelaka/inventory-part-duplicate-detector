import csv
import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.benchmarks.human_identity_review_dataset import (
    CandidatePair,
    ReviewRecord,
    write_reviewer_workbook,
)
from app.benchmarks.r18b_identity_quality_evaluation import (
    RecoveryResult,
    RecoveredRow,
    R18BRecoveryError,
    evaluate,
    recover_primary_reference,
    write_recovered_reference,
)


def _record(index):
    return ReviewRecord(
        index, f"ref-{index}", index, f"PN-{index}", f"Item {index}",
        f"Item {index}", f"Item {index}", "", "", "PCS", "Purchased", "S1",
    )


def _pairs(count=2):
    return tuple(
        CandidatePair(
            _record(index * 2), _record(index * 2 + 1), "REVIEW_SUPPORT",
            "SHADOW_LEXICAL_ONLY_OR_UNRESOLVED", False, "TEST", f"edge-{index}",
        )
        for index in range(count)
    )


def _prepare_recoverable(tmp_path: Path):
    original = tmp_path / "original.xlsx"
    submitted = tmp_path / "submitted.xlsx"
    write_reviewer_workbook(original, _pairs(), "REVIEWER_A")
    submitted.write_bytes(original.read_bytes())
    workbook = load_workbook(submitted)
    sheet = workbook["Review"]
    sheet.insert_rows(1)
    sheet.insert_cols(20, 2)
    sheet.cell(2, 20, "=I2=R2")
    sheet.cell(2, 21, "=J2=S2")
    for row in range(3, 5):
        sheet.cell(row, 20, f"=I{row}=R{row}")
        sheet.cell(row, 21, f"=J{row}=S{row}")
    headers = [cell.value for cell in sheet[2]]
    for row, label in zip(range(3, 5), ("SAME_IDENTITY", "DIFFERENT_IDENTITY")):
        sheet.cell(row, headers.index("identity_label") + 1, label)
        sheet.cell(row, headers.index("confidence") + 1, "HIGH")
        sheet.cell(row, headers.index("reason_code") + 1, "DESCRIPTION_EVIDENCE")
    sheet.cell(3, headers.index("reviewer_comment") + 1, "  literal note  ")
    workbook.save(submitted)
    source = load_workbook(original, read_only=True)["Review"]
    expected = {row[0].value for row in source.iter_rows(min_row=2, max_row=3)}
    return original, submitted, expected


def test_moved_header_and_extra_non_authoritative_formulas_are_recoverable(tmp_path: Path):
    original, submitted, expected = _prepare_recoverable(tmp_path)
    result = recover_primary_reference(original, submitted, expected)
    assert result.header_row == 2
    assert result.source_cells_compared == 36
    assert result.source_cells_different == 0
    assert len(result.extra_columns) == 2
    assert all(item["formula_cell_count"] == 3 for item in result.extra_columns)
    assert [row.identity_label for row in result.rows] == [
        "SAME_IDENTITY", "DIFFERENT_IDENTITY"
    ]
    assert result.rows[0].reviewer_comment == "  literal note  "


def test_pair_id_set_must_match_exactly(tmp_path: Path):
    original, submitted, expected = _prepare_recoverable(tmp_path)
    workbook = load_workbook(submitted)
    workbook["Review"].cell(3, 1, "R18-FOREIGN")
    workbook.save(submitted)
    with pytest.raises(R18BRecoveryError, match="PAIR_ID_INTEGRITY"):
        recover_primary_reference(original, submitted, expected)


def test_source_mismatch_fails_closed(tmp_path: Path):
    original, submitted, expected = _prepare_recoverable(tmp_path)
    workbook = load_workbook(submitted)
    sheet = workbook["Review"]
    headers = [cell.value for cell in sheet[2]]
    sheet.cell(3, headers.index("record_a_description") + 1, "altered")
    workbook.save(submitted)
    with pytest.raises(R18BRecoveryError, match="REVIEWER_VISIBLE_SOURCE_CHANGED"):
        recover_primary_reference(original, submitted, expected)


@pytest.mark.parametrize("header", ("record_a_description", "identity_label"))
def test_formula_in_authoritative_field_fails_closed(tmp_path: Path, header):
    original, submitted, expected = _prepare_recoverable(tmp_path)
    workbook = load_workbook(submitted)
    sheet = workbook["Review"]
    headers = [cell.value for cell in sheet[2]]
    sheet.cell(3, headers.index(header) + 1, "=1+1")
    workbook.save(submitted)
    with pytest.raises(R18BRecoveryError, match="FORMULA_CONTAMINATION"):
        recover_primary_reference(original, submitted, expected)


def test_invalid_human_label_fails_closed(tmp_path: Path):
    original, submitted, expected = _prepare_recoverable(tmp_path)
    workbook = load_workbook(submitted)
    sheet = workbook["Review"]
    headers = [cell.value for cell in sheet[2]]
    sheet.cell(3, headers.index("identity_label") + 1, "MAYBE")
    workbook.save(submitted)
    with pytest.raises(R18BRecoveryError, match="HUMAN_INPUT_INVALID"):
        recover_primary_reference(original, submitted, expected)


def test_derived_reference_is_deterministic(tmp_path: Path):
    original, submitted, expected = _prepare_recoverable(tmp_path)
    recovery = recover_primary_reference(original, submitted, expected)
    first = tmp_path / "first.csv"
    second = tmp_path / "second.csv"
    assert write_recovered_reference(first, recovery) == write_recovered_reference(second, recovery)
    assert first.read_bytes() == second.read_bytes()


def _result(rows):
    return RecoveryResult(
        rows=tuple(rows), header_row=2, logical_column_mapping={}, extra_columns=(),
        source_cells_compared=len(rows) * 18, source_cells_different=0,
        submission_sha256="a" * 64, original_template_sha256="b" * 64,
    )


def test_matrix_panel_confidence_stratum_and_generic_accounting(tmp_path: Path):
    labels = ("SAME_IDENTITY", "DIFFERENT_IDENTITY", "INSUFFICIENT_INFORMATION")
    classes = ("STRONG_SUPPORT", "REVIEW_SUPPORT", "NON_GROUPABLE", "CANNOT_LINK")
    primary_rows = []
    secondary_rows = []
    mapping_rows = []
    database = tmp_path / "evidence.db"
    connection = sqlite3.connect(database)
    connection.execute(
        "CREATE TABLE identity_evidence_edge_snapshot ("
        "evidence_run_id INTEGER, evidence_fingerprint TEXT, generic_evidence_json TEXT, "
        "technical_evidence_json TEXT, classification_reason_codes_json TEXT, "
        "protected_conflicts_json TEXT)"
    )
    index = 0
    for label in labels:
        for edge_class in classes:
            pair_id = f"R18-{index:016X}"
            primary_rows.append(RecoveredRow(
                pair_id, label, ("HIGH", "MEDIUM", "LOW")[index % 3],
                "OTHER", "",
            ))
            secondary_rows.append(RecoveredRow(
                pair_id, "DIFFERENT_IDENTITY", "MEDIUM", "DIFFERENT_OBJECT", "",
            ))
            mapping_rows.append({
                "review_pair_id": pair_id,
                "panel_membership": "DIAGNOSTIC|EVALUATION" if index < 2 else "EVALUATION",
                "diagnostic_family": f"CONTROL_{index}" if index < 2 else "",
                "evaluation_stratum": "CANNOT_LINK",
                "current_gf4_edge_class": edge_class,
                "r16_shadow_bucket": "SHADOW_LEXICAL_ONLY_OR_UNRESOLVED",
                "protected_conflict_present": str(edge_class == "CANNOT_LINK").lower(),
                "evidence_fingerprint": f"edge-{index}",
            })
            connection.execute(
                "INSERT INTO identity_evidence_edge_snapshot VALUES "
                "(7, ?, ?, '{}', '[]', '[]')",
                (f"edge-{index}", '{"generic_guard_reason":"GENERIC_DESCRIPTION"}' if index % 2 else '{}'),
            )
            index += 1
    connection.commit()
    connection.close()
    mapping = tmp_path / "mapping.csv"
    with mapping.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=tuple(mapping_rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(mapping_rows)
    output = tmp_path / "evaluation.csv"
    result = evaluate(_result(primary_rows), _result(secondary_rows), mapping, database, output)

    assert result["evaluation_pair_count"] == 12
    assert result["diagnostic_pair_count"] == 2
    assert result["panel_overlap_count"] == 2
    assert sum(sum(row.values()) for row in result["matrix_counts"].values()) == 12
    assert all(value == 1 for row in result["matrix_counts"].values() for value in row.values())
    assert result["secondary_review_status"] == "DEGRADED_ONE_CLASS_RESPONSE"
    assert result["sampling_stratum_counts"] == {"CANNOT_LINK": 12}
    assert result["generic_lexical_analysis"]["SAME_IDENTITY__HIGH"]
