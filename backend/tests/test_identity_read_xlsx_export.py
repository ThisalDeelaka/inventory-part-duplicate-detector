import csv
import dataclasses
import io
import json
from datetime import datetime
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from app.db.models import DuplicateScan
from app.identity_read.adapters import adapt_g2_v2_to_identity_read_snapshot
from app.identity_read.key_codec import serialize_versioned_identity_group_key
from app.llm.groq_provider import GroqLLMProvider
from app.services.identity_group_review_service import (
    GroupReviewDecision,
    VersionedIdentityGroupReviewService,
)
from app.services.identity_read_export_service import authority_selected_system_group_rows
from app.services.identity_read_service import (
    IdentityReadAuthorityInconsistent,
    IdentityReadService,
)
from app.services.identity_read_xlsx_export_service import (
    DETAILED_DATA_COLUMNS,
    GROUP_INDEX_COLUMNS,
    REVIEW_GROUP_COLUMNS,
    SHEET_ORDER,
    TECHNICAL_REFERENCE_COLUMNS,
    TECHNICAL_REFERENCE_HEADER_ROW,
    WORKBOOK_NOTICE,
    _member_pair_columns,
    _score_percentage,
    _write_review_groups,
    authority_selected_system_groups_to_xlsx,
    reason_for_group_status,
)
from test_group_first_backend_inversion import authoritative_group, review_scan
from test_identity_read_contracts import all_records, v2_source


def _workbook(content):
    return load_workbook(io.BytesIO(content), data_only=False)


def _rows(sheet):
    return list(sheet.iter_rows(min_row=2, values_only=True))


def _dict_rows(sheet):
    headers = tuple(cell.value for cell in sheet[1])
    return [dict(zip(headers, row, strict=True)) for row in _rows(sheet)]


def _technical_rows(sheet):
    headers = tuple(cell.value for cell in sheet[TECHNICAL_REFERENCE_HEADER_ROW])
    return [
        dict(zip(headers, row, strict=True))
        for row in sheet.iter_rows(
            min_row=TECHNICAL_REFERENCE_HEADER_ROW + 1, values_only=True
        )
    ]


def _add_scan(db, scan_id=21):
    db.add(DuplicateScan(
        id=scan_id,
        scan_name="XLSX synthetic",
        threshold=60,
        status="COMPLETED",
        model_version="deterministic-v1",
        selected_fields="[]",
    ))
    db.commit()


def test_client_workbook_contract_semantics_merges_and_review(db, client):
    scan = review_scan(db)
    snapshot, group = authoritative_group(db, scan)
    refs = tuple(member.stable_record_reference for member in group.members)
    long_comment = "Review deferred while the client validates equipment context. " * 4
    VersionedIdentityGroupReviewService(db).create_review(
        scan_id=scan.id,
        key=group.versioned_group_key,
        decision_type=GroupReviewDecision.UNSURE,
        reviewer="xlsx-reviewer",
        submitted_members=refs,
        comment=long_comment,
    )

    csv_response = client.get(
        f"/api/scans/{scan.id}/identity-read/system-groups/export.csv"
    )
    response = client.get(
        f"/api/scans/{scan.id}/identity-read/system-groups/export.xlsx"
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"] == (
        f'attachment; filename="scan-{scan.id}-system-groups.xlsx"'
    )
    workbook = _workbook(response.content)
    assert tuple(workbook.sheetnames) == SHEET_ORDER
    assert workbook.active.title == "Overview"

    overview = workbook["Overview"]
    overview_text = " ".join(
        str(cell.value) for row in overview.iter_rows() for cell in row
        if cell.value is not None
    )
    assert "Inventory Duplicate Review Report" in overview_text
    assert "Candidate groups generated for human review" in overview_text
    assert "SCAN INFORMATION" in overview_text
    assert "FINDINGS AT A GLANCE" in overview_text
    assert "ADDITIONAL FINDINGS REQUIRING ATTENTION" in overview_text
    assert "HUMAN REVIEW PROGRESS" in overview_text
    assert "HOW TO USE THIS WORKBOOK" in overview_text
    assert "REPORT DETAILS / TECHNICAL FOOTER" in overview_text
    assert WORKBOOK_NOTICE in overview_text
    assert overview["E6"].value == "Records Analysed"
    assert overview["E7"].value == snapshot.canonical_record_count
    assert overview["A19"].value == "Total Candidate Groups"
    assert overview["A20"].value == snapshot.group_count
    assert overview["A33"].value == "Reviewed"
    assert overview["G33"].value == "1 of 1"
    assert overview["A37"].value == "Deferred by Reviewer"
    assert overview["G37"].value == 1
    assert "Deferred / Unreviewed" not in overview_text
    assert overview["A49"].value == "Scan ID"
    assert overview["C49"].value == scan.id
    occupied = set()
    for merged in overview.merged_cells.ranges:
        for row_number in range(merged.min_row, merged.max_row + 1):
            for column_number in range(merged.min_col, merged.max_col + 1):
                coordinate = (row_number, column_number)
                assert coordinate not in occupied
                occupied.add(coordinate)

    review = workbook["Review Groups"]
    index = workbook["Group Index"]
    flat = workbook["Detailed Data"]
    technical = workbook["Technical Reference"]
    assert tuple(cell.value for cell in review[1]) == REVIEW_GROUP_COLUMNS
    assert REVIEW_GROUP_COLUMNS == (
        "Group", "Match Strength", "Match Band", "Group Sites",
        "Member Number", "Part Number", "Description", "Inventory UOM",
        "Part Type", "Commodity Group 01",
        "Commodity Group 02", "Safety Code", "Accounting Group",
        "Product Code", "Product Family", "Product Category", "HSN/SAC Code",
        "Site", "Review Consideration", "Why This Group Exists",
        "Relationship Evidence", "Part Relationships", "Pair Match Scores",
        "Description Similarity", "Wording Similarity", "Human Decision",
        "Human Comment",
    )
    assert "Evidence Tier" not in REVIEW_GROUP_COLUMNS
    assert REVIEW_GROUP_COLUMNS[:6] == (
        "Group", "Match Strength", "Match Band", "Group Sites",
        "Member Number", "Part Number",
    )
    assert REVIEW_GROUP_COLUMNS[-2:] == ("Human Decision", "Human Comment")
    assert REVIEW_GROUP_COLUMNS.count("Human Decision") == 1
    assert "Why This Group Exists" in REVIEW_GROUP_COLUMNS
    assert "Relationship Evidence" in REVIEW_GROUP_COLUMNS
    assert "were grouped because" in review.cell(
        2, REVIEW_GROUP_COLUMNS.index("Why This Group Exists") + 1
    ).value
    review_consideration = review.cell(
        2, REVIEW_GROUP_COLUMNS.index("Review Consideration") + 1
    ).value
    assert "Compare the" in review_consideration
    relationship_text = review.cell(
        2, REVIEW_GROUP_COLUMNS.index("Relationship Evidence") + 1
    ).value
    assert "%" in relationship_text
    assert "/100" not in relationship_text
    assert "Pair match:" in relationship_text
    assert "What matched:" in relationship_text
    assert "What to check:" not in relationship_text
    assert "Technical code:" not in relationship_text
    assert "Fuzzy lexical score" not in relationship_text
    assert "Inventory UOM relationship" not in relationship_text
    assert "Review Support" in relationship_text
    assert "caused" not in relationship_text.lower()
    assert tuple(cell.value for cell in index[1]) == GROUP_INDEX_COLUMNS
    assert tuple(cell.value for cell in flat[1]) == DETAILED_DATA_COLUMNS
    assert tuple(
        cell.value for cell in technical[TECHNICAL_REFERENCE_HEADER_ROW]
    ) == TECHNICAL_REFERENCE_COLUMNS
    assert review.freeze_panes == "G2"
    assert index.freeze_panes == "E2"
    assert flat.freeze_panes == "A2"
    assert technical.freeze_panes == f"A{TECHNICAL_REFERENCE_HEADER_ROW + 1}"
    assert review.cell(
        2, REVIEW_GROUP_COLUMNS.index("Match Strength") + 1
    ).number_format == '0.0"%"'
    assert index["C2"].number_format == '0.0"%"'
    assert not flat.merged_cells.ranges
    assert flat.auto_filter.ref is None
    assert flat.tables["SystemGroupData"].ref == (
        f"A1:{get_column_letter(len(DETAILED_DATA_COLUMNS))}{flat.max_row}"
    )

    csv_rows = list(csv.DictReader(io.StringIO(csv_response.text)))
    detail_rows = _dict_rows(flat)
    tech_rows = _technical_rows(technical)
    assert len(detail_rows) == len(tech_rows) == len(csv_rows) == group.member_count
    assert {row["Group"] for row in detail_rows} == {"CG-000001"}
    assert "Review Status" not in detail_rows[0]
    assert "Evidence" not in detail_rows[0]
    assert "Evidence Tier" not in detail_rows[0]
    assert {row["Human Decision"] for row in detail_rows} == {
        "Deferred for later review"
    }
    assert {row["Human Comment"] for row in detail_rows} == {long_comment}
    assert flat["E"][1].alignment.wrap_text is True
    assert max(flat.row_dimensions[row].height for row in range(2, flat.max_row + 1)) <= 42

    canonical = serialize_versioned_identity_group_key(group.versioned_group_key)
    assert {row["Canonical Group ID"] for row in tech_rows} == {canonical}
    assert {row["Stable Record Reference"] for row in tech_rows} == {
        row["stable_record_reference"] for row in csv_rows
    }
    assert {row["Source Row"] for row in tech_rows} == {
        int(row["source_row_reference"]) for row in csv_rows
    }
    for primary in (review, index, flat):
        assert "Canonical Group ID" not in tuple(cell.value for cell in primary[1])
        assert "Stable Record Reference" not in tuple(cell.value for cell in primary[1])

    expected_merges = {
        f"{get_column_letter(REVIEW_GROUP_COLUMNS.index(column) + 1)}2:"
        f"{get_column_letter(REVIEW_GROUP_COLUMNS.index(column) + 1)}"
        f"{group.member_count + 1}"
        for column in (
            "Group", "Match Strength", "Match Band", "Group Sites",
            "Review Consideration", "Why This Group Exists",
            "Relationship Evidence",
        )
    }
    assert {str(item) for item in review.merged_cells.ranges} == expected_merges
    assert [row["Member Number"] for row in _dict_rows(review)] == list(
        range(1, group.member_count + 1)
    )
    member_rows = _dict_rows(review)
    assert all(row["Part Relationships"] for row in member_rows)
    for field in (
        "Pair Match Scores", "Description Similarity", "Wording Similarity",
    ):
        assert all(row[field].endswith("%") for row in member_rows)
    why_text = review.cell(
        2, REVIEW_GROUP_COLUMNS.index("Why This Group Exists") + 1
    ).value
    assert "%" in why_text
    assert "/100" not in why_text
    assert not any(
        "/100" in cell.value
        for workbook_sheet in workbook.worksheets
        for row in workbook_sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    )
    freeze_column = review["G2"].column
    assert not any(
        merged.min_col < freeze_column <= merged.max_col
        for merged in review.merged_cells.ranges
    )
    pair_start = REVIEW_GROUP_COLUMNS.index("Part Relationships") + 1
    for column in range(pair_start, pair_start + 4):
        cell = review.cell(2, column)
        assert cell.alignment.wrap_text is True
        assert cell.alignment.vertical == "top"
        assert not any(merged.min_col <= column <= merged.max_col for merged in review.merged_cells.ranges)


def test_match_strength_xlsx_presentation_is_group_scoped_and_auditable(db):
    scan = review_scan(db)
    workbook = _workbook(authority_selected_system_groups_to_xlsx(db, scan.id))

    overview_text = " ".join(
        str(cell.value) for row in workbook["Overview"].iter_rows() for cell in row
        if cell.value is not None
    )
    assert "High Match (90–100)" in overview_text
    assert "Moderate Match (60–<90)" in overview_text
    assert "Borderline Match (0–<60)" in overview_text
    assert "not a probability of duplication" in overview_text
    assert "does not replace human review" in overview_text
    assert "Unscored" not in overview_text

    detailed_headers = tuple(cell.value for cell in workbook["Detailed Data"][1])
    assert "Evidence" not in detailed_headers
    assert "Evidence Tier" not in detailed_headers
    assert "Match Strength" not in detailed_headers
    assert "Match Band" not in detailed_headers
    assert detailed_headers == DETAILED_DATA_COLUMNS

    technical_text = " ".join(
        str(cell.value)
        for row in workbook["Technical Reference"].iter_rows(
            min_row=1, max_row=TECHNICAL_REFERENCE_HEADER_ROW - 1
        )
        for cell in row if cell.value is not None
    )
    for required in (
        "DETERMINISTIC_MATCH_STRENGTH_V2",
        "exact persisted deterministic supporting-edge score",
        "deterministic linear interpolation",
        "weakest_member_anchor",
        "support_density",
        "round_half_even",
        "High Match 90–100",
        "90 and 60 numeric boundaries",
        "authoritative signed Evidence Tier",
        "deterministic-group-explanation-v1",
        "deterministic-pair-explanation-read-model-v1",
        "deterministic-pair-explanation-v1",
        "PARTIAL_LEGACY",
        "No evaluator rerun and no LLM/provider call",
        "Created Date is not included",
        "not duplicate probability",
        "AI confidence",
        "human review decision",
        "GF4/GF5 authority input",
    ):
        assert required in technical_text

    assert tuple(workbook.sheetnames) == SHEET_ORDER
    assert not any(
        cell.data_type == "f"
        for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row
    )


def test_three_member_group_is_one_visual_block_with_distinct_members(db, monkeypatch):
    _add_scan(db)
    snapshot = adapt_g2_v2_to_identity_read_snapshot(
        v2_source(), all_records(), source_projection_run_id=202,
        source_orchestration_run_id=9,
    )
    monkeypatch.setattr(
        IdentityReadService,
        "load_identity_read_snapshot",
        lambda _self, _scan_id: snapshot,
    )
    workbook = _workbook(authority_selected_system_groups_to_xlsx(db, 21))
    sheet = workbook["Review Groups"]
    member_count = snapshot.groups[0].member_count
    assert member_count >= 3
    group_end_row = sheet.max_row
    group_columns = (
        "Group", "Match Strength", "Match Band", "Group Sites",
        "Review Consideration", "Why This Group Exists", "Relationship Evidence",
    )
    assert {
        f"{get_column_letter(REVIEW_GROUP_COLUMNS.index(column) + 1)}2:"
        f"{get_column_letter(REVIEW_GROUP_COLUMNS.index(column) + 1)}{group_end_row}"
        for column in group_columns
    }.issubset({str(item) for item in sheet.merged_cells.ranges})
    rows = _dict_rows(sheet)
    member_rows = [row for row in rows if row["Member Number"] is not None]
    assert len(member_rows) == member_count
    assert [row["Member Number"] for row in member_rows] == list(
        range(1, member_count + 1)
    )
    assert len({row["Part Number"] for row in member_rows}) == member_count
    for row in rows:
        evidence_values = (
            row["Pair Match Scores"], row["Description Similarity"],
            row["Wording Similarity"],
        )
        if row["Part Relationships"]:
            assert all(evidence_values)
        else:
            assert evidence_values == (
                "Not available", "Not available", "Not available",
            )
    pair_start = REVIEW_GROUP_COLUMNS.index("Part Relationships") + 1
    assert not any(
        merged.min_col <= pair_start + 3 and merged.max_col >= pair_start
        for merged in sheet.merged_cells.ranges
    )
    assert sheet.freeze_panes == "G2"
    assert not workbook["Detailed Data"].merged_cells.ranges
    assert workbook["Detailed Data"].max_row == member_count + 1


def test_member_pair_columns_are_incident_deterministic_and_score_aligned():
    assert _score_percentage(70.09) == "70.1%"
    assert _score_percentage(92.50) == "92.5%"
    assert _score_percentage(100.00) == "100.0%"
    assert _score_percentage(None) == "Not available"
    for member_count in (2, 3, 5):
        references = tuple(
            f"SITE-{index + 1}::record-{index + 1}"
            for index in range(member_count)
        )
        displays = tuple(
            f"DUP-100 — site {index + 1} pump"
            for index in range(member_count)
        )
        relationships = []
        pair_details = []
        expected_rows = {reference: [] for reference in references}
        for left in range(member_count):
            for right in range(left + 1, member_count):
                score = None if (left, right) == (0, 1) else 80 + left + right
                rendered_score = (
                    "Not available" if score is None else f"{score:.1f}%"
                )
                description = 0.0 if (left, right) == (0, 2) else 90 + left
                wording = None if (left, right) == (0, 1) else 95 + right
                relationships.append(SimpleNamespace(
                    relationship_id=f"pair-{left}-{right}",
                    left_record_reference=references[left],
                    right_record_reference=references[right],
                    left_display_identity=displays[left],
                    right_display_identity=displays[right],
                    deterministic_score=score,
                    signed_relationship="REVIEW_SUPPORT",
                ))
                supporting_items = [SimpleNamespace(
                    source_field="component_scores.description_similarity",
                    numeric_value=description,
                )]
                if wording is not None:
                    supporting_items.append(SimpleNamespace(
                        source_field="component_scores.fuzzy_score",
                        numeric_value=wording,
                    ))
                pair_details.append(SimpleNamespace(
                    relationship_id=f"pair-{left}-{right}",
                    supporting_items=tuple(supporting_items),
                ))
                expected = (
                    rendered_score,
                    (
                        "Not available"
                        if member_count == 2 else f"{description:.1f}%"
                    ),
                    (
                        "Not available"
                        if member_count == 2 or wording is None
                        else f"{wording:.1f}%"
                    ),
                )
                expected_rows[references[left]].append((references[right], expected))
                expected_rows[references[right]].append((references[left], expected))

        if member_count == 2:
            pair_details = []

        rendered = _member_pair_columns(SimpleNamespace(
            relationships=tuple(reversed(relationships)),
            pair_explanations=tuple(pair_details),
        ))
        assert set(rendered) == set(references)
        for reference in references:
            pair_rows = rendered[reference]
            assert len(pair_rows) == member_count - 1
            assert [row[1:] for row in pair_rows] == [
                expected for _other, expected in sorted(expected_rows[reference])
            ]
        if member_count == 3:
            zero_rows = [
                row for rows in rendered.values() for row in rows
                if row[2] == "0.0%"
            ]
            assert len(zero_rows) == 2


def test_member_pair_columns_use_stable_identity_not_duplicate_display_text():
    shared_display = "DUP-200 — shared pump with a deliberately long persisted display description " * 2
    explanation = SimpleNamespace(relationships=(
        SimpleNamespace(
            relationship_id="site-a-site-b",
            left_record_reference="SITE-A::row-1",
            right_record_reference="SITE-B::row-1",
            left_display_identity=shared_display,
            right_display_identity=shared_display,
            deterministic_score=92.5,
            signed_relationship="STRONG_SUPPORT",
        ),
        SimpleNamespace(
            relationship_id="not-supporting",
            left_record_reference="SITE-A::row-1",
            right_record_reference="SITE-C::row-1",
            left_display_identity=shared_display,
            right_display_identity=shared_display,
            deterministic_score=99.0,
            signed_relationship="CANNOT_LINK",
        ),
    ), pair_explanations=(SimpleNamespace(
        relationship_id="site-a-site-b",
        supporting_items=(),
    ),))

    rendered = _member_pair_columns(explanation)
    assert set(rendered) == {"SITE-A::row-1", "SITE-B::row-1"}
    assert len(rendered["SITE-A::row-1"][0][0]) > 100
    assert rendered["SITE-A::row-1"][0][1:] == (
        "92.5%", "Not available", "Not available",
    )
    assert rendered["SITE-B::row-1"][0][1:] == (
        "92.5%", "Not available", "Not available",
    )


def test_three_member_pair_rows_align_and_merge_without_changing_group_evidence():
    workbook = Workbook()
    sheet = workbook.active
    references = ("SITE-A::1", "SITE-B::1", "SITE-C::1")
    long_pair = "DUP-1 — pump with a deliberately long relationship label " * 3
    pair_rows = {
        references[0]: (
            (long_pair, "92.5%", "95.0%", "100.0%"),
            ("DUP-1 ↔ DUP-3", "90.0%", "0.0%", "98.5%"),
        ),
        references[1]: (
            ("DUP-2 ↔ DUP-1", "92.5%", "95.0%", "100.0%"),
            ("DUP-2 ↔ DUP-3", "91.0%", "Not available", "99.0%"),
        ),
        references[2]: (
            ("DUP-3 ↔ DUP-1", "90.0%", "0.0%", "98.5%"),
            ("DUP-3 ↔ DUP-2", "91.0%", "Not available", "99.0%"),
        ),
    }
    relationship_evidence = "persisted group relationship text — unchanged"
    presentation = {
        "label": "CG-000001",
        "evidence": "Review Evidence",
        "match_strength": 91.0,
        "match_band": "High Match",
        "sites": "SITE-A, SITE-B, SITE-C",
        "review_consideration": "Compare item details.",
        "deterministic_group_summary": "Three records have supporting pairs.",
        "relationship_evidence": relationship_evidence,
        "human_decision": "Not yet reviewed",
        "human_comment": "",
    }
    member_rows = tuple(
        {
            "stable_record_reference": reference,
            "part_no": "DUP-1",
            "description": f"Pump at site {index + 1}",
            "site_or_contract": f"SITE-{chr(65 + index)}",
        }
        for index, reference in enumerate(references)
    )

    _write_review_groups(sheet, ({
        "presentation": presentation,
        "member_rows": member_rows,
        "member_pair_columns": pair_rows,
    },))

    assert sheet.max_row == 7
    relationship_column = REVIEW_GROUP_COLUMNS.index("Relationship Evidence") + 1
    assert sheet.cell(2, relationship_column).value == relationship_evidence
    merges = {str(item) for item in sheet.merged_cells.ranges}
    group_columns = (
        "Group", "Match Strength", "Match Band", "Group Sites",
        "Review Consideration", "Why This Group Exists", "Relationship Evidence",
    )
    assert {
        f"{get_column_letter(REVIEW_GROUP_COLUMNS.index(column) + 1)}2:"
        f"{get_column_letter(REVIEW_GROUP_COLUMNS.index(column) + 1)}7"
        for column in group_columns
    }.issubset(merges)
    member_column = REVIEW_GROUP_COLUMNS.index("Member Number") + 1
    assert sum(
        merged.min_col == member_column and merged.max_row - merged.min_row == 1
        for merged in sheet.merged_cells.ranges
    ) == 3
    pair_start = REVIEW_GROUP_COLUMNS.index("Part Relationships") + 1
    decision_column = REVIEW_GROUP_COLUMNS.index("Human Decision") + 1
    comment_column = REVIEW_GROUP_COLUMNS.index("Human Comment") + 1
    assert decision_column == len(REVIEW_GROUP_COLUMNS) - 1
    assert comment_column == len(REVIEW_GROUP_COLUMNS)
    # Each member gets its own Human Decision / Human Comment cell.
    for column in (decision_column, comment_column):
        assert sum(
            merged.min_col == column and merged.max_row - merged.min_row == 1
            for merged in sheet.merged_cells.ranges
        ) == 3
    validations = sheet.data_validations.dataValidation
    assert len(validations) == 1
    assert validations[0].formula1 == '"Original Part,Duplicate Part,Valid Duplicate"'
    assert str(validations[0].sqref) == (
        f"{get_column_letter(decision_column)}2:{get_column_letter(decision_column)}7"
    )
    for row_number in range(2, 8):
        values = tuple(
            sheet.cell(row_number, column).value
            for column in range(pair_start, pair_start + 4)
        )
        assert all(value not in (None, "") for value in values)
        assert sheet.cell(row_number, pair_start).alignment.wrap_text is True
        assert sheet.row_dimensions[row_number].height <= 60
    assert not any(
        merged.min_col <= pair_start + 3 and merged.max_col >= pair_start
        for merged in sheet.merged_cells.ranges
    )
    assert sheet.freeze_panes == "G2"


def test_unreviewed_candidate_requires_human_review_and_reason_is_concise(db, client):
    scan = review_scan(db)
    workbook = _workbook(client.get(
        f"/api/scans/{scan.id}/identity-read/system-groups/export.xlsx"
    ).content)
    rows = _dict_rows(workbook["Group Index"])
    assert all("Review Status" not in row for row in rows)
    assert {row["Human Decision"] for row in rows} == {"Not yet reviewed"}
    assert all("Why Suggested" not in row for row in rows)
    assert all(
        "enough matching information to review these records together"
        in row["Review Consideration"]
        for row in rows
    )
    assert all(
        "persisted evaluator" not in row["Review Consideration"]
        for row in rows
    )
    overview = workbook["Overview"]
    assert overview["A33"].value == "Reviewed"
    assert overview["G33"].value == "0 of 1"
    assert overview["A34"].value == "Awaiting Review"
    assert overview["G34"].value == 1
    assert overview["A37"].value == "Deferred by Reviewer"
    assert overview["G37"].value == 0
    assert overview["E28"].value == "Deferred Families"


def test_repeated_generation_is_semantically_and_visually_deterministic(db):
    scan = review_scan(db)
    generated = [
        _workbook(authority_selected_system_groups_to_xlsx(db, scan.id))
        for _ in range(3)
    ]
    for sheet_name in ("Review Groups", "Group Index", "Detailed Data", "Technical Reference"):
        values = [list(book[sheet_name].iter_rows(values_only=True)) for book in generated]
        assert values[0] == values[1] == values[2]
    merges = [
        tuple(sorted(str(item) for item in book["Review Groups"].merged_cells.ranges))
        for book in generated
    ]
    assert merges[0] == merges[1] == merges[2]
    assert generated[0]["Technical Reference"][
        f"I{TECHNICAL_REFERENCE_HEADER_ROW + 1}"
    ].value == reason_for_group_status(
        "POSSIBLE_DUPLICATE_GROUP_REVIEW"
    )


def test_overview_uses_persisted_scan_metadata_and_ui_ordered_condition_labels(db):
    scan = review_scan(db)
    scan.started_at = datetime(2026, 9, 14, 8, 0, 0)
    scan.completed_at = datetime(2026, 9, 14, 12, 34, 56)
    scan.total_records = 5327
    scan.selected_fields = json.dumps(["UNIT_MEAS", "CONTRACT"])
    db.commit()

    overview = _workbook(
        authority_selected_system_groups_to_xlsx(db, scan.id)
    )["Overview"]
    assert overview["A7"].value == "14 Sep 2026, 12:34"
    assert overview["A11"].value == "Site\nInventory UOM"
    assert overview["E7"].value == 5327
    assert overview["E11"].value == "IFS APP Test"

    scan.completed_at = None
    db.commit()
    overview = _workbook(
        authority_selected_system_groups_to_xlsx(db, scan.id)
    )["Overview"]
    assert overview["A7"].value == "14 Sep 2026, 08:00"

    for selected_fields, expected in (
        (["UNIT_MEAS"], "Inventory UOM"),
        ([], "None selected"),
        (["FUTURE_CLIENT_FIELD"], "Future Client Field"),
    ):
        scan.selected_fields = json.dumps(selected_fields)
        db.commit()
        overview = _workbook(
            authority_selected_system_groups_to_xlsx(db, scan.id)
        )["Overview"]
        assert overview["A11"].value == expected
        assert "Site" not in str(overview["A11"].value)


def test_overview_findings_match_authoritative_projection_counts(db):
    scan = review_scan(db)
    snapshot = IdentityReadService(db).load_identity_read_snapshot(scan.id)
    overview = _workbook(
        authority_selected_system_groups_to_xlsx(db, scan.id)
    )["Overview"]
    findings = {
        overview["A19"].value: overview["A20"].value,
        overview["C19"].value: overview["C20"].value,
        overview["F19"].value: overview["F20"].value,
        overview["A23"].value: overview["A24"].value,
        overview["E23"].value: overview["E24"].value,
        overview["A28"].value: overview["A29"].value,
        overview["E28"].value: overview["E29"].value,
    }
    assert findings == {
        "Total Candidate Groups": snapshot.group_count,
        "Stronger Evidence": snapshot.likely_group_count,
        "Review Evidence": snapshot.review_group_count,
        "Conflicting Families": snapshot.conflict_count,
        "Deferred Families": snapshot.deferred_count,
        "Records in Candidate Groups": sum(
            group.member_count for group in snapshot.groups
        ),
        "Unassigned Records": snapshot.unassigned_count,
    }


def test_xlsx_export_does_not_mutate_request_or_group_semantics(db):
    scan = review_scan(db)
    before = IdentityReadService(db).load_identity_read_snapshot(scan.id)
    selected_fields = scan.selected_fields
    authority_selected_system_groups_to_xlsx(db, scan.id)
    db.expire_all()
    after = IdentityReadService(db).load_identity_read_snapshot(scan.id)

    assert db.get(DuplicateScan, scan.id).selected_fields == selected_fields
    assert after.snapshot_fingerprint == before.snapshot_fingerprint
    assert [
        (group.status, tuple(member.stable_record_reference for member in group.members))
        for group in after.groups
    ] == [
        (group.status, tuple(member.stable_record_reference for member in group.members))
        for group in before.groups
    ]
    assert after.conflicts == before.conflicts
    assert after.deferred_work_units == before.deferred_work_units
    assert after.unassigned_records == before.unassigned_records


def test_data_level_projection_is_logically_equivalent(db):
    scan = review_scan(db)
    _, raw_rows = authority_selected_system_group_rows(db, scan.id)
    workbook = _workbook(authority_selected_system_groups_to_xlsx(db, scan.id))
    details = _dict_rows(workbook["Detailed Data"])
    technical = _technical_rows(workbook["Technical Reference"])
    assert len(raw_rows) == len(details) == len(technical)
    labels_by_key = {}
    for raw in raw_rows:
        labels_by_key.setdefault(
            raw["group_key"], f"CG-{len(labels_by_key) + 1:06d}"
        )
    for raw, detail, tech in zip(raw_rows, details, technical, strict=True):
        assert detail["Group"] == labels_by_key[raw["group_key"]]
        assert tech["Group"] == detail["Group"]
        assert tech["Canonical Group ID"] == raw["group_key"]
        assert tech["Stable Record Reference"] == raw["stable_record_reference"]
        assert tech["Source Row"] == raw["source_row_reference"]
        assert tech["Part Number"] == raw["part_no"] == detail["Part Number"]
        assert detail["Description"] == raw["description"]
        assert detail["Site"] == raw["site_or_contract"]
        assert detail["Inventory UOM"] == raw["uom"]
        assert detail["Part Type"] == raw["part_type"]
        assert detail["Commodity Group 01"] == raw["commodity_group_01"]
        assert detail["Commodity Group 02"] == raw["commodity_group_02"]
        assert detail["Safety Code"] == raw["safety_code"]
        assert detail["Accounting Group"] == raw["accounting_group"]
        assert detail["Product Code"] == raw["product_code"]
        assert detail["Product Family"] == raw["product_family"]
        assert detail["Product Category"] == raw["product_category"]
        assert detail["HSN/SAC Code"] == raw["hsn_sac"]


def test_empty_state_is_friendly_and_structurally_valid(db, monkeypatch):
    _add_scan(db)
    snapshot = adapt_g2_v2_to_identity_read_snapshot(
        v2_source(), all_records(), source_projection_run_id=202,
        source_orchestration_run_id=9,
    )
    empty = dataclasses.replace(
        snapshot, groups=(), group_count=0, likely_group_count=0,
        review_group_count=0,
    )
    monkeypatch.setattr(
        IdentityReadService,
        "load_identity_read_snapshot",
        lambda _self, _scan_id: empty,
    )
    workbook = _workbook(authority_selected_system_groups_to_xlsx(db, 21))
    assert tuple(workbook.sheetnames) == SHEET_ORDER
    assert workbook["Overview"]["A20"].value == 0
    assert workbook["Review Groups"]["A2"].value == (
        "No candidate groups were generated for this scan."
    )
    assert workbook["Group Index"].max_row == 1
    assert workbook["Detailed Data"].max_row == 1
    assert workbook["Technical Reference"].max_row == TECHNICAL_REFERENCE_HEADER_ROW


def test_formula_like_inventory_values_remain_literal_and_source_immutable(
    db, client, monkeypatch
):
    _add_scan(db)
    original = adapt_g2_v2_to_identity_read_snapshot(
        v2_source(), all_records(), source_projection_run_id=202,
        source_orchestration_run_id=9,
    )
    original_member = original.groups[0].members[0]
    changed_member = dataclasses.replace(
        original_member, part_no="=1+1", description="+ABC", contract="-XYZ", uom="@PART"
    )
    changed_group = dataclasses.replace(
        original.groups[0], members=(changed_member,) + original.groups[0].members[1:]
    )
    snapshot = dataclasses.replace(original, groups=(changed_group,))
    monkeypatch.setattr(
        IdentityReadService,
        "load_identity_read_snapshot",
        lambda _self, _scan_id: snapshot,
    )
    response = client.get("/api/scans/21/identity-read/system-groups/export.xlsx")
    assert response.status_code == 200
    row = _workbook(response.content)["Detailed Data"][2]
    for index, expected in ((5, "=1+1"), (6, "+ABC"), (7, "-XYZ"), (8, "@PART")):
        assert row[index].value == expected
        assert row[index].data_type == "s"
    assert original_member.part_no != changed_member.part_no


def test_failed_and_incomplete_orchestration_fail_closed(db, client):
    for status in ("FAILED", "RUNNING"):
        scan = review_scan(db)
        db.execute(text(
            "UPDATE scan_orchestration_run SET status = :status WHERE scan_id = :scan_id"
        ), {"status": status, "scan_id": scan.id})
        db.commit()
        assert client.get(
            f"/api/scans/{scan.id}/identity-read/system-groups/export.xlsx"
        ).status_code == 409


def test_corrupt_authority_fails_closed(db, client, monkeypatch):
    scan = review_scan(db)
    monkeypatch.setattr(
        IdentityReadService,
        "load_identity_read_snapshot",
        lambda _self, _scan_id: (_ for _ in ()).throw(
            IdentityReadAuthorityInconsistent("corrupt persisted authority")
        ),
    )
    assert client.get(
        f"/api/scans/{scan.id}/identity-read/system-groups/export.xlsx"
    ).status_code == 422


def test_authority_missing_projection_is_not_ready(db, client):
    scan = review_scan(db)
    db.execute(text(
        "UPDATE scan_orchestration_stage_result SET source_run_reference = NULL "
        "WHERE orchestration_run_id = (SELECT id FROM scan_orchestration_run "
        "WHERE scan_id = :scan_id) AND stage_id = 'G2_V2_PROJECTION'"
    ), {"scan_id": scan.id})
    db.commit()
    response = client.get(
        f"/api/scans/{scan.id}/identity-read/system-groups/export.xlsx"
    )
    assert response.status_code == 409
    assert "G2_V2_PROJECTION_MISSING" in response.json()["detail"]


def test_cross_scan_membership_is_not_combined(db, client):
    first = review_scan(db)
    second = review_scan(db)
    _, first_group = authoritative_group(db, first)
    _, second_group = authoritative_group(db, second)
    workbook = _workbook(client.get(
        f"/api/scans/{second.id}/identity-read/system-groups/export.xlsx"
    ).content)
    canonical_ids = {
        row["Canonical Group ID"]
        for row in _technical_rows(workbook["Technical Reference"])
    }
    assert canonical_ids == {
        serialize_versioned_identity_group_key(second_group.versioned_group_key)
    }
    assert serialize_versioned_identity_group_key(
        first_group.versioned_group_key
    ) not in canonical_ids


def test_provider_calls_are_zero(db, client, monkeypatch):
    scan = review_scan(db)

    async def forbidden_provider_call(*_args, **_kwargs):
        raise AssertionError("XLSX export must not call an LLM provider")

    monkeypatch.setattr(GroqLLMProvider, "complete_json", forbidden_provider_call)
    assert client.get(
        f"/api/scans/{scan.id}/identity-read/system-groups/export.xlsx"
    ).status_code == 200


def test_no_unsupported_claims_secrets_formulas_or_writeback(db, client):
    scan = review_scan(db)
    workbook = _workbook(client.get(
        f"/api/scans/{scan.id}/identity-read/system-groups/export.xlsx"
    ).content)
    forbidden = (
        "ground truth", "api key", "credential", "password", "merge target",
        "likely duplicate", "confirmed duplicate", "high-confidence duplicate",
        "accuracy", "confidence %",
    )
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f"
                if isinstance(cell.value, str):
                    assert not any(term in cell.value.lower() for term in forbidden)
                    if "probability" in cell.value.lower():
                        assert (
                            "not a probability" in cell.value.lower()
                            or "not duplicate probability" in cell.value.lower()
                        )


def test_selected_matching_columns_are_green_and_others_default(db):
    scan = review_scan(db)
    scan.selected_fields = json.dumps(["CONTRACT", "UNIT_MEAS"])
    db.commit()
    workbook = _workbook(authority_selected_system_groups_to_xlsx(db, scan.id))
    for sheet_name in ("Review Groups", "Detailed Data"):
        sheet = workbook[sheet_name]
        headers = {cell.value: cell for cell in sheet[1]}
        for selected in ("Part Number", "Description", "Site", "Inventory UOM"):
            header = headers[selected]
            assert header.fill.fgColor.rgb.endswith("548235")
            assert sheet.cell(2, header.column).fill.fgColor.rgb.endswith("C6EFCE")
        for other in (
            "Part Type", "Commodity Group 01", "Safety Code", "HSN/SAC Code",
        ):
            header = headers[other]
            assert header.fill.fgColor.rgb.endswith("1F4E78")
            assert not sheet.cell(2, header.column).fill.fgColor.rgb.endswith("C6EFCE")


def test_part_type_is_green_only_when_its_condition_is_selected(db):
    scan = review_scan(db)
    scan.selected_fields = json.dumps(["TYPE_CODE"])
    db.commit()
    sheet = _workbook(
        authority_selected_system_groups_to_xlsx(db, scan.id)
    )["Review Groups"]
    headers = {cell.value: cell for cell in sheet[1]}
    assert headers["Part Type"].fill.fgColor.rgb.endswith("548235")
    assert headers["Site"].fill.fgColor.rgb.endswith("1F4E78")


def test_overview_explains_green_and_blue_columns(db):
    scan = review_scan(db)
    overview = _workbook(
        authority_selected_system_groups_to_xlsx(db, scan.id)
    )["Overview"]
    assert overview["A10"].value == "Duplicate-checking Conditions"
    assert overview["A14"].value == "COLUMN COLOUR KEY"
    assert overview["A15"].value == "Green"
    assert overview["A15"].fill.fgColor.rgb.endswith("548235")
    assert "Duplicate-checking condition" in overview["C15"].value
    assert overview["A16"].value == "Blue"
    assert overview["A16"].fill.fgColor.rgb.endswith("1F4E78")
    assert "not used as duplicate-checking conditions" in overview["C16"].value
    assert overview["A18"].value == "FINDINGS AT A GLANCE"


def test_part_relationships_show_part_numbers_only():
    explanation = SimpleNamespace(relationships=(
        SimpleNamespace(
            relationship_id="a-b",
            left_record_reference="S::1",
            right_record_reference="S::2",
            left_display_identity="LITHIUM-ION BATTERY — Lithium-Ion Battery",
            right_display_identity="LITHIUM-ION BATTERY 02-1 — LITHIUM-ION BATTERY 02-1",
            deterministic_score=91.0,
            signed_relationship="STRONG_SUPPORT",
        ),
    ), pair_explanations=())
    rendered = _member_pair_columns(explanation, {
        "S::1": "LITHIUM-ION BATTERY", "S::2": "LITHIUM-ION BATTERY 02-1",
    })
    assert rendered["S::1"][0][0] == "LITHIUM-ION BATTERY ↔ LITHIUM-ION BATTERY 02-1"
    assert rendered["S::2"][0][0] == "LITHIUM-ION BATTERY 02-1 ↔ LITHIUM-ION BATTERY"
